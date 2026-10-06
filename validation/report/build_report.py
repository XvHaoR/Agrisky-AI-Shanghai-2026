from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SUMMARY_PATH = ROOT / "validation" / "results" / "summary_metrics.json"
FIGURES = ROOT / "validation" / "figures"
OUTPUT = Path(__file__).resolve().parent / "穹野智保公开数据与文献验证报告.docx"

NAVY = "17315F"
TEAL = "087F6B"
BLUE = "2E74B5"
LIGHT = "F2F6F8"
MID = "D8E3EA"
INK = "152238"
MUTED = "5E6F82"
GOLD = "A26700"


def set_font(run, size=10.5, bold=False, color=INK, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        add_text(paragraph, bold_prefix, bold=True)
        add_text(paragraph, text[len(bold_prefix) :])
    else:
        add_text(paragraph, text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.38)
    paragraph.paragraph_format.first_line_indent = Inches(-0.19)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2
    add_text(paragraph, text)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    add_text(paragraph, text, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=BLUE if level < 3 else NAVY)
    return paragraph


def add_callout(doc, label, text, fill=LIGHT, accent=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    add_text(p, f"{label}  ", bold=True, color=accent)
    add_text(p, text)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, MID)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, header, size=9, bold=True, color=NAVY)
    mark_header_row(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, str(value), size=9)
    set_table_geometry(table, widths)
    return table


def add_figure(doc, filename, caption, width=6.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    picture = run.add_picture(str(FIGURES / filename), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", filename)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    add_text(cap, caption, size=9, color=MUTED)


def pct(value):
    return f"{value * 100:.1f}%"


def build_document():
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    production = summary["test"]["production"]
    calibrated = summary["test"]["calibrated"]
    fusion = summary["test"]["sar_feature_fusion_v2"]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after in ((1, 16, 16, 8), (2, 13, 12, 6), (3, 11.5, 8, 4)):
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "穹野智保 Agrisky AI | 公开数据与文献验证", size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "复赛验证材料 | 2026-08-29", size=8, color=MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(kicker, "AGRISKY AI · SEMIFINAL VALIDATION", size=10, bold=True, color=TEAL)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(10)
    add_text(title, "穹野智保公开数据与文献验证报告", size=25, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    add_text(subtitle, "农业保险遥感理赔辅助系统 · SAR 灾后水体识别复验", size=13, color=MUTED)
    add_callout(
        doc,
        "核心结论",
        "在 Sen1Floods11 官方独立测试划分抽取的 30 个灾后样本上，SAR 特征融合 v2 获得微观 F1 79.6%、IoU 66.1%，水体面积比例中位绝对误差 2.09 个百分点。结果支持其作为人工复核前的水体证据提取候选方案，但不构成作物损失或赔款准确性的业务验证。",
    )
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["报告版本", "v1.0"],
            ["复验日期", "2026-08-29"],
            ["代码基线", summary["repository_commit"][:12]],
            ["随机种子", summary["seed"]],
            ["数据集", "Sen1Floods11 v1.1 手工标注子集"],
            ["测试规模", "30 个官方 test 样本，测试集不参与训练或阈值选择"],
        ],
        [1900, 7460],
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    add_text(p, "穹野智保 Agrisky AI 项目组", size=11, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "GOAI 2026「AI+金融」赛道复赛支撑材料", size=10, color=MUTED)
    p.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "摘要", 1)
    add_body(doc, "本报告面向穹野智保 Agrisky AI 复赛阶段的定损准确性与业务价值论证。由于短期内无法完成保险机构真实理赔案件的地面真值采集，本阶段采用公开数据复验、权威文献交叉验证和可审计的业务情景测算三层证据结构。")
    add_body(doc, "公开复验使用 Sen1Floods11 v1.1 手工标注洪水事件子集，严格保留官方 train/validation/test 边界。100 个 train 样本用于训练，20 个 validation 样本用于选择概率阈值，30 个 test 样本仅用于最终一次评估。生产基线复现项目当前 Sentinel-1 VH 中值滤波与 -16 dB 固定阈值逻辑；候选 v2 在 VV、VH、极化差及多尺度局部统计特征上训练 ExtraTrees 分类器。")
    add_body(doc, "独立测试结果显示，v2 相比生产基线将微观 F1 从 52.7% 提升至 79.6%，微观 IoU 从 35.7% 提升至 66.1%，面积比例中位绝对误差从 30.59 个百分点降至 2.09 个百分点。同时平均单幅 512×512 影像推理耗时由约 0.15 秒增加至约 0.53 秒，仍处于辅助判读可接受量级。")
    add_callout(doc, "适用边界", "Sen1Floods11 标签表示灾后水体/非水体。该复验不能区分永久水体与新增淹没水体，也未验证作物受损程度、保险责任成立或赔款金额。系统输出必须保留来源、算法版本和人工复核节点。", fill="FFF7E8", accent=GOLD)

    add_heading(doc, "一、验证目标与证据分层", 1)
    add_heading(doc, "1.1 评委意见对应", 2)
    add_body(doc, "评委要求围绕不同灾害类型建立地面查勘真值集、报告关键误差指标，并量化理赔周期、单案成本和人工复核工作量改善。本阶段优先回答可在公开证据下严谨回答的问题：SAR 灾后水体识别是否可运行、误差有多大、较当前生产基线是否改善，以及后续真实业务验证应如何设计。")
    add_heading(doc, "1.2 三层证据结构", 2)
    add_bullet(doc, "A 类：本项目复验结果。由仓库内脚本、清单、CSV、JSON 和图表可重复生成。")
    add_bullet(doc, "B 类：第三方研究结论。用于说明公开领域中遥感洪水制图、农业指数保险的已知能力和限制。")
    add_bullet(doc, "C 类：业务情景测算。仅提供参数化公式与示例，不冒充保险机构真实试点成效。")

    add_heading(doc, "二、系统技术对象与验证范围", 1)
    add_heading(doc, "2.1 穹野智保理赔辅助链路", 2)
    add_body(doc, "系统以在册保单边界和案件为业务主线，串联材料校验、卫星初筛、长势评估、合规校验、规则评级、报告草稿、人工审核和归档。Agent 负责理解自然语言指令、读取案件与证据、调用受控工具并推进辅助流程；最终赔付和归档仍保留人工审核。")
    add_heading(doc, "2.2 本次复验对象", 2)
    add_body(doc, "本次只验证 SAR 模块中的灾后水体像素识别。生产基线采用 Sentinel-1 VH 极化、约 50 米中值滤波和 -16 dB 固定阈值。候选 v2 使用 VV、VH、VV-VH、均值以及 3/7/11 像素窗口局部均值与标准差，共 16 维特征，由 ExtraTrees 完成像素分类。")
    add_heading(doc, "2.3 不在本次结论内的内容", 2)
    add_bullet(doc, "永久水体与新增洪水的时序差分准确率。")
    add_bullet(doc, "旱灾、冰雹、台风、虫害、霜冻等灾种的地面真值准确率。")
    add_bullet(doc, "作物减产率、保险责任判定和赔款金额的真实业务准确率。")
    add_bullet(doc, "保险机构真实理赔周期、单案成本和人工工时节约比例。")

    add_heading(doc, "三、数据与实验设计", 1)
    add_heading(doc, "3.1 数据来源", 2)
    add_body(doc, "Sen1Floods11 是公开的全球洪水遥感数据集，包含 Sentinel-1 SAR 影像及人工水体标签。本次从官方手工标注划分中抽取 150 个样本，全部为 512×512 像素；标签 -1 为无效区域、0 为非水体、1 为水体。")
    add_table(
        doc,
        ["用途", "官方划分", "样本数", "是否参与参数选择"],
        [
            ["训练", "train", 100, "用于拟合 v2"],
            ["校准", "validation", 20, "用于选择固定阈值和概率阈值"],
            ["最终评估", "test", 30, "否；仅用于最终报告"],
        ],
        [1800, 1800, 1400, 4360],
    )
    add_heading(doc, "3.2 比较方案", 2)
    add_table(
        doc,
        ["方案", "规则/模型", "参数来源"],
        [
            ["生产基线", "VH 中值滤波 + -16 dB", "当前项目生产逻辑"],
            ["阈值校准", "VH 中值滤波 + -18 dB", "20 个 validation 样本"],
            ["SAR 特征融合 v2", "ExtraTrees + 16 维极化/纹理特征，阈值 0.70", "100 train 拟合；20 validation 定阈值"],
        ],
        [1900, 4300, 3160],
    )
    add_heading(doc, "3.3 指标定义", 2)
    add_bullet(doc, "Precision：预测为水体的像素中，真实水体所占比例。")
    add_bullet(doc, "Recall：真实水体像素中，被识别出的比例。")
    add_bullet(doc, "F1：Precision 与 Recall 的调和平均。")
    add_bullet(doc, "IoU：预测水体与真实水体交集占并集的比例。")
    add_bullet(doc, "面积比例误差：预测水体占有效像素比例减去真实水体比例，单位为百分点。")
    add_body(doc, "微观指标先汇总全部像素再计算，反映总体像素性能；宏观指标先按样本计算再平均，避免大水体样本完全主导结果。无真实水体且无预测水体时部分比率指标未定义，统计脚本保留该语义。")

    add_heading(doc, "四、复验结果", 1)
    add_heading(doc, "4.1 独立测试集总体结果", 2)
    add_table(
        doc,
        ["方案", "微观 Precision", "微观 Recall", "微观 F1", "微观 IoU", "面积误差中位数", "耗时/幅"],
        [
            ["生产基线", pct(production["micro"]["precision"]), pct(production["micro"]["recall"]), pct(production["micro"]["f1"]), pct(production["micro"]["iou"]), f'{production["median_abs_area_error_pp"]:.2f} pp', f'{production["mean_runtime_ms"]:.0f} ms'],
            ["阈值校准", pct(calibrated["micro"]["precision"]), pct(calibrated["micro"]["recall"]), pct(calibrated["micro"]["f1"]), pct(calibrated["micro"]["iou"]), f'{calibrated["median_abs_area_error_pp"]:.2f} pp', f'{calibrated["mean_runtime_ms"]:.0f} ms'],
            ["SAR 特征融合 v2", pct(fusion["micro"]["precision"]), pct(fusion["micro"]["recall"]), pct(fusion["micro"]["f1"]), pct(fusion["micro"]["iou"]), f'{fusion["median_abs_area_error_pp"]:.2f} pp', f'{fusion["mean_runtime_ms"]:.0f} ms'],
        ],
        [1500, 1300, 1300, 1100, 1100, 1560, 1500],
    )
    add_body(doc, "v2 的主要收益来自显著抑制固定阈值在城市、地形阴影和低后向散射区域产生的假阳性。与基线相比，微观 F1 提升 26.9 个百分点，IoU 提升 30.4 个百分点，面积比例中位绝对误差下降约 93.2%。代价是单幅平均推理时间增加约 0.38 秒。")
    add_figure(doc, "01_test_metrics.png", "图 1  独立测试集宏观指标对比。v2 提升精确率、F1 与 IoU，但召回率低于高召回的固定阈值基线。")
    add_heading(doc, "4.2 阈值选择与独立性", 2)
    add_body(doc, "v2 概率阈值仅在 20 个 validation 样本上从 0.30 至 0.90 搜索，以宏观 F1 最高为选择规则，最终阈值为 0.70。30 个 test 样本未参与训练或阈值选择。")
    add_figure(doc, "02_threshold_sensitivity.png", "图 2  固定 VH 阈值在校准集上的敏感性。该图用于复现生产基线的参数局限。")
    add_heading(doc, "4.3 面积估计与事件差异", 2)
    add_figure(doc, "03_area_agreement.png", "图 3  独立测试样本的真实水体比例与预测水体比例。虚线表示理想一致。")
    add_figure(doc, "04_event_area_error.png", "图 4  不同洪水事件的面积比例误差分布。事件级差异提示跨区域部署仍需本地校准。")
    add_heading(doc, "4.4 像素混淆与典型案例", 2)
    add_figure(doc, "05_confusion_matrix.png", "图 5  生产基线与 v2 的归一化混淆矩阵。")
    add_figure(doc, "06_example_overlays.png", "图 6  v2 在独立测试集中的最好、中位和最差样本。右列显示真阳性、假阳性和假阴性空间分布。", width=6.35)
    add_callout(doc, "结果解读", "v2 证明了公开数据上的改进方向，但最差样本仍存在明显漏检或误检。复赛展示应将其定位为‘提高人工查勘优先级与证据提取效率的辅助能力’，而非无人值守自动定损。")

    add_heading(doc, "五、公开文献交叉验证", 1)
    add_body(doc, "公开文献用于判断本项目结果是否处于合理技术区间，并说明商业部署的通用边界。下列数值均为第三方研究结果，不是穹野智保项目实测。")
    add_table(
        doc,
        ["来源", "公开结论", "对本项目的含义"],
        [
            ["Sen1Floods11, CVPR Workshops 2020", "提供全球洪水事件的地理配准 Sentinel-1 影像和手工水体标签", "支持跨事件公开复验，但标签不是保险赔案真值"],
            ["CEMS Rapid Mapping 验证研究, 2025", "18 个事件中洪水产品用户精度中位数约 84.0%、生产者精度中位数约 72.9%", "遥感洪水图具备辅助价值，但误差与事件差异必须披露"],
            ["CEMS 服务评估, 2025", "产品通常可在数小时内形成，但受云、覆盖和事件条件限制", "系统可强调时效辅助，不应承诺所有案件均自动完成"],
            ["GEE/NDVI 农业保险研究, 2020", "公开案例表明遥感指数可支持农业保险评估并缩短数据等待时间", "支持长势证据链方向，仍需本地作物与保单规则校准"],
            ["世界银行农业保险遥感报告", "NDVI 与 SAR 应用增加，但每个实施项目都需要验证", "公开复验是起点，真实机构试点仍是后续必要环节"],
        ],
        [2000, 3800, 3560],
    )

    add_heading(doc, "六、业务价值的可审计测算方法", 1)
    add_heading(doc, "6.1 不伪造试点数据", 2)
    add_body(doc, "当前没有保险机构真实接入和工时记录，因此不报告‘理赔周期缩短 X%’‘单案成本降低 Y%’等实测结论。复赛阶段采用参数化计算表，由评委或合作机构代入本机构数据即可复算。")
    add_heading(doc, "6.2 建议测算公式", 2)
    add_bullet(doc, "单案节约工时 = 传统查勘工时 -（自动处理时间 + 人工复核时间）。")
    add_bullet(doc, "单案成本变化 = 人员工时成本变化 + 现场出勤成本变化 + 云计算与影像成本变化。")
    add_bullet(doc, "理赔周期变化 = 传统结案时长 - 系统辅助结案时长。")
    add_bullet(doc, "人工复核负担 = 进入人工队列案件数 × 单案复核时长；应按低、中、高风险分层统计。")
    add_heading(doc, "6.3 可验证的技术代理指标", 2)
    add_body(doc, "在无机构试点时，可先报告系统运行层指标：单幅推理耗时、自动生成证据包耗时、失败率、需要人工复核的案件比例、误报/漏报数量，以及报告字段完整率。上述指标可由系统日志直接生成，不应替代真实业务价值指标。")

    add_heading(doc, "七、合规与产品边界", 1)
    add_bullet(doc, "数据授权：仅处理获得合法授权的保单边界、案件材料和遥感数据；公开数据按原许可证和引用要求使用。")
    add_bullet(doc, "决策边界：Agent 和遥感算法输出为理赔辅助建议，不自动形成最终赔付决定。")
    add_bullet(doc, "可追溯性：保留数据来源、时间窗、算法版本、参数、工具调用和人工审核记录。")
    add_bullet(doc, "最小必要：避免向模型发送与案件无关的个人信息；生产接入时应配置脱敏、权限和留存策略。")
    add_bullet(doc, "人工兜底：高风险、证据冲突、数据缺失和超出模型适用范围的案件必须进入人工复核。")

    add_heading(doc, "八、局限性与下一阶段计划", 1)
    add_heading(doc, "8.1 当前局限", 2)
    add_bullet(doc, "样本规模为 30 个独立测试芯片，适合复赛阶段技术复验，不足以替代统计功效充分的商业验证。")
    add_bullet(doc, "标签是灾后水体，未通过灾前水体掩膜分离新增洪水。")
    add_bullet(doc, "公开数据不是中国农业保险承保地块，地域、作物和成像条件存在域偏移。")
    add_bullet(doc, "v2 当前是可复现候选方案，尚未替换线上生产推理链。")
    add_heading(doc, "8.2 复赛阶段落地顺序", 2)
    add_table(
        doc,
        ["优先级", "任务", "验收标准"],
        [
            ["P0", "冻结公开复验清单、脚本、结果和算法版本", "任何环境可用同一随机种子重现 CSV/JSON/图表"],
            ["P0", "将 v2 封装为版本化推理服务并接入系统开关", "同一输入可在基线/v2 间可追溯切换"],
            ["P1", "加入灾前永久水体掩膜和时序变化检测", "报告新增淹没范围，并保留灾前/灾后证据"],
            ["P1", "建立最小地面真值模板", "至少记录灾种、地块、受损面积、查勘日期和证据来源"],
            ["P2", "与保险机构开展受控回溯试点", "盲测误差、工时、周期、成本和人工复核量均可审计"],
        ],
        [1100, 4700, 3560],
    )

    add_heading(doc, "九、可复现性与材料索引", 1)
    add_table(
        doc,
        ["材料", "仓库位置", "用途"],
        [
            ["数据清单", "validation/manifest.csv", "固定样本、划分和来源 URL"],
            ["逐样本结果", "validation/results/sample_metrics.csv", "每个样本、每种方案的像素与面积指标"],
            ["阈值结果", "validation/results/calibration_metrics.csv", "固定阈值校准结果"],
            ["v2 校准结果", "validation/results/ml_calibration_metrics.csv", "概率阈值选择过程"],
            ["汇总指标", "validation/results/summary_metrics.json", "报告数值的机器可读来源"],
            ["验证脚本", "validation/sen1floods11_benchmark.py", "下载、训练、评估和作图"],
            ["自动化测试", "tests/test_public_validation.py", "指标、划分和算法关键逻辑测试"],
        ],
        [1700, 3600, 4060],
    )
    add_body(doc, "复现命令：安装 validation/requirements.txt 后，执行 python -m validation.sen1floods11_benchmark all。原始公开影像默认保存于 validation/data/，不提交 Git；结果文件和图表提交版本控制。")

    add_heading(doc, "参考资料", 1)
    sources = [
        "[1] Bonafilia et al. Sen1Floods11: A Georeferenced Dataset to Train and Test Deep Learning Flood Algorithms for Sentinel-1. CVPR Workshops, 2020. https://openaccess.thecvf.com/content_CVPRW_2020/html/w11/Bonafilia_Sen1Floods11_A_Georeferenced_Dataset_to_Train_and_Test_Deep_Learning_CVPRW_2020_paper.html",
        "[2] Sen1Floods11 repository and dataset documentation. https://github.com/cloudtostreet/Sen1Floods11",
        "[3] Validation of Copernicus Emergency Management Service flood products across 18 events. Science of Remote Sensing, 2025. https://doi.org/10.1016/j.srs.2025.100210",
        "[4] Evaluation of the Copernicus Emergency Management Service rapid flood mapping service. Remote Sensing of Environment, 2025. https://doi.org/10.1016/j.rse.2025.115108",
        "[5] Remote-sensing and Google Earth Engine based agricultural insurance study. Remote Sensing, 2020, 12(18), 3031. https://www.mdpi.com/2072-4292/12/18/3031",
        "[6] Rice crop insurance using remotely sensed vegetation indices. Natural Hazards and Earth System Sciences, 2020. https://nhess.copernicus.org/articles/20/345/2020/",
        "[7] World Bank. Remote sensing applications for agricultural insurance. https://documents1.worldbank.org/curated/en/099021424123514286/pdf/P1720441022bcb0721a2711387ac6ccd404.pdf",
        "[8] WorldFloods dataset. Zenodo. https://doi.org/10.5281/zenodo.8153514",
    ]
    for source in sources:
        add_body(doc, source)

    doc.core_properties.title = "穹野智保公开数据与文献验证报告"
    doc.core_properties.subject = "Agrisky AI SAR 公共数据复验与业务价值证据"
    doc.core_properties.author = "Agrisky AI 项目组"
    doc.core_properties.keywords = "Agrisky AI, 农业保险, SAR, Sen1Floods11, Agent, 遥感理赔"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
