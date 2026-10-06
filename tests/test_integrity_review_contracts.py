"""Focused integrity regressions for report and analysis authority boundaries.

These tests intentionally exercise the security contracts at their persistence
and controlled-download edges.  They do not mock the checks under test.
"""

from __future__ import annotations

import copy
import hashlib
import json
import sqlite3
import sys
import threading
import uuid
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook


sys.path.insert(0, str(Path(__file__).parent.parent / "api_gateway"))

import main  # noqa: E402


def _boundary(offset: float = 0.0) -> dict:
    x = 113.0 + offset
    y = 34.0 + offset
    return {
        "type": "Polygon",
        "coordinates": [
            [[x, y], [x + 0.01, y], [x + 0.01, y + 0.01], [x, y + 0.01], [x, y]]
        ],
    }


def _insert_case(
    claim_id: str,
    policy_id: str | None,
    state: str,
    *,
    crop_type: str = "rice",
) -> None:
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (
            claim_id,
            policy_id,
            state,
            "flood",
            "2026-06-30",
            crop_type,
            "PLOT-INTEGRITY",
            "2026-07-14T00:00:00Z",
        ),
    )
    con.commit()
    con.close()


def _register_policy(policy_id: str, *, address: str = "原承保地址") -> dict:
    return main._save_policy(
        policy_id,
        "完整性测试投保人",
        "rice",
        address,
        _boundary(),
        f"{policy_id}:v1",
    )


def _valid_snapshot() -> dict:
    boundary_sha256 = "a" * 64
    return {
        "case": {
            "policy_id": "POL-SNAPSHOT",
            "policy_version_id": "POL-SNAPSHOT:v1",
            "policy_boundary_sha256": boundary_sha256,
            "policy_boundary_hash_scheme": "canonical_geometry_v1",
        },
        "satellite": {
            "status": "success",
            "damage_ratio": 0.2,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
        "growth": {},
        "historical_ndvi": {},
        "parcel_growth": {},
        "loss_assessment": {},
        "compliance": {
            "status": "success",
            "damage_ratio": 0.2,
            "insured_area_mu": 10.0,
            "valid_damage_area_mu": 2.0,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
        "rule": {"risk_level": "medium", "review_required": True},
        "payout": {"status": "success", "payout_amount_yuan": 1000.0},
        "_snapshot_revisions": {},
    }


def _artifact(kind: str, path: Path) -> dict:
    return {
        "kind": kind,
        "filename": path.name,
        "url": main._output_url(path),
        "download_url": f"/api/v1/cases/fixture/artifacts/generation/{path.name}",
        "media_type": "application/octet-stream",
        "size_bytes": path.stat().st_size,
        "sha256": main._sha256_file(path),
    }


def _ready_report_payload(root: Path) -> dict:
    report_path = root / "claim.docx"
    excel_path = root / "assessment.xlsx"
    bundle_path = root / "reports.zip"
    root.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_paragraph("完整性测试报告")
    document.save(report_path)

    workbook = Workbook()
    workbook.active["A1"] = "完整性测试"
    workbook.save(excel_path)

    with zipfile.ZipFile(bundle_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", "{}")

    artifacts = [
        _artifact("claim_report", report_path),
        _artifact("assessment_excel", excel_path),
        _artifact("bundle", bundle_path),
    ]
    return {
        "status": "success",
        "claim_id": "CLAIM-READY",
        "report_docx_url": main._output_url(report_path),
        "excel_report_url": main._output_url(excel_path),
        "bundle_zip_url": main._output_url(bundle_path),
        "template_version": "v1.1",
        "generation_id": "v1.1-ready-generation",
        "snapshot_sha256": "1" * 64,
        "bundle_sha256": main._sha256_file(bundle_path),
        "growth_status": "not_available",
        "historical_ndvi_status": "not_available",
        "parcel_growth_status": "not_available",
        "artifacts": artifacts,
    }


def _seed_archived_artifact(
    root: Path,
    *,
    with_receipt: bool,
    receipt_mutation: str | None = None,
) -> tuple[str, str, str, bytes]:
    policy_id = f"POL-DOWNLOAD-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-DOWNLOAD-{uuid.uuid4().hex[:8]}"
    generation_id = f"v1-approved-{uuid.uuid4().hex[:8]}"
    review_id = uuid.uuid4().hex
    _register_policy(policy_id)
    _insert_case(claim_id, policy_id, "RULE_DONE")

    claim_key = hashlib.sha256(claim_id.encode("utf-8")).hexdigest()[:20]
    generation_dir = root / "reports" / claim_key / generation_id
    generation_dir.mkdir(parents=True)
    artifact_path = generation_dir / "claim-report.docx"
    content = b"receipt-anchored-approved-artifact"
    artifact_path.write_bytes(content)
    report_artifact = {
        "kind": "claim_report",
        "filename": artifact_path.name,
        "url": main._output_url(artifact_path),
        "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": len(content),
        "sha256": hashlib.sha256(content).hexdigest(),
    }
    main._save_result(
        claim_id,
        "report",
        {
            "claim_id": claim_id,
            "generation_id": generation_id,
            "snapshot_sha256": "1" * 64,
            "bundle_sha256": "2" * 64,
            "artifacts": [report_artifact],
        },
    )

    con = main._db()
    con.execute("UPDATE cases SET state = 'ARCHIVED' WHERE claim_id = ?", (claim_id,))
    con.commit()
    con.close()

    if with_receipt:
        receipt_artifact = {
            key: value
            for key, value in report_artifact.items()
            if key in {"kind", "filename", "media_type", "size_bytes", "sha256"}
        }
        receipt = {
            "schema_version": "1.0",
            "review_id": review_id,
            "claim_id": claim_id,
            "decision": "approved",
            "actor": "test:auditor",
            "reviewed_at_utc": "2026-07-14T00:00:00Z",
            "final_state": "ARCHIVED",
            "report": {
                "generation_id": generation_id,
                "template_version": "v1.1",
                "snapshot_sha256": "1" * 64,
                "bundle_sha256": "2" * 64,
                "manifest_sha256": "3" * 64,
                "artifacts": [receipt_artifact],
            },
        }
        if receipt_mutation == "claim_mismatch":
            receipt["claim_id"] = "CLAIM-OTHER"
        elif receipt_mutation == "generation_mismatch":
            receipt["report"]["generation_id"] = "v1-other-generation"
        elif receipt_mutation == "missing_sha256":
            receipt_artifact.pop("sha256")
        elif receipt_mutation == "missing_size":
            receipt_artifact.pop("size_bytes")
        elif receipt_mutation == "not_approved":
            receipt["decision"] = "rejected"

        receipt_json = json.dumps(receipt, ensure_ascii=False, sort_keys=True, indent=2)
        receipt_sha256 = hashlib.sha256(receipt_json.encode("utf-8")).hexdigest()
        if receipt_mutation == "bad_receipt_digest":
            receipt_sha256 = "f" * 64
        con = main._db()
        con.execute(
            "INSERT INTO review_receipts "
            "(review_id, claim_id, decision, generation_id, receipt_json, receipt_sha256, "
            "idempotency_key, created_at) VALUES (?,?,?,?,?,?,?,?)",
            (
                review_id,
                claim_id,
                "approved",
                generation_id,
                receipt_json,
                receipt_sha256,
                f"download:{review_id}",
                "2026-07-14T00:00:00Z",
            ),
        )
        con.commit()
        con.close()

    return claim_id, generation_id, artifact_path.name, content


def test_case_analysis_pointer_must_equal_registered_immutable_json(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    policy_id = f"POL-ANALYSIS-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-ANALYSIS-{uuid.uuid4().hex[:8]}"
    boundary_sha256 = _register_policy(policy_id)["boundary_sha256"]
    _insert_case(claim_id, policy_id, "SCREENING_DONE")

    task_id = f"parcel-growth-{uuid.uuid4().hex[:12]}"
    task_root = main._analysis_task_root(claim_id, "parcel_growth", task_id)
    task_root.mkdir(parents=True)
    immutable_result = {
        "task_id": task_id,
        "boundary_sha256": boundary_sha256,
        "features": [],
    }
    result_path = task_root / "result.json"
    result_path.write_text(
        json.dumps(immutable_result, ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )
    registered = main._register_analysis_artifacts(
        claim_id,
        "parcel_growth",
        task_id,
        task_root,
        [("parcel_growth_result", "result.json", "application/json")],
    )[0]
    pointer = {
        **immutable_result,
        "_result_artifact_sha256": registered["sha256"],
        "_result_artifact_size_bytes": registered["size_bytes"],
    }
    assert main._analysis_result_response(claim_id, "parcel_growth", pointer)["features"] == []

    tampered_pointer = copy.deepcopy(pointer)
    tampered_pointer["features"] = [{"feature_id": "tampered"}]
    with pytest.raises(main.HTTPException) as error:
        main._analysis_result_response(claim_id, "parcel_growth", tampered_pointer)
    assert error.value.status_code == 409


def test_parcel_analyzer_must_report_the_exact_source_raster(tmp_path: Path) -> None:
    raster = tmp_path / "ndvi_clip.tif"
    raster.write_bytes(b"authoritative-raster-bytes")
    expected = main._sha256_file(raster)
    with pytest.raises(main.HTTPException) as mismatch:
        main._assert_parcel_raster_binding(
            {"raster": {"filename": raster.name, "sha256": "b" * 64}},
            raster,
            expected,
        )
    assert mismatch.value.status_code == 409

    valid_result = {"raster": {"filename": raster.name, "sha256": expected}}
    main._assert_parcel_raster_binding(valid_result, raster, expected)
    raster.write_bytes(b"replaced-after-analysis")
    with pytest.raises(main.HTTPException) as replaced:
        main._assert_parcel_raster_binding(valid_result, raster, expected)
    assert replaced.value.status_code == 409


def test_policy_reference_and_terms_update_are_serialized(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A claim insertion racing a terms update must have one valid serial order."""

    policy_id = f"POL-RACE-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-RACE-{uuid.uuid4().hex[:8]}"
    _register_policy(policy_id, address="old-address")
    original_db = main._db
    allow_reference = threading.Event()
    reference_finished = threading.Event()
    reference_outcome: dict[str, str] = {}
    update_outcome: dict[str, object] = {}

    class InterceptConnection:
        def __init__(self, connection: sqlite3.Connection):
            self._connection = connection
            self._released = False

        def execute(self, sql: str, parameters=()):
            cursor = self._connection.execute(sql, parameters)
            if (
                not self._released
                and "SELECT confirmation_id FROM parcel_confirmations" in sql
            ):
                self._released = True
                allow_reference.set()
                if not reference_finished.wait(timeout=5):
                    raise RuntimeError("timed out waiting for the racing reference")
            return cursor

        def __getattr__(self, name: str):
            return getattr(self._connection, name)

    def connection_factory():
        connection = original_db()
        if threading.current_thread().name == "policy-updater":
            return InterceptConnection(connection)
        return connection

    monkeypatch.setattr(main, "_db", connection_factory)

    def insert_reference() -> None:
        if not allow_reference.wait(timeout=5):
            reference_outcome["status"] = "timeout"
            reference_finished.set()
            return
        connection = sqlite3.connect(str(main.DB_PATH), timeout=0.2)
        try:
            connection.execute(
                "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
                (
                    claim_id,
                    policy_id,
                    "MATERIAL_CHECK",
                    "flood",
                    "2026-06-30",
                    "rice",
                    "PLOT-RACE",
                    "2026-07-14T00:00:00Z",
                ),
            )
            connection.commit()
            reference_outcome["status"] = "inserted"
        except sqlite3.OperationalError as exc:
            connection.rollback()
            reference_outcome["status"] = "locked"
            reference_outcome["error"] = str(exc)
        finally:
            connection.close()
            reference_finished.set()

    def update_terms() -> None:
        try:
            update_outcome["result"] = main._save_policy(
                policy_id,
                "完整性测试投保人",
                "rice",
                "new-address",
                _boundary(),
                f"{policy_id}:v1",
            )
        except Exception as exc:  # noqa: BLE001 - the contract accepts a rejected writer
            update_outcome["error"] = exc

    reference_thread = threading.Thread(target=insert_reference, name="claim-inserter")
    update_thread = threading.Thread(target=update_terms, name="policy-updater")
    reference_thread.start()
    update_thread.start()
    reference_thread.join(timeout=10)
    update_thread.join(timeout=10)
    assert not reference_thread.is_alive()
    assert not update_thread.is_alive()
    assert reference_outcome.get("status") in {"inserted", "locked"}

    if reference_outcome["status"] == "inserted":
        assert "error" in update_outcome, "referenced terms were overwritten after the claim committed"
        expected_address = "old-address"
    else:
        assert "result" in update_outcome, update_outcome.get("error")
        # The update serialized first; create the reference after it commits.
        _insert_case(claim_id, policy_id, "MATERIAL_CHECK")
        expected_address = "new-address"

    con = original_db()
    row = con.execute(
        "SELECT address FROM policies WHERE policy_id = ?", (policy_id,)
    ).fetchone()
    reference = con.execute(
        "SELECT policy_id FROM cases WHERE claim_id = ?", (claim_id,)
    ).fetchone()
    con.close()
    assert row["address"] == expected_address
    assert reference["policy_id"] == policy_id


@pytest.mark.parametrize("operation", ["update", "delete"])
def test_review_receipts_are_database_append_only(operation: str) -> None:
    review_id = uuid.uuid4().hex
    claim_id = f"CLAIM-RECEIPT-{uuid.uuid4().hex[:8]}"
    receipt = {
        "schema_version": "1.0",
        "review_id": review_id,
        "claim_id": claim_id,
        "decision": "approved",
        "report": {"generation_id": "v1-receipt"},
    }
    receipt_json = json.dumps(receipt, ensure_ascii=False, sort_keys=True, indent=2)
    con = main._db()
    con.execute(
        "INSERT INTO review_receipts "
        "(review_id, claim_id, decision, generation_id, receipt_json, receipt_sha256, "
        "idempotency_key, created_at) VALUES (?,?,?,?,?,?,?,?)",
        (
            review_id,
            claim_id,
            "approved",
            "v1-receipt",
            receipt_json,
            hashlib.sha256(receipt_json.encode("utf-8")).hexdigest(),
            f"append-only:{review_id}",
            "2026-07-14T00:00:00Z",
        ),
    )
    con.commit()
    try:
        with pytest.raises(sqlite3.IntegrityError):
            if operation == "update":
                con.execute(
                    "UPDATE review_receipts SET decision = 'rejected' WHERE review_id = ?",
                    (review_id,),
                )
            else:
                con.execute("DELETE FROM review_receipts WHERE review_id = ?", (review_id,))
    finally:
        con.rollback()
        con.close()


def test_archived_artifact_download_requires_an_approved_receipt_anchor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id, generation_id, filename, _content = _seed_archived_artifact(
        tmp_path, with_receipt=False
    )
    with TestClient(main.app) as api:
        response = api.get(
            f"/api/v1/cases/{claim_id}/artifacts/{generation_id}/{filename}"
        )
    assert response.status_code == 409


def test_archived_case_cannot_bypass_receipt_via_preliminary_excel(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-PRELIMINARY-{uuid.uuid4().hex[:8]}"
    _insert_case(claim_id, None, "ARCHIVED")
    reports = tmp_path / "reports"
    reports.mkdir(parents=True)
    path = reports / f"{claim_id}.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "must be receipt anchored"
    workbook.save(path)
    main._save_result(
        claim_id,
        "excel_report",
        {
            "storage_url": main._output_url(path),
            "sha256": main._sha256_file(path),
            "size_bytes": path.stat().st_size,
        },
    )
    with TestClient(main.app) as api:
        response = api.get(f"/api/v1/cases/{claim_id}/preliminary-excel")
    assert response.status_code == 409


def test_archived_artifact_download_accepts_a_verified_receipt_anchor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id, generation_id, filename, content = _seed_archived_artifact(
        tmp_path, with_receipt=True
    )
    with TestClient(main.app) as api:
        response = api.get(
            f"/api/v1/cases/{claim_id}/artifacts/{generation_id}/{filename}"
        )
    assert response.status_code == 200, response.text
    assert response.content == content


@pytest.mark.parametrize(
    "receipt_mutation",
    [
        "claim_mismatch",
        "generation_mismatch",
        "missing_sha256",
        "missing_size",
        "bad_receipt_digest",
        "not_approved",
    ],
)
def test_archived_artifact_download_rejects_invalid_receipt_evidence(
    receipt_mutation: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id, generation_id, filename, _content = _seed_archived_artifact(
        tmp_path,
        with_receipt=True,
        receipt_mutation=receipt_mutation,
    )
    with TestClient(main.app) as api:
        response = api.get(
            f"/api/v1/cases/{claim_id}/artifacts/{generation_id}/{filename}"
        )
    assert response.status_code == 409


@pytest.mark.parametrize(
    "invalid_part",
    [
        "generation_id",
        "snapshot_sha256",
        "bundle_sha256",
        "all_artifacts",
        "claim_report_artifact",
        "assessment_excel_artifact",
        "bundle_artifact",
    ],
)
def test_frozen_report_readiness_requires_complete_identity_and_artifact_registry(
    invalid_part: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    payload = _ready_report_payload(tmp_path / "reports" / "ready")
    assert main._report_payload_ready(payload) is True
    invalid = copy.deepcopy(payload)
    if invalid_part in {"generation_id", "snapshot_sha256", "bundle_sha256"}:
        invalid.pop(invalid_part)
    elif invalid_part == "all_artifacts":
        invalid["artifacts"] = []
    else:
        kind = invalid_part.removesuffix("_artifact")
        invalid["artifacts"] = [
            item for item in invalid["artifacts"] if item["kind"] != kind
        ]
    assert main._report_payload_ready(invalid) is False


@pytest.mark.parametrize("scheme", [None, "legacy_geometry_v0"])
@pytest.mark.parametrize(
    "provenance_step", ["satellite", "compliance", "loss_assessment", "growth"]
)
def test_report_snapshot_rejects_missing_or_unknown_boundary_hash_scheme(
    provenance_step: str,
    scheme: str | None,
) -> None:
    snapshot = _valid_snapshot()
    boundary_sha256 = snapshot["case"]["policy_boundary_sha256"]
    if provenance_step == "loss_assessment":
        snapshot[provenance_step] = {
            "status": "success",
            "yield_loss_ratio": 0.2,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": scheme,
        }
    elif provenance_step == "growth":
        snapshot["growth"] = {
            "task_id": "growth-boundary-scheme",
            "raster": {
                "ndvi_meta": {
                    "source": "gee",
                    "boundary_sha256": boundary_sha256,
                    "boundary_hash_scheme": scheme,
                }
            },
        }
    else:
        snapshot[provenance_step]["boundary_hash_scheme"] = scheme

    with pytest.raises(main.HTTPException) as error:
        main._validate_report_snapshot(snapshot)
    assert error.value.status_code == 409


@pytest.mark.parametrize(
    "invalid_payout",
    [True, float("nan"), float("inf"), float("-inf"), -0.01],
    ids=["bool", "nan", "positive-infinity", "negative-infinity", "negative"],
)
def test_report_snapshot_rejects_non_finite_boolean_or_negative_payout(
    invalid_payout: object,
) -> None:
    snapshot = _valid_snapshot()
    snapshot["payout"]["payout_amount_yuan"] = invalid_payout
    with pytest.raises(main.HTTPException) as error:
        main._validate_report_snapshot(snapshot)
    assert error.value.status_code == 409


@pytest.mark.parametrize("frozen_state", ["REPORT_DRAFTED", "HUMAN_REVIEW", "ARCHIVED"])
def test_payout_cannot_mutate_after_report_freeze(
    frozen_state: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-PAYOUT-FROZEN-{uuid.uuid4().hex[:8]}"
    _insert_case(claim_id, None, "COMPLIANCE_DONE")
    main._save_result(
        claim_id,
        "compliance",
        {"damage_ratio": 0.8, "insured_area_mu": 100.0, "valid_damage_area_mu": 80.0},
    )
    original_payout = {"status": "success", "payout_amount_yuan": 123.45, "marker": "frozen"}
    main._save_result(claim_id, "payout", original_payout)
    con = main._db()
    con.execute("UPDATE cases SET state = ? WHERE claim_id = ?", (frozen_state, claim_id))
    con.commit()
    con.close()

    with TestClient(main.app) as api:
        response = api.post(
            "/api/v1/tools/run_payout_estimate", json={"claim_id": claim_id}
        )
    assert response.status_code == 409
    assert main._get_result(claim_id, "payout") == original_payout


def test_payout_state_check_and_write_are_one_transaction(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-PAYOUT-RACE-{uuid.uuid4().hex[:8]}"
    _insert_case(claim_id, None, "COMPLIANCE_DONE")
    main._save_result(
        claim_id,
        "compliance",
        {"damage_ratio": 0.8, "insured_area_mu": 100.0, "valid_damage_area_mu": 80.0},
    )
    original = {"status": "success", "payout_amount_yuan": 123.45, "marker": "original"}
    main._save_result(claim_id, "payout", original)
    original_db = main._db
    freeze_started = threading.Event()
    freeze_finished = threading.Event()

    class DelayedBegin:
        def __init__(self, connection: sqlite3.Connection):
            self.connection = connection
            self.delayed = False

        def execute(self, sql: str, parameters=()):
            if not self.delayed and sql.strip().upper() == "BEGIN IMMEDIATE":
                self.delayed = True
                freeze_started.set()
                assert freeze_finished.wait(timeout=5)
            return self.connection.execute(sql, parameters)

        def __getattr__(self, name: str):
            return getattr(self.connection, name)

    monkeypatch.setattr(main, "_db", lambda: DelayedBegin(original_db()))

    def freeze() -> None:
        assert freeze_started.wait(timeout=5)
        connection = sqlite3.connect(str(main.DB_PATH))
        connection.execute(
            "UPDATE cases SET state = 'REPORT_DRAFTED' WHERE claim_id = ?", (claim_id,)
        )
        connection.commit()
        connection.close()
        freeze_finished.set()

    worker = threading.Thread(target=freeze)
    worker.start()
    with TestClient(main.app) as api:
        response = api.post("/api/v1/tools/run_payout_estimate", json={"claim_id": claim_id})
    worker.join(timeout=10)
    assert not worker.is_alive()
    assert response.status_code == 409
    assert main._get_result(claim_id, "payout") == original


def test_legacy_client_damage_geometry_route_is_rejected_and_non_authoritative(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    policy_id = f"POL-LEGACY-DAMAGE-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-LEGACY-DAMAGE-{uuid.uuid4().hex[:8]}"
    _register_policy(policy_id)
    _insert_case(claim_id, policy_id, "SCREENING_DONE")

    with TestClient(main.app) as api:
        response = api.post(
            "/api/v1/tools/run_compliance_calc",
            json={
                "claim_id": claim_id,
                "damage_geojson": _boundary(0.001),
                # Required only by the legacy schema.  The route must not become
                # safe merely because this client-supplied duplicate is omitted.
                "insured_geom": _boundary(),
            },
        )
    assert 400 <= response.status_code < 500
    assert main._get_result(claim_id, "compliance") is None
    con = main._db()
    state = con.execute(
        "SELECT state FROM cases WHERE claim_id = ?", (claim_id,)
    ).fetchone()["state"]
    con.close()
    assert state == "SCREENING_DONE"
