"""Regression tests for the current Agrisky AI release."""

import hashlib
import io
import json
import os
import sys
import tempfile
import time
import types
import uuid
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from docx import Document

_review_root = Path(tempfile.gettempdir()) / "agrisky-latest-tests"
_review_root.mkdir(parents=True, exist_ok=True)
_review_id = uuid.uuid4().hex
os.environ["AGRISKY_DB_PATH"] = str(_review_root / f"agrisky-latest-test-{_review_id}.db")
os.environ["AGRISKY_OUTPUT_ROOT"] = str(_review_root / f"agrisky-latest-output-{_review_id}")
os.environ["AGRISKY_ALLOW_MOCK_REMOTE_SENSING"] = "true"
sys.path.insert(0, str(Path(__file__).parent.parent / "api_gateway"))

from fastapi.testclient import TestClient  # noqa: E402

import main  # noqa: E402


client = TestClient(main.app, raise_server_exceptions=False)


def _docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _demo_boundary_sha256() -> str:
    boundary = main._policy_roi("POL-2026-001")
    assert boundary is not None
    return main._geometry_sha256(boundary)


def _artifact_integrity(path: Path) -> dict:
    return {
        "filename": path.name,
        "size_bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def test_growth_by_claim_missing_claim_is_not_sql_server_error():
    response = client.post(
        "/api/v1/tools/run_growth_analysis_by_claim",
        json={"claim_id": "CLAIM-MISSING"},
    )
    assert response.status_code == 404


def test_policy_upsert_preserves_existing_holder_account():
    policy_id = f"POL-UPSERT-{uuid.uuid4().hex[:8]}"
    con = main._db()
    con.execute(
        "INSERT INTO policies "
        "(policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at, holder_account) "
        "VALUES (?,?,?,?,?,?,?,?) "
        "ON CONFLICT(policy_id) DO UPDATE SET holder_account = excluded.holder_account",
        (policy_id, "原投保人", "rice", "旧地址", "{}", 1.0, "2026-07-11", "nonghu"),
    )
    con.commit()
    con.close()

    main._save_policy(
        policy_id,
        "更新投保人",
        "rice",
        "新地址",
        {
            "type": "Polygon",
            "coordinates": [[[113, 34], [113.1, 34], [113.1, 34.1], [113, 34.1], [113, 34]]],
        },
    )

    con = main._db()
    holder = con.execute(
        "SELECT holder_account FROM policies WHERE policy_id = ?", (policy_id,)
    ).fetchone()["holder_account"]
    con.close()
    assert holder == "nonghu"


def test_referenced_policy_terms_and_boundary_are_immutable():
    policy_id = f"POL-FROZEN-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-FROZEN-{uuid.uuid4().hex[:8]}"
    boundary = {
        "type": "Polygon",
        "coordinates": [[[113, 34], [113.01, 34], [113.01, 34.01], [113, 34.01], [113, 34]]],
    }
    main._save_policy(policy_id, "投保人", "rice", "原地址", boundary)
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "MATERIAL_CHECK", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-13"),
    )
    con.commit()
    con.close()

    changed = {
        "type": "Polygon",
        "coordinates": [[[114, 35], [114.01, 35], [114.01, 35.01], [114, 35.01], [114, 35]]],
    }
    with pytest.raises(main.HTTPException) as error:
        main._save_policy(policy_id, "投保人", "rice", "新地址", changed)
    assert error.value.status_code == 409
    assert main._save_policy(policy_id, "投保人", "rice", "原地址", boundary)["boundary_sha256"]


def test_admin_soft_delete_case_then_policy_preserves_source_records():
    suffix = uuid.uuid4().hex[:8]
    policy_id = f"POL-DELETE-{suffix}"
    claim_id = f"CLAIM-DELETE-{suffix}"
    boundary = {
        "type": "Polygon",
        "coordinates": [[[113, 34], [113.01, 34], [113.01, 34.01], [113, 34.01], [113, 34]]],
    }
    main._save_policy(policy_id, "待删除投保人", "rice", "测试地址", boundary)
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "MATERIAL_CHECK", "flood", "2026-07-01", "rice", "PLOT-DEL", "2026-07-30"),
    )
    con.commit()
    con.close()

    blocked = client.delete(f"/api/v1/policies/{policy_id}")
    assert blocked.status_code == 409
    assert "先在案件队列删除关联案件" in blocked.json()["detail"]

    deleted_case = client.delete(f"/api/v1/cases/{claim_id}")
    assert deleted_case.status_code == 200
    assert deleted_case.json()["evidence_retained"] is True
    assert client.get(f"/api/v1/cases/{claim_id}").status_code == 404
    assert claim_id not in {item["claim_id"] for item in client.get("/api/v1/cases").json()["items"]}

    deleted_policy = client.delete(f"/api/v1/policies/{policy_id}")
    assert deleted_policy.status_code == 200
    assert deleted_policy.json()["evidence_retained"] is True
    assert client.get(f"/api/v1/policies/{policy_id}").status_code == 404
    assert policy_id not in {item["policy_id"] for item in client.get("/api/v1/policies").json()["items"]}

    con = main._db()
    assert con.execute("SELECT 1 FROM cases WHERE claim_id = ?", (claim_id,)).fetchone()
    assert con.execute("SELECT 1 FROM policies WHERE policy_id = ?", (policy_id,)).fetchone()
    markers = con.execute(
        "SELECT entity_type FROM deleted_records WHERE entity_id IN (?, ?) ORDER BY entity_type",
        (claim_id, policy_id),
    ).fetchall()
    con.close()
    assert [row["entity_type"] for row in markers] == ["case", "policy"]


def test_material_documents_are_extracted_grounded_and_gate_validation(monkeypatch):
    monkeypatch.setattr(main, "REQUIRE_CLAIM_DOCUMENTS", True)
    monkeypatch.setattr(main, "llm_fields", lambda _lines, _document_type: ([], None))
    suffix = uuid.uuid4().hex[:8].upper()
    policy_id = f"POL-MATERIAL-{suffix}"
    claim_id = f"CLAIM-MATERIAL-{suffix}"
    boundary = {
        "type": "Polygon",
        "coordinates": [[[113, 34], [113.01, 34], [113.01, 34.01], [113, 34.01], [113, 34]]],
    }
    main._save_policy(policy_id, "材料测试合作社", "rice", "测试地址", boundary)
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "MATERIAL_CHECK", "flood", "2026-07-28", "rice", "PLOT-1", "2026-07-30"),
    )
    con.commit()
    con.close()

    payloads = {
        "policy_document": _docx_bytes(f"保单号：{policy_id}", "作物：水稻", "地块编号：PLOT-1"),
        "claim_notice": _docx_bytes(
            f"保单号：{policy_id}", "出险日期：2026年07月28日", "灾害类型：洪涝"
        ),
        "damage_certificate": _docx_bytes("查勘记录", "灾害类型：洪涝", "作物：水稻"),
    }
    for document_type, content in payloads.items():
        response = client.post(
            f"/api/v1/cases/{claim_id}/documents",
            data={"document_type": document_type},
            files={
                "document": (
                    f"{document_type}.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["extraction"]["field_count"] >= 1

    review = client.get(f"/api/v1/cases/{claim_id}/material-review")
    assert review.status_code == 200
    assert review.json()["passed"] is True
    assert len(review.json()["documents"]) == 3
    assert any(field["source_ref"].startswith("p1-l") for field in review.json()["fields"])

    validated = client.post("/api/v1/tools/validate_materials", json={"claim_id": claim_id})
    assert validated.status_code == 200
    assert validated.json()["passed"] is True
    assert validated.json()["quality_flags"]["document_understanding"] == "pass"


def test_optional_documents_do_not_block_claim_workflow(monkeypatch):
    monkeypatch.setattr(main, "REQUIRE_CLAIM_DOCUMENTS", False)
    suffix = uuid.uuid4().hex[:8].upper()
    policy_id = f"POL-OPTIONAL-{suffix}"
    claim_id = f"CLAIM-OPTIONAL-{suffix}"
    boundary = {
        "type": "Polygon",
        "coordinates": [[[113, 34], [113.01, 34], [113.01, 34.01], [113, 34.01], [113, 34]]],
    }
    main._save_policy(policy_id, "演示合作社", "rice", "测试地址", boundary)
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "MATERIAL_CHECK", "flood", "2026-07-28", "rice", "PLOT-1", "2026-07-30"),
    )
    con.commit()
    con.close()

    review = client.get(f"/api/v1/cases/{claim_id}/material-review")
    assert review.status_code == 200
    assert review.json()["required"] is False
    assert review.json()["required_document_types"] == []
    assert review.json()["missing_document_types"] == []
    assert review.json()["passed"] is True

    validated = client.post("/api/v1/tools/validate_materials", json={"claim_id": claim_id})
    assert validated.status_code == 200
    assert validated.json()["passed"] is True
    assert client.get(f"/api/v1/cases/{claim_id}").json()["state"] == "PREPROCESS_READY"


def test_material_review_only_uses_latest_completed_extraction_fields():
    suffix = uuid.uuid4().hex[:8].upper()
    claim_id = f"CLAIM-FIELDS-{suffix}"
    document_id = f"DOC-FIELDS-{suffix}"
    older_id = f"DEX-OLD-{suffix}"
    latest_id = f"DEX-NEW-{suffix}"
    con = main._db()
    con.executemany(
        "INSERT INTO document_extraction_runs "
        "(extraction_id, document_id, status, started_at, completed_at) VALUES (?,?,?,?,?)",
        [
            (older_id, document_id, "completed", "2026-07-30T10:00:00", "2026-07-30T10:00:01"),
            (latest_id, document_id, "completed", "2026-07-30T11:00:00", "2026-07-30T11:00:01"),
        ],
    )
    con.executemany(
        "INSERT INTO document_fields "
        "(field_id, extraction_id, document_id, claim_id, field_name, normalized_value_json, "
        "raw_text, confidence, page_number, bbox_json, source_ref, extractor, created_at) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        [
            (
                f"FIELD-OLD-{suffix}", older_id, document_id, claim_id, "loss_date",
                json.dumps("2025-08-31"), "2025-08-31", 0.9, 1, None, "p1-l3",
                "llm_json", "2026-07-30T10:00:01",
            ),
            (
                f"FIELD-NEW-{suffix}", latest_id, document_id, claim_id, "survey_date",
                json.dumps("2025-08-31"), "2025-08-31", 0.85, 1, None, "p1-l3",
                "llm_json", "2026-07-30T11:00:01",
            ),
        ],
    )
    con.commit()
    fields = main._latest_material_fields(con, claim_id)
    con.close()
    assert [field["field_name"] for field in fields] == ["survey_date"]


def test_agent_run_endpoint_persists_progress_and_result(monkeypatch):
    from agent import agent_runtime

    def fake_agent(messages, **kwargs):
        callback = kwargs.get("event_callback")
        if callback:
            callback("tool_started", {"tool": "analyze_claims", "args": {}})
            callback("tool_completed", {"tool": "analyze_claims", "status": "success"})
        return {
            "messages": [*messages, {"role": "assistant", "content": "队列分析完成"}],
            "trace": [{"tool": "analyze_claims", "result": {"status": "success"}}],
            "reply": "队列分析完成",
            "error": False,
        }

    monkeypatch.setattr(agent_runtime, "run_agent_api", fake_agent)
    created = client.post(
        "/api/v1/agent/runs",
        json={"messages": [{"role": "user", "content": "分析案件队列"}]},
    )
    assert created.status_code == 202
    run_id = created.json()["run_id"]
    payload = None
    for _ in range(50):
        response = client.get(f"/api/v1/agent/runs/{run_id}")
        assert response.status_code == 200
        payload = response.json()
        if payload["status"] in {"completed", "failed", "cancelled"}:
            break
        time.sleep(0.02)
    assert payload is not None
    assert payload["status"] == "completed"
    assert payload["response"]["reply"] == "队列分析完成"
    assert [event["event_type"] for event in payload["events"]] == [
        "run_started",
        "tool_started",
        "tool_completed",
        "run_completed",
    ]


def test_report_failure_does_not_advance_case(monkeypatch):
    claim_id = "CLAIM-REPORT-FAIL-LATEST"
    boundary_sha256 = _demo_boundary_sha256()
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, "POL-2026-001", "RULE_DONE", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()
    main._save_result(
        claim_id,
        "satellite",
        {
            "status": "success",
            "damage_ratio": 0.2,
            "suspected_damage_area_mu": 20.0,
            "confidence": "high",
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "compliance",
        {
            "status": "success",
            "damage_ratio": 0.2,
            "insured_area_mu": 100.0,
            "valid_damage_area_mu": 20.0,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "rule",
        {"risk_level": "medium", "review_required": True, "rule_version": "test_v1"},
    )
    main._save_result(
        claim_id,
        "payout",
        {"status": "success", "payout_amount_yuan": 1000.0},
    )

    import report_generator

    def fail_renderer(*args, **kwargs):
        raise RuntimeError("forced renderer failure")

    monkeypatch.setattr(report_generator, "generate_claim_report", fail_renderer)
    response = client.post("/api/v1/tools/generate_report", json={"claim_id": claim_id})

    con = main._db()
    state = con.execute("SELECT state FROM cases WHERE claim_id = ?", (claim_id,)).fetchone()["state"]
    con.close()
    assert response.status_code == 500
    assert state == "RULE_DONE"


def test_outputs_are_protected_when_api_key_auth_is_enabled(monkeypatch):
    monkeypatch.setattr(main, "AUTH_ENABLED", True)
    monkeypatch.setattr(main, "API_KEYS", {"viewer-key": "viewer"})
    response = client.get("/outputs/not-present")
    assert response.status_code == 401
    monkeypatch.setattr(main, "AUTH_ENABLED", False)


def test_invalid_portal_token_cannot_run_agent(monkeypatch):
    fake_runtime = types.SimpleNamespace(run_agent_api=lambda *args, **kwargs: {"trace": []})
    monkeypatch.setitem(sys.modules, "agent.agent_runtime", fake_runtime)
    response = client.post(
        "/api/v1/agent/chat",
        headers={"x-agrisky-token": "invalid-token"},
        json={"messages": []},
    )
    assert response.status_code == 401


def test_satellite_screening_uses_policy_boundary_not_client_roi(monkeypatch):
    claim_id = "CLAIM-ROI-SCOPE-LATEST"
    policy_id = "POL-ROI-SCOPE-LATEST"
    server_roi = {
        "type": "Polygon",
        "coordinates": [[[120, 30], [120.1, 30], [120.1, 30.1], [120, 30.1], [120, 30]]],
    }
    client_roi = {
        "type": "Polygon",
        "coordinates": [[[1, 1], [2, 1], [2, 2], [1, 2], [1, 1]]],
    }
    con = main._db()
    con.execute(
        "INSERT OR REPLACE INTO policies "
        "(policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (policy_id, "holder", "rice", "address", __import__("json").dumps(server_roi), 10.0, "2026-07-11"),
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "PREPROCESS_READY", "flood", "2026-07-01", "rice", "PLOT-ROI", "2026-07-11"),
    )
    con.commit()
    con.close()

    captured = {}
    import space_engine.sar_flood as sar_flood

    def fake_calculate_flood_ratio(**kwargs):
        captured["roi"] = kwargs["roi_geojson"]
        return {"status": "success", "flooded_area_mu": 1, "damage_ratio": 0.1, "confidence": "high"}

    monkeypatch.setattr(sar_flood, "calculate_flood_ratio", fake_calculate_flood_ratio)
    monkeypatch.setattr(main, "_ensure_screening_thumbnails", lambda *args: (None, None))
    response = client.post(
        "/api/v1/tools/run_satellite_screening",
        json={"claim_id": claim_id, "roi_geojson": client_roi, "start_date": "2026-06-01", "end_date": "2026-07-01"},
    )
    assert response.status_code == 200
    assert captured["roi"] == server_roi


def test_material_validation_requires_registered_boundary():
    claim_id = f"CLAIM-MATERIAL-{uuid.uuid4().hex[:8]}"
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, "POL-NOT-REGISTERED", "MATERIAL_CHECK", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()

    response = client.post("/api/v1/tools/validate_materials", json={"claim_id": claim_id})
    assert response.status_code == 200
    assert response.json()["passed"] is False
    assert "insured_geom" in response.json()["missing_fields"]

    con = main._db()
    state = con.execute("SELECT state FROM cases WHERE claim_id = ?", (claim_id,)).fetchone()["state"]
    con.close()
    assert state == "MATERIAL_CHECK"


def test_compliance_estimate_uses_server_boundary_and_ratio():
    policy_id = f"POL-COMPLIANCE-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-COMPLIANCE-{uuid.uuid4().hex[:8]}"
    server_roi = {
        "type": "Polygon",
        "coordinates": [[[113.0, 34.0], [113.01, 34.0], [113.01, 34.01], [113.0, 34.01], [113.0, 34.0]]],
    }
    client_roi = {
        "type": "Polygon",
        "coordinates": [[[1.0, 1.0], [20.0, 1.0], [20.0, 20.0], [1.0, 20.0], [1.0, 1.0]]],
    }
    con = main._db()
    con.execute(
        "INSERT INTO policies (policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (policy_id, "holder", "rice", "address", __import__("json").dumps(server_roi), None, "2026-07-11"),
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "SCREENING_DONE", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()
    main._save_result(claim_id, "satellite", {"damage_ratio": 0.0})

    response = client.post(
        "/api/v1/tools/run_compliance_estimate",
        json={"claim_id": claim_id, "roi_geojson": client_roi},
    )
    assert response.status_code == 200
    result = response.json()
    assert result["damage_ratio"] == 0.0
    assert result["insured_area_mu"] == main._geojson_area_mu(server_roi)


def test_compliance_estimate_rejects_missing_authoritative_ratio():
    policy_id = f"POL-NO-RATIO-{uuid.uuid4().hex[:8]}"
    claim_id = f"CLAIM-NO-RATIO-{uuid.uuid4().hex[:8]}"
    roi = {
        "type": "Polygon",
        "coordinates": [[[113.0, 34.0], [113.01, 34.0], [113.01, 34.01], [113.0, 34.01], [113.0, 34.0]]],
    }
    con = main._db()
    con.execute(
        "INSERT INTO policies (policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (policy_id, "holder", "rice", "address", __import__("json").dumps(roi), None, "2026-07-11"),
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "SCREENING_DONE", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()

    response = client.post("/api/v1/tools/run_compliance_estimate", json={"claim_id": claim_id})
    assert response.status_code == 409

    con = main._db()
    state = con.execute("SELECT state FROM cases WHERE claim_id = ?", (claim_id,)).fetchone()["state"]
    con.close()
    assert state == "SCREENING_DONE"


def test_rule_engine_uses_persisted_values_and_is_idempotent():
    claim_id = f"CLAIM-RULE-{uuid.uuid4().hex[:8]}"
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, "POL-2026-001", "COMPLIANCE_DONE", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()
    main._save_result(
        claim_id,
        "compliance",
        {"status": "success", "damage_ratio": 0.55, "insured_area_mu": 10.0, "valid_damage_area_mu": 5.5},
    )

    first = client.post(
        "/api/v1/tools/run_rule_engine",
        json={"claim_id": claim_id, "damage_ratio": 0.0, "crop_type": "wheat", "estimated_payout_yuan": 0},
    )
    assert first.status_code == 200
    assert first.json()["risk_level"] == "high"
    assert "damage_ratio=0.55" in first.json()["rule_trace"]
    assert "crop_type=rice" in first.json()["rule_trace"]

    second = client.post(
        "/api/v1/tools/run_rule_engine",
        json={"claim_id": claim_id, "damage_ratio": 0.01, "crop_type": "corn"},
    )
    assert second.status_code == 200
    assert second.json() == first.json()


def test_rule_engine_completed_state_without_result_is_conflict():
    claim_id = f"CLAIM-RULE-MISSING-{uuid.uuid4().hex[:8]}"
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, "POL-2026-001", "RULE_DONE", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()

    response = client.post("/api/v1/tools/run_rule_engine", json={"claim_id": claim_id})
    assert response.status_code == 409


def test_human_review_refuses_to_archive_without_registered_report():
    claim_id = f"CLAIM-REVIEW-{uuid.uuid4().hex[:8]}"
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, "POL-2026-001", "REPORT_DRAFTED", "flood", "2026-07-01", "rice", "PLOT-1", "2026-07-11"),
    )
    con.commit()
    con.close()

    response = client.post(
        f"/api/v1/cases/{claim_id}/human_review",
        json={"decision": "approved", "generation_id": "v1-missing", "comment": "approve"},
    )
    assert response.status_code == 409
    state = client.get(f"/api/v1/cases/{claim_id}").json()["state"]
    assert state == "REPORT_DRAFTED"
    audit = client.get(f"/api/v1/audit/{claim_id}").json()
    assert not any(item["action"].startswith("human_review_") for item in audit)


def test_gee_service_account_path_is_used(monkeypatch, tmp_path):
    import space_engine.gee_auth as gee_auth

    credential_file = tmp_path / "service-account.json"
    credential_file.write_text("{}", encoding="utf-8")
    sentinel = object()
    captured = {}
    monkeypatch.setenv("GEE_CREDENTIALS", str(credential_file))
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    monkeypatch.setattr(gee_auth.ee, "ServiceAccountCredentials", lambda _email, path: sentinel)
    monkeypatch.setattr(
        gee_auth.ee,
        "Initialize",
        lambda **kwargs: captured.update(kwargs),
    )
    monkeypatch.setattr(gee_auth, "_gee_initialized", False)
    monkeypatch.setattr(gee_auth, "_gee_failed_at", 0.0)

    assert gee_auth.init_gee(project_id="project-test", max_retries=1) is True
    assert captured["credentials"] is sentinel
    assert captured["project"] == "project-test"
    assert captured["http_transport"] is not None


def test_feature_collection_normalization_keeps_all_parcels():
    from shapely.geometry import shape

    feature_collection = {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature",
                "properties": {"field_id": "A"},
                "geometry": {
                    "type": "Polygon",
                    "coordinates": [[[113.0, 34.0], [113.01, 34.0], [113.01, 34.01], [113.0, 34.01], [113.0, 34.0]]],
                },
            },
            {
                "type": "Feature",
                "properties": {"field_id": "B"},
                "geometry": {
                    "type": "Polygon",
                    "coordinates": [[[113.02, 34.0], [113.03, 34.0], [113.03, 34.01], [113.02, 34.01], [113.02, 34.0]]],
                },
            },
        ],
    }
    normalized = main._normalize_geometry(feature_collection)
    assert normalized["type"] == "MultiPolygon"
    assert len(shape(normalized).geoms) == 2


def test_fixed_ndvi_thresholds_and_inclusive_end_date():
    import numpy as np
    from growth_analysis import _inclusive_gee_end_date, compute_class_breaks, validate_ndvi_stats

    values = np.array([0.1, 0.35, 0.5, 0.7, 0.9], dtype="float32")
    assert compute_class_breaks(values, "fixed", 5, {"min": 0.1, "max": 0.9}) == [0.3, 0.45, 0.6, 0.75, 1.0]
    assert _inclusive_gee_end_date("2026-07-12") == "2026-07-13"
    validate_ndvi_stats({"min": -1.0, "max": 1.0})
    try:
        validate_ndvi_stats({"min": -0.2, "max": 8500})
    except ValueError as exc:
        assert "比例因子" in str(exc)
    else:
        raise AssertionError("未缩放的整数 NDVI 必须被拒绝")


def test_growth_window_defaults_to_loss_quarter():
    assert main._growth_observation_window("2025-08-30", None, None) == ("2025-07-01", "2025-09-30")


def test_production_growth_does_not_silently_fallback_to_synthetic(monkeypatch, tmp_path):
    import growth_analysis

    monkeypatch.setenv("AGRISKY_ALLOW_MOCK_REMOTE_SENSING", "false")
    monkeypatch.setattr(
        growth_analysis,
        "export_gee_ndvi_raster",
        lambda **_kwargs: (_ for _ in ()).throw(RuntimeError("forced GEE failure")),
    )
    try:
        growth_analysis.prepare_auto_ndvi_raster(
            boundary_path=tmp_path / "boundary.geojson",
            work_dir=tmp_path,
            out_tif=tmp_path / "ndvi.tif",
            ndvi_source="auto",
        )
    except RuntimeError as exc:
        assert "禁止降级为模拟数据" in str(exc)
    else:
        raise AssertionError("production mode must not create synthetic NDVI")


def test_generate_report_bundles_persisted_growth_not_client_payload():
    from docx import Document
    from openpyxl import load_workbook

    claim_id = f"CLAIM-GROWTH-BUNDLE-{uuid.uuid4().hex[:8]}"
    task_id = f"growth-test-{uuid.uuid4().hex[:8]}"
    boundary_sha256 = _demo_boundary_sha256()
    growth_dir = main.OUTPUT_ROOT / "growth" / task_id
    growth_dir.mkdir(parents=True, exist_ok=True)
    ndvi_tif = growth_dir / "ndvi_clip.tif"
    ndvi_tif.write_bytes(b"hash-bound-test-ndvi-raster")
    class_preview = growth_dir / "class_preview.png"
    Image.new("RGB", (16, 16), "green").save(class_preview)
    growth_docx = growth_dir / "growth_report.docx"
    document = Document()
    document.add_paragraph("服务端持久化长势报告")
    document.save(growth_docx)

    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, "POL-2026-001", "RULE_DONE", "flood", "2025-08-30", "rice", "PLOT-1", "2026-07-13"),
    )
    con.commit()
    con.close()
    main._save_result(
        claim_id,
        "satellite",
        {
            "damage_ratio": 0.2,
            "confidence": "high",
            "image_count": 3,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "compliance",
        {
            "status": "success",
            "damage_ratio": 0.2,
            "insured_area_mu": 100.0,
            "valid_damage_area_mu": 20.0,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "rule",
        {"risk_level": "medium", "review_required": False, "rule_trace": [], "rule_version": "default_v1"},
    )
    main._save_result(
        claim_id,
        "growth",
        {
            "status": "success",
            "schema_version": "agrisky.growth-analysis/v1",
            "algorithm_version": "ndvi-five-level-v1",
            "task_id": task_id,
            "method": "fixed",
            "n_classes": 5,
            "total_area_mu": 100.0,
            "valid_pixel_count": 10,
            "class_breaks": [0.3, 0.45, 0.6, 0.75, 1.0],
            "raster": {
                "ndvi_source": "gee",
                "ndvi_source_label": "GEE Sentinel-2 SR NDVI",
                "ndvi_clip_sha256": _artifact_integrity(ndvi_tif)["sha256"],
                "ndvi_clip_size_bytes": ndvi_tif.stat().st_size,
                "ndvi_meta": {
                    "source": "gee",
                    "boundary_sha256": boundary_sha256,
                    "boundary_hash_scheme": "canonical_geometry_v1",
                },
            },
            "summary": [
                {"value": 5, "label": "优", "count": 10, "ratio": 1.0, "area_mu": 100.0, "color": "#2f9e44"}
            ],
            "outputs": {
                "ndvi_clip_tif": f"/outputs/growth/{task_id}/ndvi_clip.tif",
                "class_preview_png": f"/outputs/growth/{task_id}/class_preview.png",
                "report_docx": f"/outputs/growth/{task_id}/growth_report.docx",
            },
            "artifact_integrity": {
                "ndvi_clip_tif": _artifact_integrity(ndvi_tif),
                "class_preview_png": _artifact_integrity(class_preview),
                "report_docx": _artifact_integrity(growth_docx),
            },
            "message": "ok",
        },
    )

    response = client.post(
        "/api/v1/tools/generate_report",
        json={
            "claim_id": claim_id,
            "data": {"growth": {"summary": [{"value": 1, "label": "伪造差", "area_mu": 9999, "ratio": 1.0}]}},
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["growth_report_docx_url"].endswith(f"{claim_id}-growth.docx")
    assert payload["excel_report_url"].endswith(f"{claim_id}.xlsx")
    assert payload["bundle_zip_url"].endswith(f"{claim_id}-reports.zip")

    claim_report = main.OUTPUT_ROOT / "reports" / f"{claim_id}.docx"
    report_text = "\n".join(p.text for p in Document(claim_report).paragraphs)
    assert "伪造差" not in report_text
    assert "固定阈值" in report_text

    workbook = load_workbook(main.OUTPUT_ROOT / "reports" / f"{claim_id}.xlsx", read_only=True)
    assert "长势分析" in workbook.sheetnames
    workbook.close()

    with zipfile.ZipFile(main.OUTPUT_ROOT / "reports" / f"{claim_id}-reports.zip") as archive:
        names = set(archive.namelist())
    assert f"{claim_id}.docx" in names
    assert f"{claim_id}-growth.docx" in names
    assert f"{claim_id}.xlsx" in names
    assert "manifest.json" in names
