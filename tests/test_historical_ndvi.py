"""Focused tests for historical quarterly NDVI planning, schema, and reports."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from PIL import Image
from pydantic import ValidationError

from historical_ndvi import (
    MAX_SCENE_PROVENANCE_ITEMS,
    HistoricalNDVIResult,
    HistoricalNDVISource,
    QuarterlyNDVIObservation,
    QuarterlySceneProvenance,
    _build_scene_provenance,
    load_historical_ndvi_result,
    plan_quarter_windows,
    save_historical_ndvi_result,
)
from sentinel2_ndvi import (
    S2_CLOUD_MASK_VERSION,
    add_ndvi_band,
    build_s2_sr_collection,
    mask_s2_sr,
)


def test_quarter_plan_keeps_partial_and_future_quarters_explicit() -> None:
    windows = plan_quarter_windows(2025, 2026, as_of_date="2026-07-12")
    reference_span = plan_quarter_windows(2022, 2026, as_of_date="2026-07-12")

    assert len(windows) == 8
    assert len(reference_span) == 20
    assert [item.status for item in reference_span].count("complete") == 18
    assert [item.status for item in reference_span].count("partial") == 1
    assert [item.status for item in reference_span].count("not_reached") == 1
    assert windows[5].label == "2026-Q2"
    assert windows[5].status == "complete"
    assert windows[5].observation_end_date == date(2026, 6, 30)
    assert windows[6].status == "partial"
    assert windows[6].start_date == date(2026, 7, 1)
    assert windows[6].calendar_end_date == date(2026, 9, 30)
    assert windows[6].observation_end_date == date(2026, 7, 12)
    assert windows[7].status == "not_reached"
    assert windows[7].observation_end_date is None


@pytest.mark.parametrize(
    ("start_year", "end_year", "message"),
    [
        (2016, 2020, "不得早于 2017"),
        (2025, 2024, "不能早于"),
        (2017, 2037, "最多支持 20 年"),
    ],
)
def test_quarter_plan_rejects_invalid_or_abusive_ranges(
    start_year: int,
    end_year: int,
    message: str,
) -> None:
    with pytest.raises(ValueError, match=message):
        plan_quarter_windows(start_year, end_year, as_of_date="2026-07-12")


class _Condition:
    def __init__(self, equals: set[int] | None = None, dilations: list[tuple[int, str]] | None = None):
        self.equals = equals or set()
        self.dilations = dilations or []
        self.negated = False

    def Or(self, other: "_Condition") -> "_Condition":
        return _Condition(self.equals | other.equals, self.dilations + other.dilations)

    def focalMax(self, *, radius: int, units: str) -> "_Condition":
        return _Condition(set(self.equals), [*self.dilations, (radius, units)])

    def Not(self) -> "_Condition":
        result = _Condition(set(self.equals), list(self.dilations))
        result.negated = True
        return result


class _Band:
    def eq(self, value: int) -> _Condition:
        return _Condition({value})


class _FakeImage:
    def __init__(self) -> None:
        self.updated_mask: _Condition | None = None
        self.normalized_bands: list[str] | None = None
        self.renamed: str | None = None
        self.added_band: object | None = None

    def select(self, name: str) -> _Band:
        assert name == "SCL"
        return _Band()

    def updateMask(self, mask: _Condition) -> "_FakeImage":
        self.updated_mask = mask
        return self

    def normalizedDifference(self, bands: list[str]) -> "_FakeImage":
        self.normalized_bands = bands
        return self

    def rename(self, name: str) -> "_FakeImage":
        self.renamed = name
        return self

    def addBands(self, band: object) -> "_FakeImage":
        self.added_band = band
        return self


def test_shared_sentinel_mask_excludes_all_classes_and_dilates_edges() -> None:
    image = _FakeImage()
    result = mask_s2_sr(image)

    assert result is image
    assert image.updated_mask is not None
    assert image.updated_mask.equals == {0, 1, 2, 3, 8, 9, 10, 11}
    assert image.updated_mask.dilations == [(20, "meters")]
    assert image.updated_mask.negated is True
    assert S2_CLOUD_MASK_VERSION == "s2-scl-v2-20m"

    add_ndvi_band(image)
    assert image.normalized_bands == ["B8", "B4"]
    assert image.renamed == "ndvi"
    assert image.added_band is image


class _FakeCollection:
    def __init__(self) -> None:
        self.calls: list[tuple] = []

    def filterDate(self, start: str, end: str) -> "_FakeCollection":
        self.calls.append(("filterDate", start, end))
        return self

    def filterBounds(self, roi: object) -> "_FakeCollection":
        self.calls.append(("filterBounds", roi))
        return self

    def filter(self, expression: object) -> "_FakeCollection":
        self.calls.append(("filter", expression))
        return self

    def map(self, function: object) -> "_FakeCollection":
        self.calls.append(("map", getattr(function, "__name__", "?")))
        return self


class _FakeFilter:
    @staticmethod
    def lt(field: str, value: float) -> tuple[str, str, float]:
        return "lt", field, value


class _FakeEE:
    Filter = _FakeFilter

    def __init__(self) -> None:
        self.collection = _FakeCollection()

    def ImageCollection(self, collection_id: str) -> _FakeCollection:
        assert collection_id == "COPERNICUS/S2_SR_HARMONIZED"
        return self.collection


def test_collection_builder_uses_inclusive_conversion_input_and_shared_functions() -> None:
    ee = _FakeEE()
    roi = object()
    collection = build_s2_sr_collection(
        ee,
        roi=roi,
        start_date="2026-07-01",
        end_exclusive="2026-07-13",
        max_cloud_pct=30,
    )

    assert collection.calls == [
        ("filterDate", "2026-07-01", "2026-07-13"),
        ("filterBounds", roi),
        ("filter", ("lt", "CLOUDY_PIXEL_PERCENTAGE", 30)),
        ("map", "mask_s2_sr"),
        ("map", "add_ndvi_band"),
    ]


class _ImmediateInfo:
    def __init__(self, value: object) -> None:
        self.value = value

    def getInfo(self) -> object:
        return self.value


class _FakeSceneCollection:
    def __init__(self, scene_ids: list[str], product_ids: list[str]) -> None:
        self.scene_ids = scene_ids
        self.product_ids = product_ids
        self.requested: list[str] = []

    def sort(self, property_name: str) -> "_FakeSceneCollection":
        assert property_name == "system:index"
        pairs = sorted(zip(self.scene_ids, self.product_ids, strict=True))
        self.scene_ids = [item[0] for item in pairs]
        self.product_ids = [item[1] for item in pairs]
        return self

    def aggregate_array(self, property_name: str) -> _ImmediateInfo:
        self.requested.append(property_name)
        if property_name == "system:index":
            return _ImmediateInfo(self.scene_ids)
        if property_name == "PRODUCT_ID":
            return _ImmediateInfo(self.product_ids)
        raise AssertionError(property_name)


def test_scene_provenance_is_ordered_complete_bounded_and_hash_bound() -> None:
    collection = _FakeSceneCollection(
        ["SCENE-C", "SCENE-A", "SCENE-B"],
        ["PRODUCT-C", "PRODUCT-A", "PRODUCT-B"],
    )
    provenance = _build_scene_provenance(collection, 3)

    assert collection.requested == ["system:index", "PRODUCT_ID"]
    assert [item.scene_id for item in provenance.scenes] == [
        "SCENE-A",
        "SCENE-B",
        "SCENE-C",
    ]
    assert [item.product_id for item in provenance.scenes] == [
        "PRODUCT-A",
        "PRODUCT-B",
        "PRODUCT-C",
    ]
    assert provenance.scene_count == 3
    assert provenance.metadata_complete is True
    assert provenance.schema_version == "agrisky.historical-ndvi-scene-set/v1"
    assert provenance.source_id == "google-earth-engine"
    assert provenance.algorithm_id == "agrisky.sentinel2.quarterly-ndvi"
    assert len(provenance.scene_set_sha256) == 64

    same_set = _build_scene_provenance(
        _FakeSceneCollection(
            ["SCENE-B", "SCENE-C", "SCENE-A"],
            ["PRODUCT-B", "PRODUCT-C", "PRODUCT-A"],
        ),
        3,
    )
    assert same_set.scene_set_sha256 == provenance.scene_set_sha256

    tampered = provenance.model_dump(mode="json")
    tampered["scenes"][0]["product_id"] = "PRODUCT-TAMPERED"
    with pytest.raises(ValidationError, match="scene_set_sha256"):
        QuarterlySceneProvenance.model_validate(tampered)

    with pytest.raises(RuntimeError, match="GEE_SCENE_SET_TOO_LARGE"):
        _build_scene_provenance(
            _FakeSceneCollection([], []),
            MAX_SCENE_PROVENANCE_ITEMS + 1,
        )


def test_public_failure_redacts_exception_details_and_has_distinct_cache_semantics() -> None:
    window = plan_quarter_windows(2025, 2025, as_of_date="2026-07-12")[0]
    secret = "https://earthengine.example/thumb?api_key=do-not-leak"
    failed = QuarterlyNDVIObservation(
        window=window,
        status="failed",
        error_code=secret,
        error_message=secret,
    )
    no_imagery = QuarterlyNDVIObservation(
        window=window,
        status="no_imagery",
        scene_provenance=_build_scene_provenance(_FakeSceneCollection([], []), 0),
    )

    failed_payload = failed.model_dump(mode="json")
    serialized = json.dumps(failed_payload, sort_keys=True)
    assert failed.error_code == "GEE_QUARTER_FAILED"
    assert failed.outcome_code == "HISTORICAL_NDVI_PROCESSING_FAILED"
    assert failed.cache_disposition == "retryable_failure"
    assert "error_message" not in failed_payload
    assert "api_key" not in serialized
    assert "earthengine.example" not in serialized

    assert no_imagery.error_code is None
    assert no_imagery.outcome_code == "HISTORICAL_NDVI_NO_IMAGERY"
    assert no_imagery.cache_disposition == "negative_no_imagery"
    assert no_imagery.scene_provenance is not None
    assert no_imagery.scene_provenance.scene_count == 0
    assert no_imagery.scene_provenance.scene_set_sha256 != "0" * 64


@pytest.mark.parametrize("status", ["complete", "partial", "no_imagery", "no_valid_pixels"])
def test_reached_quarters_require_scene_provenance(status: str) -> None:
    windows = plan_quarter_windows(2025, 2026, as_of_date="2026-07-12")
    window = next(item for item in windows if item.status == status) if status in {"complete", "partial"} else windows[0]
    payload: dict[str, object] = {"window": window, "status": status, "image_count": 0}
    if status in {"complete", "partial"}:
        payload.update(
            {
                "image_count": 1,
                "median_ndvi": 0.5,
                "mean_ndvi": 0.5,
                "stddev_ndvi": 0.1,
                "min_ndvi": 0.2,
                "max_ndvi": 0.8,
                "valid_pixel_count": 8,
                "roi_pixel_count": 10,
                "valid_pixel_coverage": 0.8,
            }
        )
    elif status == "no_valid_pixels":
        payload.update(
            {
                "image_count": 1,
                "valid_pixel_count": 0,
                "roi_pixel_count": 10,
                "valid_pixel_coverage": 0.0,
            }
        )
    with pytest.raises(ValidationError, match="场景来源"):
        QuarterlyNDVIObservation.model_validate(payload)


def _file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sample_result(root: Path) -> HistoricalNDVIResult:
    windows = plan_quarter_windows(2025, 2026, as_of_date="2026-07-12")
    observations: list[QuarterlyNDVIObservation] = []
    for index, window in enumerate(windows):
        if window.status == "not_reached":
            observations.append(QuarterlyNDVIObservation(window=window, status="not_reached"))
            continue
        relative = f"quarter_maps/{window.year}-Q{window.quarter}.png"
        image_path = root / relative
        image_path.parent.mkdir(parents=True, exist_ok=True)
        Image.new(
            "RGB",
            (360, 240),
            color=(45 + index * 8, 115 + index * 5, 62 + index * 7),
        ).save(image_path)
        median = 0.32 + index * 0.065
        image_count = 8 + index
        provenance = _build_scene_provenance(
            _FakeSceneCollection(
                [f"SCENE-{index:02d}-{item:03d}" for item in range(image_count)],
                [f"PRODUCT-{index:02d}-{item:03d}" for item in range(image_count)],
            ),
            image_count,
        )
        observations.append(
            QuarterlyNDVIObservation(
                window=window,
                status=window.status,
                image_count=image_count,
                median_ndvi=median,
                mean_ndvi=median - 0.01,
                stddev_ndvi=0.08,
                min_ndvi=max(-1.0, median - 0.25),
                max_ndvi=min(1.0, median + 0.20),
                valid_pixel_count=850,
                roi_pixel_count=1000,
                valid_pixel_coverage=0.85,
                quarter_map_png=relative,
                quarter_map_sha256=_file_sha256(image_path),
                scene_provenance=provenance,
            )
        )
    return HistoricalNDVIResult(
        task_id="history-test-2026",
        scope_label="测试地块",
        parcel_id="PARCEL-001",
        feature_id="feature-001",
        boundary_geometry_sha256="a" * 64,
        generated_at=datetime(2026, 7, 12, 12, 0, tzinfo=timezone.utc),
        as_of_date=date(2026, 7, 12),
        start_year=2025,
        end_year=2026,
        source=HistoricalNDVISource(max_cloud_pct=30, scale_m=10),
        quarters=observations,
    )


def test_persisted_result_is_utf8_atomic_round_trip_and_rejects_duplicates(tmp_path: Path) -> None:
    result = _sample_result(tmp_path)
    path = tmp_path / "result.json"
    metadata = save_historical_ndvi_result(result, path)
    loaded = load_historical_ndvi_result(path)

    assert loaded == result
    assert "测试地块" in path.read_text(encoding="utf-8")
    assert metadata["sha256"] == _file_sha256(path)
    assert not list(tmp_path.glob(".result.json.*.tmp"))
    assert loaded.source.provenance_schema_version == "agrisky.historical-ndvi-source/v1"
    assert loaded.source.result_schema_version == loaded.schema_version
    assert loaded.source.source_id == "google-earth-engine"
    assert loaded.source.algorithm_id == "agrisky.sentinel2.quarterly-ndvi"

    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["quarters"][1] = payload["quarters"][0]
    with pytest.raises(ValidationError, match="重复|完整季度槽位"):
        HistoricalNDVIResult.model_validate(payload)


def test_report_outputs_trend_year_panels_docx_and_blank_future_quarter(tmp_path: Path) -> None:
    pytest.importorskip("matplotlib")
    from docx import Document
    from historical_ndvi_report import generate_historical_ndvi_report

    result = _sample_result(tmp_path)
    updated = generate_historical_ndvi_report(result, tmp_path, artifact_root=tmp_path)

    assert (tmp_path / "historical_ndvi_trend.png").is_file()
    assert (tmp_path / "yearly_panels" / "2025-quarters.png").is_file()
    assert (tmp_path / "yearly_panels" / "2026-quarters.png").is_file()
    with Image.open(tmp_path / "historical_ndvi_trend.png") as trend_image:
        assert trend_image.mode == "RGB"
    with Image.open(tmp_path / "yearly_panels" / "2025-quarters.png") as panel_image:
        assert panel_image.mode == "RGB"
    report_path = tmp_path / "historical_ndvi_report.docx"
    assert report_path.is_file()
    assert updated.outputs.report_docx == "historical_ndvi_report.docx"
    assert load_historical_ndvi_result(tmp_path / "result.json") == updated

    document = Document(report_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "2025—2026 年季度 NDVI 时序监测报告" in text
    assert "未来季度保持空白" in text
    assert "不做月均替代或时间插值" in text
    assert "人工复核前不得作为最终定损结论" in "\n".join(
        paragraph.text for section in document.sections for paragraph in section.footer.paragraphs
    )
