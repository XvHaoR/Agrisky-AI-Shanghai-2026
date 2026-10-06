"""报告导出格式、完整性与受控下载回归测试。"""

from __future__ import annotations

import hashlib
import json
import sqlite3
import sys
import time
import types
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "api_gateway"))

import main
from excel_report import generate_excel_report
from growth_analysis import GROWTH_ANALYSIS_ALGORITHM_VERSION, GROWTH_ANALYSIS_SCHEMA_VERSION
from report_generator import generate_claim_report


_REPORT_BOUNDARY = {
    "type": "Polygon",
    "coordinates": [
        [[113.0, 34.0], [113.01, 34.0], [113.01, 34.01], [113.0, 34.01], [113.0, 34.0]]
    ],
}


def _document_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def _growth_payload(*, missing_first_row: bool = False) -> dict:
    rows = [
        {"value": 5, "label": "优", "area_mu": 40.0, "ratio": 0.4, "count": 400, "color": "#2F9E44"},
        {"value": 4, "label": "良", "area_mu": 30.0, "ratio": 0.3, "count": 300, "color": "#8CC152"},
        {"value": 3, "label": "中", "area_mu": 20.0, "ratio": 0.2, "count": 200, "color": "#E4C441"},
        {"value": 2, "label": "一般", "area_mu": 7.0, "ratio": 0.07, "count": 70, "color": "#EF8F35"},
        {"value": 1, "label": "差", "area_mu": 3.0, "ratio": 0.03, "count": 30, "color": "#D64545"},
    ]
    if missing_first_row:
        rows[-1]["area_mu"] = None
        rows[-1]["ratio"] = None
    return {
        "status": "success",
        "task_id": "growth-export-test",
        "method": "fixed",
        "n_classes": 5,
        "total_area_mu": 100.0,
        "valid_pixel_count": 1000,
        "class_breaks": [0.3, 0.45, 0.6, 0.75, 1.0],
        "raster": {
            "crs": "EPSG:32649",
            "nodata": -9999,
            "roi_pixel_count": 1200,
            "valid_pixel_count": 1000,
            "valid_pixel_coverage": 0.833333,
            "ndvi_source": "gee",
            "ndvi_source_label": "GEE Sentinel-2 SR NDVI",
            "ndvi_meta": {
                "source": "gee",
                "source_label": "GEE Sentinel-2 SR NDVI",
                "collection": "COPERNICUS/S2_SR_HARMONIZED",
                "formula": "NDVI = (B8 - B4) / (B8 + B4)",
                "start_date": "2026-06-01",
                "end_date": "2026-06-30",
                "max_cloud_pct": 30,
                "cloud_mask": "SCL 排除 0/1/2/3/8/9/10/11；云、阴影和雪边缘外扩 20 m",
                "cloud_mask_version": "s2-scl-v2-20m",
                "scale": 10,
                "image_count": 4,
            },
        },
        "summary": rows,
        "outputs": {},
        "message": "ok",
    }


def test_docx_marks_missing_growth_values_and_updates_fields(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    generate_claim_report(
        "CLAIM-DOCX-FORMAT",
        {
            "case": {"holder_name": "测试投保人", "crop_type": "rice"},
            "satellite": {
                "source_label": "GEE Sentinel-1 SAR",
                "start_date": "2026-05-01",
                "end_date": "2026-05-31",
                "damage_ratio": 0.2,
                "suspected_damage_area_mu": 20,
                "confidence": "high",
                "image_count": 3,
            },
            "growth": _growth_payload(missing_first_row=True),
            "loss_assessment": {
                "status": "success",
                "yield_loss_ratio": 0.2,
                "product_breakdown": [
                    {
                        "name_cn": "测试遥感产品",
                        "sensor": "S2",
                        "native_res_m": 10,
                        "decline_score": 0.2,
                        "weight": 1.0,
                        "contribution": 0.2,
                        "data_source": "authoritative-provider-source-with-version-and-provenance-2026-very-long",
                    }
                ],
            },
            "payout": {
                "crop_type": "rice",
                "sum_insured_per_mu": 1000,
                "insured_area_mu": 100,
                "total_sum_insured_yuan": 100000,
                "yield_loss_ratio": 0.2,
                "deductible_threshold": 0.1,
                "tier_label": "测试档",
                "payout_factor": 0.1,
                "payout_amount_yuan": 10000,
                "trace": ["最终以审核回执绑定的报告生成号为准"],
            },
            "_report_meta": {"generation_id": "v1.1-test", "snapshot_sha256": "a" * 64},
        },
        str(output),
    )

    text = _document_text(output)
    assert "起始含、截止不含" in text
    assert "部分长势等级的面积缺失" in text
    assert "部分长势等级的占比缺失" in text
    assert "无法可靠汇总“差/一般”区域" in text
    assert "未包含可用的本地 NDVI 分级预览图" in text
    assert "COPERNICUS/S2_SR_HARMONIZED" in text
    assert "云、阴影和雪边缘外扩 20 m" in text
    assert "83.3%" in text

    with zipfile.ZipFile(output) as archive:
        settings = archive.read("word/settings.xml").decode("utf-8")
        footer = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "updateFields" in settings
    assert "人工审核稿" in footer
    assert "keepNext" in document_xml

    document = Document(output)
    trace_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if "最终以审核回执绑定的报告生成号为准" in paragraph.text
    )
    assert document.paragraphs[trace_index].paragraph_format.keep_together is True
    assert document.paragraphs[trace_index - 1].paragraph_format.keep_with_next is True
    payout_table = next(
        table for table in document.tables
        if any("预估赔款（元）" in cell.text for row in table.rows for cell in row.cells)
    )
    assert all(
        paragraph.paragraph_format.keep_with_next is True
        for cell in payout_table.rows[-1].cells
        for paragraph in cell.paragraphs
    )


def test_excel_uses_numeric_cells_and_blocks_formula_injection(tmp_path: Path) -> None:
    output = tmp_path / "report.xlsx"
    generated_at = datetime(2026, 7, 13, 8, 30, 45, tzinfo=timezone(timedelta(hours=8)))
    generate_excel_report(
        "CLAIM-XLSX-FORMAT",
        {
            "case": {"holder_name": "=HYPERLINK(\"https://example.invalid\")", "crop_type": "rice"},
            "satellite": {
                "start_date": "2026-05-01",
                "end_date": "2026-05-31",
                "damage_ratio": 0.25,
                "confidence": "high",
                "image_count": 2,
            },
            "growth": _growth_payload(missing_first_row=True),
            "loss_assessment": {
                "status": "success",
                "yield_loss_ratio": 0.2,
                "product_breakdown": [
                    {
                        "name_cn": "测试遥感产品",
                        "sensor": "S2",
                        "native_res_m": 10,
                        "decline_score": 0.2,
                        "weight": 1.0,
                        "contribution": 0.2,
                        "data_source": "authoritative-provider-source-with-version-and-provenance-2026-very-long",
                    }
                ],
            },
        },
        str(output),
        generated_at=generated_at,
    )

    workbook = load_workbook(output, data_only=False)
    assert "结论依据映射" in workbook.sheetnames
    assert "依据索引" in workbook.sheetnames
    assert any(cell.value == "L-01§12" for cell in workbook["依据索引"]["A"])
    assert any(cell.value == "L-06§15" for cell in workbook["依据索引"]["A"])
    assert any("遥感查勘与灾损评估" == cell.value for cell in workbook["结论依据映射"]["A"])
    assessment = workbook["受灾评估"]
    holder_row = next(cell.row for cell in assessment["A"] if cell.value == "投保人")
    ratio_row = next(cell.row for cell in assessment["A"] if cell.value == "初筛受损比例")
    window_row = next(cell.row for cell in assessment["A"] if cell.value == "观测窗口")
    assert assessment.cell(holder_row, 2).value.startswith("'=")
    assert assessment.cell(holder_row, 2).data_type == "s"
    assert assessment.cell(ratio_row, 2).value == 0.25
    assert assessment.cell(ratio_row, 2).number_format == "0.0%"
    assert "起始含、截止不含" in assessment.cell(window_row, 2).value
    assert assessment.freeze_panes == "A4"
    assert assessment.print_area
    assert assessment.page_margins.top >= 0.8
    assert assessment.page_margins.header <= 0.3
    assert assessment.row_breaks.brk
    product_row = next(cell.row for cell in assessment["A"] if cell.value == "测试遥感产品")
    assert assessment.row_dimensions[product_row].height > 24

    growth = workbook["长势分析"]
    mask_row = next(cell.row for cell in growth["A"] if cell.value == "像元质量掩膜")
    assert f"A{mask_row}:B{mask_row}" in {str(cell_range) for cell_range in growth.merged_cells.ranges}
    assert growth.cell(mask_row, 3).value.startswith("SCL 排除")
    header_row = next(cell.row for cell in growth["A"] if cell.value == "等级")
    level_one_row = next(
        row for row in range(header_row + 1, growth.max_row + 1) if growth.cell(row, 1).value == 1
    )
    total_row = next(cell.row for cell in growth["A"] if cell.value == "合计")
    assert growth.cell(level_one_row, 4).value == "—"
    assert growth.cell(level_one_row, 5).value == "—"
    assert growth.cell(total_row, 4).value == "—"
    assert growth.cell(total_row, 5).value == "—"
    assert any("数据质量警告" in str(cell.value) for cell in growth["A"])
    assert growth.auto_filter.ref.endswith(str(level_one_row))
    assert growth.print_area
    assert workbook.properties.created == datetime(2026, 7, 13, 0, 30, 45)
    workbook.close()


def test_growth_quality_rules_detect_duplicates_and_each_break_error(tmp_path: Path) -> None:
    payload = _growth_payload()
    payload["summary"].append(
        {"value": 5, "label": "优", "area_mu": 0.0, "ratio": 0.0, "count": 0, "color": "#2F9E44"}
    )
    payload["class_breaks"] = [0.3, 0.2, 1.2]
    payload["valid_pixel_count"] = 999
    payload["raster"]["valid_pixel_coverage"] = 0.6
    docx_path = tmp_path / "quality.docx"
    xlsx_path = tmp_path / "quality.xlsx"
    data = {"case": {}, "growth": payload}
    generate_claim_report("CLAIM-QUALITY", data, str(docx_path))
    generate_excel_report("CLAIM-QUALITY", data, str(xlsx_path))

    docx_text = _document_text(docx_path)
    assert "阈值数量与等级数不一致" in docx_text
    assert "阈值超出 [-1, 1]" in docx_text
    assert "有效像元覆盖率仅 60.00%" in docx_text
    assert "阈值未严格递增" in docx_text
    assert "长势等级存在重复记录：5 级" in docx_text
    assert "各等级像元数合计 1,000" in docx_text

    workbook = load_workbook(xlsx_path)
    warning_text = "\n".join(str(cell.value or "") for cell in workbook["长势分析"]["A"])
    assert "阈值数量与等级数不一致" in warning_text
    assert "阈值超出 [-1, 1]" in warning_text
    assert "有效像元覆盖率仅 60.00%" in warning_text
    assert "阈值未严格递增" in warning_text
    assert "长势等级存在重复记录：5 级" in warning_text
    assert "各等级像元数合计 1,000" in warning_text
    workbook.close()


def test_excel_always_contains_explicit_growth_status_sheet(tmp_path: Path) -> None:
    output = tmp_path / "missing-growth.xlsx"
    generate_excel_report("CLAIM-NO-GROWTH", {"case": {}}, str(output))
    workbook = load_workbook(output)
    assert "长势分析" in workbook.sheetnames
    assert "未执行长势分析" in str(workbook["长势分析"]["A4"].value)
    assert any(
        cell.value == "长势结果状态" for cell in workbook["导出说明"]["A"]
    )
    workbook.close()


def _seed_report_case(claim_id: str, policy_id: str, growth: dict | None = None) -> None:
    boundary_sha256 = main._geometry_sha256(_REPORT_BOUNDARY)
    policy_version_id = f"{policy_id}:v1"
    con = main._db()
    con.execute(
        "INSERT INTO policies "
        "(policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at, policy_version_id) "
        "VALUES (?,?,?,?,?,?,?,?)",
        (
            policy_id,
            "报告测试投保人",
            "rice",
            "测试地点",
            json.dumps(_REPORT_BOUNDARY),
            100.0,
            "2026-07-13",
            policy_version_id,
        ),
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "RULE_DONE", "flood", "2026-06-30", "rice", "PLOT-REPORT", "2026-07-13"),
    )
    con.commit()
    con.close()
    main._save_result(
        claim_id,
        "satellite",
        {
            "source": "gee",
            "source_label": "GEE Sentinel-1 SAR",
            "start_date": "2026-05-01",
            "end_date": "2026-05-31",
            "damage_ratio": 0.2,
            "suspected_damage_area_mu": 20.0,
            "confidence": "high",
            "image_count": 3,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "compliance",
        {
            "status": "success",
            "damage_ratio": 0.2,
            "insured_area_mu": 100.0,
            "valid_damage_area_mu": 20.0,
            "excluded_area_mu": 80.0,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "rule",
        {
            "risk_level": "medium",
            "review_required": True,
            "rule_trace": ["测试规则"],
            "rule_version": "test_v1",
        },
    )
    main._save_result(
        claim_id,
        "payout",
        {
            "crop_type": "rice",
            "sum_insured_per_mu": 1000.0,
            "insured_area_mu": 100.0,
            "total_sum_insured_yuan": 100000.0,
            "yield_loss_ratio": 0.2,
            "deductible_threshold": 0.1,
            "payout_factor": 0.1,
            "tier_label": "测试档",
            "payout_amount_yuan": 10000.0,
        },
    )
    if growth:
        growth["schema_version"] = GROWTH_ANALYSIS_SCHEMA_VERSION
        growth["algorithm_version"] = GROWTH_ANALYSIS_ALGORITHM_VERSION
        growth.setdefault("raster", {}).setdefault("ndvi_meta", {}).update(
            {
                "boundary_sha256": boundary_sha256,
                "boundary_hash_scheme": "canonical_geometry_v1",
            }
        )
        main._save_result(claim_id, "growth", growth)


def _bind_report_growth_artifacts(output_root: Path, growth: dict) -> None:
    """Create the three authoritative growth artifacts required by report export."""
    task_id = growth["task_id"]
    task_dir = output_root / "growth" / task_id
    task_dir.mkdir(parents=True)

    ndvi_path = task_dir / "ndvi_clip.tif"
    ndvi_path.write_bytes(b"fixture-ndvi-raster")
    preview_path = task_dir / "class_preview.png"
    Image.new("RGB", (16, 16), "#2F9E44").save(preview_path, format="PNG")
    report_path = task_dir / "growth_report.docx"
    source_document = Document()
    source_document.add_paragraph("RAW-SOURCE-MARKER-SHOULD-NOT-BE-COPIED")
    source_document.save(report_path)

    growth["outputs"] = {
        "ndvi_clip_tif": f"/outputs/growth/{task_id}/{ndvi_path.name}",
        "class_preview_png": f"/outputs/growth/{task_id}/{preview_path.name}",
        "report_docx": f"/outputs/growth/{task_id}/{report_path.name}",
    }
    growth["raster"].update(
        {
            "ndvi_clip_sha256": hashlib.sha256(ndvi_path.read_bytes()).hexdigest(),
            "ndvi_clip_size_bytes": ndvi_path.stat().st_size,
        }
    )
    main._bind_growth_artifact_integrity(growth, task_dir)


def test_report_bundle_is_immutable_hashed_and_uses_standard_growth_appendix(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    monkeypatch.setenv("AGRISKY_OUTPUT_ROOT", str(tmp_path))
    claim_id = f"CLAIM-BUNDLE-{uuid.uuid4().hex[:8]}"
    policy_id = f"POL-BUNDLE-{uuid.uuid4().hex[:8]}"
    growth = _growth_payload()
    growth["task_id"] = f"growth-{uuid.uuid4().hex[:8]}"
    _bind_report_growth_artifacts(tmp_path, growth)
    _seed_report_case(claim_id, policy_id, growth)

    with TestClient(main.app) as api:
        first = api.post("/api/v1/tools/generate_report", json={"claim_id": claim_id})
        assert first.status_code == 200, first.text
        payload = first.json()
        bundle = main._output_artifact_path(payload["bundle_zip_url"])
        assert bundle is not None
        first_hash = hashlib.sha256(bundle.read_bytes()).hexdigest()
        first_mtime = bundle.stat().st_mtime_ns

        with zipfile.ZipFile(bundle) as archive:
            names = archive.namelist()
            assert len(names) == len(set(names))
            manifest = json.loads(archive.read("manifest.json"))
            for entry in manifest["file_integrity"]:
                content = archive.read(entry["filename"])
                assert hashlib.sha256(content).hexdigest() == entry["sha256"]
                assert len(content) == entry["size_bytes"]
            growth_name = manifest["files"]["growth_report"]
            legal_basis_name = manifest["files"]["legal_basis"]
            legal_basis = json.loads(archive.read(legal_basis_name))
            assert legal_basis["schema_version"] == "agrisky-case-legal-basis-v1.0"
            assert legal_basis["claim_id"] == claim_id
            assert any("L-02§21" in row["basis_ids"] for row in legal_basis["basis"])
            assert any(item["id"] == "L-03§18" for item in legal_basis["references"])
            extracted_growth = tmp_path / "extracted-growth.docx"
            extracted_growth.write_bytes(archive.read(growth_name))

        assert "RAW-SOURCE-MARKER-SHOULD-NOT-BE-COPIED" not in _document_text(extracted_growth)
        assert "作物长势监测报告" in _document_text(extracted_growth)
        assert payload["growth_status"] == "included"
        assert any(row["stage"] == "赔款测算" for row in payload["legal_basis"])
        assert any(item["kind"] == "legal_basis" for item in payload["artifacts"])
        assert payload["bundle_sha256"] == first_hash
        assert manifest["snapshot_sha256"] == payload["snapshot_sha256"]

        time.sleep(0.01)
        second = api.post("/api/v1/tools/generate_report", json={"claim_id": claim_id})
        assert second.status_code == 200
        assert second.json()["generation_id"] == payload["generation_id"]
        assert bundle.stat().st_mtime_ns == first_mtime
        assert hashlib.sha256(bundle.read_bytes()).hexdigest() == first_hash

        assert api.get(payload["bundle_zip_url"]).status_code == 403
        bundle_artifact = next(item for item in payload["artifacts"] if item["kind"] == "bundle")
        controlled = api.get(bundle_artifact["download_url"])
        assert controlled.status_code == 200
        assert hashlib.sha256(controlled.content).hexdigest() == first_hash

        original_bundle = bundle.read_bytes()
        bundle.write_bytes(b"tampered")
        rejected_tamper = api.post(
            f"/api/v1/cases/{claim_id}/human_review",
            json={
                "decision": "approved",
                "generation_id": payload["generation_id"],
                "comment": "完整性测试",
            },
        )
        assert rejected_tamper.status_code == 409
        assert api.get(f"/api/v1/cases/{claim_id}").json()["state"] == "REPORT_DRAFTED"
        bundle.write_bytes(original_bundle)

        review_request = {
            "decision": "approved",
            "generation_id": payload["generation_id"],
            "comment": "核对报告与附件后同意归档",
            "idempotency_key": f"approve:{claim_id}",
        }
        approved = api.post(f"/api/v1/cases/{claim_id}/human_review", json=review_request)
        assert approved.status_code == 200, approved.text
        review = approved.json()
        assert review["state"] == "ARCHIVED"
        assert review["receipt"]["report"]["snapshot_sha256"] == payload["snapshot_sha256"]
        assert review["receipt"]["report"]["bundle_sha256"] == first_hash
        receipt_download = api.get(review["receipt_download_url"])
        assert receipt_download.status_code == 200
        assert hashlib.sha256(receipt_download.content).hexdigest() == review["receipt_sha256"]
        receipt = json.loads(receipt_download.content)
        assert receipt["review_id"] == review["review_id"]
        assert receipt["decision"] == "approved"
        assert receipt["actor"] != "forged-user"
        assert api.post(f"/api/v1/cases/{claim_id}/human_review", json=review_request).json()["review_id"] == review["review_id"]
        with zipfile.ZipFile(bundle) as archive:
            assert json.loads(archive.read("manifest.json"))["report_status"] == "draft_pending_human_review"


def test_admin_session_authorizes_protected_api(monkeypatch) -> None:
    username = f"report-admin-{uuid.uuid4().hex[:8]}"
    password = "test-admin-password"
    salt = uuid.uuid4().hex
    con = main._db()
    con.execute(
        "INSERT INTO accounts (username, pwd_hash, salt, holder_name, role, created_at) VALUES (?,?,?,?,?,?)",
        (username, main._hash_pwd(password, salt), salt, "测试管理员", "admin", "2026-07-13T00:00:00"),
    )
    con.commit()
    con.close()
    monkeypatch.setattr(main, "AUTH_ENABLED", True)
    monkeypatch.setattr(main, "COOKIE_SECURE", False)

    with TestClient(main.app) as api:
        assert api.get("/api/v1/cases").status_code == 401
        login = api.post("/api/v1/auth/login", json={"username": username, "password": password})
        assert login.status_code == 200
        assert "token" not in login.json()
        assert "httponly" in login.headers.get("set-cookie", "").lower()
        assert api.get("/api/v1/cases").status_code == 200
        assert api.post("/api/v1/auth/logout").status_code == 200
        assert api.get("/api/v1/cases").status_code == 401


def test_report_directory_is_not_part_of_static_mounts(monkeypatch) -> None:
    """双斜杠/点段不能绕过受控报告下载。"""
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    mount_paths = {getattr(route, "path", None) for route in main.app.routes if route.__class__.__name__ == "Mount"}
    assert "/outputs" not in mount_paths
    assert "/outputs/growth" not in mount_paths
    assert "/outputs/screening" not in mount_paths

    route_paths = {getattr(route, "path", None) for route in main.app.routes}
    assert "/api/v1/cases/{claim_id}/evidence/growth/{task_id}/{filename}" in route_paths
    assert "/api/v1/cases/{claim_id}/evidence/screening/{filename}" in route_paths

    candidates = [
        "/outputs/reports/not-registered.docx",
        "/outputs//reports/not-registered.docx",
        "/outputs/growth/../reports/not-registered.docx",
        "/outputs/%2Freports/not-registered.docx",
        "/outputs/growth/not-registered/growth_report.docx",
        "/outputs/screening/not-registered/evidence.png",
    ]
    with TestClient(main.app) as api:
        for path in candidates:
            assert api.get(path).status_code in {403, 404}


def test_policyholder_cookie_reaches_only_scoped_agent(monkeypatch) -> None:
    username = f"holder-{uuid.uuid4().hex[:8]}"
    password = "policyholder-password"
    policy_id = f"POL-HOLDER-{uuid.uuid4().hex[:8]}"
    salt = uuid.uuid4().hex
    con = main._db()
    con.execute(
        "INSERT INTO accounts (username, pwd_hash, salt, holder_name, role, created_at) VALUES (?,?,?,?,?,?)",
        (username, main._hash_pwd(password, salt), salt, "限权投保人", "policyholder", "2026-07-13T00:00:00"),
    )
    con.execute(
        "INSERT INTO policies "
        "(policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at, holder_account) "
        "VALUES (?,?,?,?,?,?,?,?)",
        (policy_id, "限权投保人", "rice", "测试地点", None, 1.0, "2026-07-13", username),
    )
    con.commit()
    con.close()

    captured: dict = {}

    def fake_agent(messages, scope=None):
        captured["scope"] = scope
        return {"reply": "ok", "messages": messages, "trace": [], "error": None}

    monkeypatch.setitem(sys.modules, "agent.agent_runtime", types.SimpleNamespace(run_agent_api=fake_agent))
    monkeypatch.setattr(main, "AUTH_ENABLED", True)
    monkeypatch.setattr(main, "COOKIE_SECURE", False)
    with TestClient(main.app) as api:
        login = api.post("/api/v1/auth/login", json={"username": username, "password": password})
        assert login.status_code == 200
        response = api.post("/api/v1/agent/chat", json={"messages": []})
        assert response.status_code == 200, response.text
    assert captured["scope"]["username"] == username
    assert captured["scope"]["policy_ids"] == [policy_id]


def test_admin_cookie_reaches_agent_with_global_scope(monkeypatch) -> None:
    username = f"agent-admin-{uuid.uuid4().hex[:8]}"
    password = "agent-admin-password"
    salt = uuid.uuid4().hex
    con = main._db()
    con.execute(
        "INSERT INTO accounts (username, pwd_hash, salt, holder_name, role, created_at) VALUES (?,?,?,?,?,?)",
        (username, main._hash_pwd(password, salt), salt, "智能体管理员", "admin", "2026-07-22T00:00:00"),
    )
    con.commit()
    con.close()

    captured: dict = {}

    def fake_agent(messages, scope=None):
        captured["scope"] = scope
        return {"reply": "ok", "messages": messages, "trace": [], "error": None}

    monkeypatch.setitem(sys.modules, "agent.agent_runtime", types.SimpleNamespace(run_agent_api=fake_agent))
    monkeypatch.setattr(main, "AUTH_ENABLED", True)
    monkeypatch.setattr(main, "COOKIE_SECURE", False)
    with TestClient(main.app) as api:
        login = api.post("/api/v1/auth/login", json={"username": username, "password": password})
        assert login.status_code == 200
        response = api.post("/api/v1/agent/chat", json={"messages": []})
        assert response.status_code == 200, response.text
    assert captured["scope"] is None


def test_preliminary_excel_uses_controlled_download(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-EXCEL-{uuid.uuid4().hex[:8]}"
    policy_id = f"POL-EXCEL-{uuid.uuid4().hex[:8]}"
    con = main._db()
    con.execute(
        "INSERT INTO policies "
        "(policy_id, holder_name, crop_type, address, boundary_geojson, area_mu, created_at) "
        "VALUES (?,?,?,?,?,?,?)",
        (policy_id, "Excel 测试投保人", "rice", "测试地点", None, 100.0, "2026-07-13"),
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, "COMPLIANCE_DONE", "flood", "2026-06-30", "rice", "PLOT-X", "2026-07-13"),
    )
    con.commit()
    con.close()
    main._save_result(
        claim_id,
        "compliance",
        {
            "status": "success",
            "damage_ratio": 0.2,
            "insured_area_mu": 100.0,
            "valid_damage_area_mu": 20.0,
            "excluded_area_mu": 80.0,
        },
    )

    with TestClient(main.app) as api:
        generated = api.post("/api/v1/tools/generate_excel_report", json={"claim_id": claim_id})
        assert generated.status_code == 200, generated.text
        payload = generated.json()
        assert payload["excel_url"] == f"/api/v1/cases/{claim_id}/preliminary-excel"
        downloaded = api.get(payload["excel_url"])
        assert downloaded.status_code == 200
        assert zipfile.is_zipfile(__import__("io").BytesIO(downloaded.content))
        assert api.get(f"/outputs/reports/{claim_id}.xlsx").status_code == 403


def test_unknown_login_uses_dummy_hash_and_is_rate_limited(monkeypatch) -> None:
    username = f"missing-{uuid.uuid4().hex}"
    original_hash = main._hash_pwd
    seen_salts: list[str] = []

    def tracked_hash(password: str, salt: str) -> str:
        seen_salts.append(salt)
        return original_hash(password, salt)

    monkeypatch.setattr(main, "_hash_pwd", tracked_hash)
    monkeypatch.setattr(main, "_LOGIN_ACCOUNT_MAX_FAILURES", 2)
    monkeypatch.setattr(main, "_LOGIN_IP_MAX_FAILURES", 100)
    monkeypatch.setattr(main, "_LOGIN_BLOCK_SECONDS", 60)
    with main._LOGIN_RATE_LOCK:
        main._LOGIN_FAILURES.clear()
        main._LOGIN_BLOCKED_UNTIL.clear()
    try:
        with TestClient(main.app) as api:
            first = api.post("/api/v1/auth/login", json={"username": username, "password": "wrong"})
            second = api.post("/api/v1/auth/login", json={"username": username, "password": "wrong"})
        assert first.status_code == 401
        assert second.status_code == 429
        assert int(second.headers["Retry-After"]) >= 1
        assert seen_salts == [main._DUMMY_AUTH_SALT, main._DUMMY_AUTH_SALT]
    finally:
        with main._LOGIN_RATE_LOCK:
            main._LOGIN_FAILURES.clear()
            main._LOGIN_BLOCKED_UNTIL.clear()


def test_login_rejects_oversized_credentials_and_prunes_expired_buckets() -> None:
    with TestClient(main.app) as api:
        assert api.post(
            "/api/v1/auth/login",
            json={"username": "u" * 129, "password": "password"},
        ).status_code == 422
        assert api.post(
            "/api/v1/auth/login",
            json={"username": "user", "password": "p" * 1025},
        ).status_code == 422

    now = time.monotonic()
    with main._LOGIN_RATE_LOCK:
        main._LOGIN_FAILURES["account:expired"] = [now - main._LOGIN_WINDOW_SECONDS - 1]
        main._LOGIN_BLOCKED_UNTIL["account:expired-block"] = now - 1
        main._prune_login_rate_state(now)
        assert "account:expired" not in main._LOGIN_FAILURES
        assert "account:expired-block" not in main._LOGIN_BLOCKED_UNTIL


def test_human_review_state_and_audit_commit_atomically(monkeypatch) -> None:
    claim_id = f"CLAIM-REVIEW-{uuid.uuid4().hex[:8]}"
    conflicting_log_id = uuid.uuid4().hex[:8]
    review_id = uuid.uuid4().hex
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, None, "REPORT_DRAFTED", "flood", "2026-06-30", "rice", "PLOT-R", "2026-07-13"),
    )
    con.execute(
        "INSERT INTO audit_log VALUES (?,?,?,?,?,?,?,?,?)",
        (conflicting_log_id, "OTHER-CLAIM", "fixture", None, None, None, None, "test", "2026-07-13T00:00:00Z"),
    )
    con.execute(
        "INSERT INTO case_results (claim_id, step, result_json, created_at) VALUES (?,?,?,?)",
        (claim_id, "report", json.dumps({"generation_id": "v1-review-test"}), "2026-07-13T00:00:00Z"),
    )
    con.commit()
    con.close()

    generated_ids = iter([review_id, conflicting_log_id])
    monkeypatch.setattr(main.uuid, "uuid4", lambda: types.SimpleNamespace(hex=next(generated_ids)))
    monkeypatch.setattr(
        main,
        "_verify_report_for_review",
        lambda *_args: {
            "generation_id": "v1-review-test",
            "template_version": "v1.1",
            "snapshot_sha256": "1" * 64,
            "bundle_sha256": "2" * 64,
            "manifest_sha256": "3" * 64,
            "artifacts": [],
        },
    )
    with pytest.raises(sqlite3.IntegrityError):
        main._record_human_review_decision(
            claim_id,
            decision="approved",
            actor="admin",
            generation_id="v1-review-test",
        )

    con = main._db()
    state = con.execute("SELECT state FROM cases WHERE claim_id = ?", (claim_id,)).fetchone()["state"]
    audit_count = con.execute("SELECT COUNT(*) FROM audit_log WHERE claim_id = ?", (claim_id,)).fetchone()[0]
    receipt_count = con.execute("SELECT COUNT(*) FROM review_receipts WHERE claim_id = ?", (claim_id,)).fetchone()[0]
    con.close()
    assert state == "REPORT_DRAFTED"
    assert audit_count == 0
    assert receipt_count == 0
