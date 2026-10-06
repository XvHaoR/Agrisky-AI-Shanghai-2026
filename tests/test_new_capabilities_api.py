"""API integration for KML preflight, historical NDVI, parcel growth, and report bundling."""

from __future__ import annotations

import hashlib
import io
import json
import uuid
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

import main
from historical_ndvi import (
    HistoricalNDVIOutputs,
    HistoricalNDVIResult,
    HistoricalNDVISource,
    QuarterlyNDVIObservation,
    QuarterlySceneProvenance,
    _scene_set_digest,
    plan_quarter_windows,
    save_historical_ndvi_result,
)
from growth_analysis import (
    GROWTH_ANALYSIS_ALGORITHM_VERSION,
    GROWTH_ANALYSIS_SCHEMA_VERSION,
)
from parcel_growth import (
    CLASSIFICATION_SCHEME_VERSION,
    PARCEL_GROWTH_ALGORITHM_VERSION,
    PARCEL_GROWTH_SCHEMA_VERSION,
)


def _boundary() -> dict:
    return {
        "type": "Polygon",
        "coordinates": [[
            [114.0, 34.0],
            [114.001, 34.0],
            [114.001, 34.001],
            [114.0, 34.001],
            [114.0, 34.0],
        ]],
    }


def _kml_zip(*, add_ignored_file: bool = False) -> bytes:
    coordinates = "114.0,34.0 114.001,34.0 114.001,34.001 114.0,34.001 114.0,34.0"
    kml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>测试地块</name>'
        '<Placemark><name>地块一</name><Polygon><outerBoundaryIs><LinearRing><coordinates>'
        f"{coordinates}"
        '</coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>'
        '</Document></kml>'
    ).encode("utf-8")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("测试/地块一.kml", kml)
        if add_ignored_file:
            archive.writestr("测试/readme.txt", "second immutable preflight")
    return buffer.getvalue()


def _seed_case(claim_id: str, policy_id: str, *, state: str = "RULE_DONE") -> str:
    result = main._save_policy(
        policy_id,
        "接口测试合作社",
        "rice",
        "测试地区",
        _boundary(),
        f"{policy_id}:v1",
    )
    con = main._db()
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?,?)",
        (claim_id, policy_id, state, "flood", "2025-09-01", "rice", "PLOT-API", "2026-07-13"),
    )
    con.commit()
    con.close()
    return result["boundary_sha256"]


def _fake_history_job(**kwargs) -> dict:
    output_dir = Path(kwargs["output_dir"])
    windows = plan_quarter_windows(
        kwargs["start_year"], kwargs["end_year"], as_of_date=kwargs["as_of_date"]
    )
    observations = [
        QuarterlyNDVIObservation(
            window=window,
            status=window.status,
            image_count=4,
            median_ndvi=0.62,
            mean_ndvi=0.61,
            stddev_ndvi=0.05,
            min_ndvi=0.35,
            max_ndvi=0.85,
            valid_pixel_count=90,
            roi_pixel_count=100,
            valid_pixel_coverage=0.9,
            scene_provenance=QuarterlySceneProvenance(
                scene_count=4,
                scenes=(
                    records := [
                        {
                            "scene_id": f"SCENE-{window.year}-Q{window.quarter}-{index}",
                            "product_id": f"PRODUCT-{window.year}-Q{window.quarter}-{index}",
                        }
                        for index in range(4)
                    ]
                ),
                scene_set_sha256=_scene_set_digest(records),
            ),
        )
        for window in windows
    ]
    panels: dict[str, str] = {}
    for year in range(kwargs["start_year"], kwargs["end_year"] + 1):
        relative = f"yearly_panels/{year}-quarters.png"
        path = output_dir / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (120, 80), "green").save(path)
        panels[str(year)] = relative
    Image.new("RGB", (120, 80), "blue").save(output_dir / "historical_ndvi_trend.png")
    document = Document()
    document.add_heading("历史季度 NDVI 时序监测报告", level=1)
    document.save(output_dir / "historical_ndvi_report.docx")
    result = HistoricalNDVIResult(
        task_id=kwargs["task_id"],
        scope_label=kwargs["scope_label"],
        boundary_geometry_sha256=kwargs["boundary_geometry_sha256"],
        generated_at=datetime.now(timezone.utc),
        as_of_date=kwargs["as_of_date"],
        start_year=kwargs["start_year"],
        end_year=kwargs["end_year"],
        source=HistoricalNDVISource(
            max_cloud_pct=kwargs["max_cloud_pct"], scale_m=kwargs["scale_m"]
        ),
        quarters=observations,
        outputs=HistoricalNDVIOutputs(
            result_json="result.json",
            trend_chart_png="historical_ndvi_trend.png",
            yearly_panel_pngs=panels,
            report_docx="historical_ndvi_report.docx",
        ),
    )
    save_historical_ndvi_result(result, output_dir / "result.json")
    return result.model_dump(mode="json")


def _seed_growth_result(
    output_root: Path,
    boundary_sha256: str,
    *,
    task_id: str | None = None,
) -> tuple[dict, Path]:
    """Create a current, hash-bound growth result suitable for downstream APIs."""
    task_id = task_id or f"growth-{uuid.uuid4().hex[:8]}"
    task_dir = output_root / "growth" / task_id
    task_dir.mkdir(parents=True)

    ndvi_path = task_dir / "ndvi_clip.tif"
    ndvi_path.write_bytes(b"registered-test-raster")
    Image.new("RGB", (120, 80), "green").save(task_dir / "class_preview.png")
    document = Document()
    document.add_heading("Test growth report", level=1)
    document.save(task_dir / "growth_report.docx")

    ndvi_sha256 = hashlib.sha256(ndvi_path.read_bytes()).hexdigest()
    output_prefix = f"/outputs/growth/{task_id}"
    growth = {
        "schema_version": GROWTH_ANALYSIS_SCHEMA_VERSION,
        "algorithm_version": GROWTH_ANALYSIS_ALGORITHM_VERSION,
        "status": "success",
        "task_id": task_id,
        "method": "fixed",
        "n_classes": 5,
        "total_area_mu": 10.0,
        "valid_pixel_count": 100,
        "class_breaks": [0.3, 0.45, 0.6, 0.75, 1.0],
        "summary": [
            {
                "value": level,
                "label": label,
                "count": 20,
                "ratio": 0.2,
                "area_mu": 2.0,
                "color": color,
            }
            for level, label, color in [
                (1, "poor", "#d64545"),
                (2, "fair", "#ef8f35"),
                (3, "medium", "#e4c441"),
                (4, "good", "#8cc152"),
                (5, "excellent", "#2f9e44"),
            ]
        ],
        "raster": {
            "ndvi_source": "gee",
            "ndvi_clip_size_bytes": ndvi_path.stat().st_size,
            "ndvi_clip_sha256": ndvi_sha256,
            "ndvi_meta": {
                "source": "gee",
                "boundary_sha256": boundary_sha256,
                "boundary_hash_scheme": "canonical_geometry_v1",
            },
        },
        "outputs": {
            "ndvi_clip_tif": f"{output_prefix}/ndvi_clip.tif",
            "class_preview_png": f"{output_prefix}/class_preview.png",
            "report_docx": f"{output_prefix}/growth_report.docx",
        },
    }
    main._bind_growth_artifact_integrity(growth, task_dir)
    return growth, ndvi_path


def _save_report_prerequisites(claim_id: str, boundary_sha256: str, growth: dict) -> None:
    main._save_result(
        claim_id,
        "satellite",
        {
            "status": "success",
            "source": "gee",
            "damage_ratio": 0.2,
            "suspected_damage_area_mu": 2.0,
            "confidence": "high",
            "image_count": 3,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(claim_id, "growth", growth)
    main._save_result(
        claim_id,
        "compliance",
        {
            "status": "success",
            "damage_ratio": 0.2,
            "insured_area_mu": 10.0,
            "valid_damage_area_mu": 2.0,
            "excluded_area_mu": 0.0,
            "boundary_sha256": boundary_sha256,
            "boundary_hash_scheme": "canonical_geometry_v1",
        },
    )
    main._save_result(
        claim_id,
        "rule",
        {"risk_level": "medium", "review_required": True, "rule_trace": [], "rule_version": "test"},
    )
    main._save_result(
        claim_id,
        "payout",
        {
            "crop_type": "rice",
            "sum_insured_per_mu": 1000.0,
            "insured_area_mu": 10.0,
            "total_sum_insured_yuan": 10000.0,
            "yield_loss_ratio": 0.2,
            "deductible_threshold": 0.1,
            "payout_factor": 0.1,
            "tier_label": "测试档",
            "payout_amount_yuan": 1000.0,
        },
    )


def test_new_write_endpoints_are_authenticated(monkeypatch) -> None:
    monkeypatch.setattr(main, "AUTH_ENABLED", True)
    with TestClient(main.app) as api:
        response = api.post(
            "/api/v1/policies/parcels/preflight",
            files={"archive_file": ("parcels.zip", _kml_zip(), "application/zip")},
        )
        analyst_confirmation = api.post(
            "/api/v1/policies/parcels/preflights/preflight-000000000000000000000000/confirm",
            headers={"x-agrisky-api-key": main.INTERNAL_AGENT_API_KEY},
            json={
                "manifest_sha256": "0" * 64,
                "mappings": [{
                    "proposed_parcel_id": "proposed",
                    "policy_id": "POL",
                    "policy_version_id": "POL:v1",
                    "parcel_id": "PARCEL",
                }],
                "confirmation_statement": "I_CONFIRM_REVIEWED_PARCEL_MAPPINGS",
            },
        )
    assert response.status_code == 401
    assert analyst_confirmation.status_code == 403


def test_preflight_history_parcel_analysis_and_report_bundle(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    monkeypatch.setattr(main, "_run_historical_ndvi_job", _fake_history_job)
    claim_id = f"CLAIM-NEW-{uuid.uuid4().hex[:8]}"
    policy_id = f"POL-NEW-{uuid.uuid4().hex[:8]}"
    boundary_sha256 = _seed_case(claim_id, policy_id)

    growth, ndvi_path = _seed_growth_result(tmp_path, boundary_sha256)
    _save_report_prerequisites(claim_id, boundary_sha256, growth)

    def fake_parcel_job(raster_path: Path, manifest: dict, class_breaks: list[float]) -> dict:
        assert raster_path == ndvi_path
        assert manifest["parcels"][0]["mapping"]["confirmed"] is True
        return {
            "schema_version": PARCEL_GROWTH_SCHEMA_VERSION,
            "algorithm_version": PARCEL_GROWTH_ALGORITHM_VERSION,
            "classification_scheme_version": CLASSIFICATION_SCHEME_VERSION,
            "boundary_hash_scheme": "canonical_geometry_v1",
            "boundary_sha256": boundary_sha256,
            "source_crs": "EPSG:4326",
            "manifest": {"manifest_sha256": manifest["manifest_sha256"]},
            "class_breaks": class_breaks,
            "n_classes": 5,
            "valid_coverage_threshold": 0.7,
            "raster": {
                "path": str(raster_path),
                "filename": raster_path.name,
                "sha256": hashlib.sha256(raster_path.read_bytes()).hexdigest(),
                "crs": "EPSG:4326",
                "width": 10,
                "height": 10,
            },
            "features": [{"parcel_id": "PARCEL-001", "feature_id": "feature-1", "quality_status": "ok"}],
            "parcels": [{"parcel_id": "PARCEL-001", "quality_status": "ok"}],
            "overall": {"parcel_count": 1, "feature_count": 1, "valid_pixel_coverage": 0.9},
        }

    monkeypatch.setattr(main, "_run_parcel_growth_job", fake_parcel_job)

    with TestClient(main.app) as api:
        preflight = api.post(
            "/api/v1/policies/parcels/preflight",
            files={"archive_file": ("parcels.zip", _kml_zip(), "application/zip")},
        )
        assert preflight.status_code == 200, preflight.text
        preflight_payload = preflight.json()
        assert preflight_payload["direct_import_allowed"] is False
        proposed = preflight_payload["parcels"][0]["proposed_parcel_id"]
        confirmation_body = {
            "manifest_sha256": preflight_payload["manifest_sha256"],
            "confirmation_statement": "I_CONFIRM_REVIEWED_PARCEL_MAPPINGS",
            "idempotency_key": f"confirm-{uuid.uuid4().hex}",
            "mappings": [{
                "proposed_parcel_id": proposed,
                "policy_id": policy_id,
                "policy_version_id": f"{policy_id}:v1",
                "parcel_id": "PARCEL-001",
                "resolution_note": "",
            }],
        }
        stale = api.post(
            f"/api/v1/policies/parcels/preflights/{preflight_payload['preflight_id']}/confirm",
            json={**confirmation_body, "manifest_sha256": "0" * 64},
        )
        assert stale.status_code == 409
        confirmed = api.post(
            f"/api/v1/policies/parcels/preflights/{preflight_payload['preflight_id']}/confirm",
            json=confirmation_body,
        )
        assert confirmed.status_code == 200, confirmed.text
        receipt = confirmed.json()
        assert receipt["direct_import_performed"] is False
        assert receipt["confirmation_statement"] == "I_CONFIRM_REVIEWED_PARCEL_MAPPINGS"
        assert receipt["request_sha256"]
        assert receipt["mappings"][0]["resolution_note"] == ""
        receipt_download = api.get(receipt["receipt_download_url"])
        assert receipt_download.status_code == 200
        assert hashlib.sha256(receipt_download.content).hexdigest() == receipt_download.headers["x-content-sha256"]
        assert receipt_download.headers["x-receipt-sha256"] == receipt["receipt_sha256"]
        assert api.post(
            f"/api/v1/policies/parcels/preflights/{preflight_payload['preflight_id']}/confirm",
            json=confirmation_body,
        ).json()["confirmation_batch_id"] == receipt["confirmation_batch_id"]
        parcels = api.get(f"/api/v1/cases/{claim_id}/parcels")
        assert parcels.status_code == 200
        assert parcels.json()["items"][0]["parcel_id"] == "PARCEL-001"
        assert parcels.json()["items"][0]["receipt_sha256"] == receipt["receipt_sha256"]

        second_preflight = api.post(
            "/api/v1/policies/parcels/preflight",
            files={
                "archive_file": (
                    "parcels-second.zip",
                    _kml_zip(add_ignored_file=True),
                    "application/zip",
                )
            },
        ).json()
        duplicate_version = api.post(
            f"/api/v1/policies/parcels/preflights/{second_preflight['preflight_id']}/confirm",
            json={
                "manifest_sha256": second_preflight["manifest_sha256"],
                "confirmation_statement": "I_CONFIRM_REVIEWED_PARCEL_MAPPINGS",
                "mappings": [{
                    "proposed_parcel_id": second_preflight["parcels"][0]["proposed_parcel_id"],
                    "policy_id": policy_id,
                    "policy_version_id": f"{policy_id}:v1",
                    "parcel_id": "PARCEL-SECOND",
                }],
            },
        )
        assert duplicate_version.status_code == 409
        assert "已有不可变地块确认批次" in duplicate_version.json()["detail"]

        history = api.post(
            "/api/v1/tools/run_historical_ndvi_by_claim",
            json={
                "claim_id": claim_id,
                "start_year": 2025,
                "end_year": 2025,
                "as_of_date": "2025-12-31",
            },
        )
        assert history.status_code == 200, history.text
        history_payload = history.json()
        assert history_payload["boundary_geometry_sha256"] == boundary_sha256
        assert {item["kind"] for item in history_payload["artifacts"]} >= {
            "historical_ndvi_report", "historical_ndvi_result", "historical_ndvi_trend",
            "historical_ndvi_year_panel_2025",
        }
        trend = next(item for item in history_payload["artifacts"] if item["kind"] == "historical_ndvi_trend")
        downloaded = api.get(trend["download_url"])
        assert downloaded.status_code == 200
        assert hashlib.sha256(downloaded.content).hexdigest() == trend["sha256"]

        parcel_growth = api.post(
            "/api/v1/tools/run_parcel_growth_by_claim", json={"claim_id": claim_id}
        )
        assert parcel_growth.status_code == 200, parcel_growth.text
        parcel_payload = parcel_growth.json()
        assert parcel_payload["parcels"][0]["parcel_id"] == "PARCEL-001"
        assert "path" not in parcel_payload["raster"]

        report = api.post("/api/v1/tools/generate_report", json={"claim_id": claim_id})
        assert report.status_code == 200, report.text
        report_payload = report.json()
        assert report_payload["historical_ndvi_status"] == "included"
        assert report_payload["historical_ndvi_years"] == [2025]
        assert report_payload["parcel_growth_status"] == "included"
        bundle_path = main._output_artifact_path(report_payload["bundle_zip_url"])
        assert bundle_path is not None
        with zipfile.ZipFile(bundle_path) as archive:
            manifest = json.loads(archive.read("manifest.json"))
            assert manifest["historical_ndvi_status"] == "included"
            assert manifest["parcel_growth_status"] == "included"
            assert manifest["files"]["historical_ndvi_year_panel_2025"] in archive.namelist()
            assert manifest["files"]["parcel_growth_result"] in archive.namelist()
            assert manifest["parcel_growth_provenance"]["confirmation_receipts"][0]["receipt_sha256"] == receipt["receipt_sha256"]
            exported_history = json.loads(archive.read(manifest["files"]["historical_ndvi_result"]))
            assert exported_history["outputs"]["report_docx"] == manifest["files"]["historical_ndvi_report"]
            assert exported_history["outputs"]["trend_chart_png"] == manifest["files"]["historical_ndvi_trend"]
            assert exported_history["outputs"]["yearly_panel_pngs"]["2025"] == manifest["files"]["historical_ndvi_year_panel_2025"]
        approved = api.post(
            f"/api/v1/cases/{claim_id}/human_review",
            json={
                "decision": "approved",
                "generation_id": report_payload["generation_id"],
                "comment": "已核对历史与分地块附件",
            },
        )
        assert approved.status_code == 200, approved.text


def test_historical_remote_failure_never_persists_synthetic(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-HFAIL-{uuid.uuid4().hex[:8]}"
    policy_id = f"POL-HFAIL-{uuid.uuid4().hex[:8]}"
    _seed_case(claim_id, policy_id, state="SCREENING_DONE")

    def fail_history(**_kwargs):
        raise RuntimeError("GEE unavailable")

    monkeypatch.setattr(main, "_run_historical_ndvi_job", fail_history)
    with TestClient(main.app) as api:
        future = api.post(
            "/api/v1/tools/run_historical_ndvi_by_claim",
            json={
                "claim_id": claim_id,
                "start_year": 2025,
                "end_year": 2027,
                "as_of_date": "2027-12-31",
            },
        )
        response = api.post(
            "/api/v1/tools/run_historical_ndvi_by_claim",
            json={"claim_id": claim_id, "start_year": 2025, "end_year": 2025},
        )
    assert future.status_code == 422
    assert response.status_code == 502
    assert "未生成模拟结果" in response.json()["detail"]
    assert main._get_result(claim_id, "historical_ndvi") is None


def test_history_publish_is_rejected_if_report_freezes_during_remote_job(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-HRACE-{uuid.uuid4().hex[:8]}"
    policy_id = f"POL-HRACE-{uuid.uuid4().hex[:8]}"
    _seed_case(claim_id, policy_id, state="SCREENING_DONE")

    def freeze_then_finish(**kwargs):
        con = main._db()
        con.execute("UPDATE cases SET state = 'REPORT_DRAFTED' WHERE claim_id = ?", (claim_id,))
        con.commit()
        con.close()
        return _fake_history_job(**kwargs)

    monkeypatch.setattr(main, "_run_historical_ndvi_job", freeze_then_finish)
    with TestClient(main.app) as api:
        response = api.post(
            "/api/v1/tools/run_historical_ndvi_by_claim",
            json={
                "claim_id": claim_id,
                "start_year": 2025,
                "end_year": 2025,
                "as_of_date": "2025-12-31",
            },
        )
    assert response.status_code == 409
    assert "不得再发布" in response.json()["detail"]
    assert main._get_result(claim_id, "historical_ndvi") is None


def test_report_publish_rechecks_all_snapshot_revisions(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(main, "OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr(main, "AUTH_ENABLED", False)
    claim_id = f"CLAIM-RRACE-{uuid.uuid4().hex[:8]}"
    policy_id = f"POL-RRACE-{uuid.uuid4().hex[:8]}"
    boundary_sha256 = _seed_case(claim_id, policy_id, state="RULE_DONE")
    growth, _ = _seed_growth_result(
        tmp_path, boundary_sha256, task_id="growth-report-race"
    )
    _save_report_prerequisites(claim_id, boundary_sha256, growth)

    def change_revision_during_build(_claim_id, _server_data, template_version, sections):
        main._save_result(claim_id, "historical_ndvi", {"changed_during_report_build": True})
        return {
            "status": "success",
            "claim_id": claim_id,
            "generation_id": "race-generation",
            "template_version": template_version,
            "artifacts": [],
            "sections": sections,
        }

    monkeypatch.setattr(main, "_build_report_generation", change_revision_during_build)
    with TestClient(main.app) as api:
        response = api.post("/api/v1/tools/generate_report", json={"claim_id": claim_id})
    assert response.status_code == 409
    assert "权威结果发生变化" in response.json()["detail"]
    con = main._db()
    state = con.execute("SELECT state FROM cases WHERE claim_id = ?", (claim_id,)).fetchone()["state"]
    con.close()
    assert state == "RULE_DONE"
