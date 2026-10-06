"""
Agrisky AI — 报告生成回归测试
"""

import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent / "api_gateway"))

from report_generator import generate_claim_report


def _doc_text(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_report_uses_growth_payload_for_ndvi_section(tmp_path):
    output = tmp_path / "claim-report.docx"
    generate_claim_report(
        "CLAIM-TEST-001",
        {
            "case": {
                "policy_id": "POL-001",
                "disaster_type": "flood",
                "loss_date": "2026-07-18",
                "crop_type": "rice",
                "plot_id": "PLOT-001",
            },
            "satellite": {
                "suspected_damage_area_mu": 25,
                "damage_ratio": 0.25,
                "confidence": "mock",
                "image_count": 0,
            },
            "growth": {
                "summary": [
                    {"value": 1, "label": "差", "area_mu": 12.3, "ratio": 0.123, "count": 100, "color": "#D64545"},
                    {"value": 5, "label": "优", "area_mu": 45.6, "ratio": 0.456, "count": 300, "color": "#2F9E44"},
                ],
                "outputs": {},
            },
            "compliance": {
                "insured_area_mu": 100,
                "valid_damage_area_mu": 12.3,
                "excluded_area_mu": 0,
                "damage_ratio": 0.123,
                "clip_log": ["测试合规日志"],
            },
            "rule": {
                "risk_level": "medium",
                "review_required": True,
                "rule_version": "test_v1",
                "rule_trace": ["测试规则轨迹"],
            },
        },
        str(output),
    )

    text = _doc_text(output)
    assert "NDVI 作物长势分析" in text
    assert "差" in text
    assert "12.3" in text
    assert "测试合规日志" in text
    assert "结论与合规依据映射" in text
    assert "L-01§12" in text
    assert "L-02§21" in text
    assert "L-03§18" in text
    assert "L-06§15" in text
