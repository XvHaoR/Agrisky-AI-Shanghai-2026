"""
Agrisky AI — 定损报告生成器。

基于服务端案件快照生成可归档的 DOCX 报告。报告中的缺失值不会被解释为 0，
模拟遥感数据会被明确标注，所有 NDVI 结论均作为长势筛查结果而非直接灾损结论。
"""

from __future__ import annotations

import math
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

from compliance_library import reference_index, report_basis_rows


REPORT_TIMEZONE = timezone(timedelta(hours=8))
BODY_FONT_CN = "宋体"
HEADING_FONT_CN = "黑体"
BODY_FONT_LATIN = "Arial"
REPORT_GREEN = "1B4332"
LIGHT_GREEN = "E8F2EC"
MUTED_GREEN = "64736A"
WARNING_RED = "B94B46"

# 5 级长势配色
LEVEL_COLORS = {
    "差": "D64545",
    "一般": "EF8F35",
    "中": "E4C441",
    "良": "8CC152",
    "优": "2F9E44",
    1: "D64545",
    2: "EF8F35",
    3: "E4C441",
    4: "8CC152",
    5: "2F9E44",
}


def generate_claim_report(
    claim_id: str,
    data: dict,
    output_path: str,
    *,
    template_version: str = "v1.1",
    generated_at: datetime | None = None,
) -> str:
    """生成标准理赔报告 DOCX，含数据来源、NDVI 长势图表和人工审核区。"""
    generated_at = generated_at or datetime.now(REPORT_TIMEZONE)
    case = data.get("case", {}) or {}
    sat = data.get("satellite", {}) or {}
    growth = data.get("ndvi") or data.get("growth", {}) or {}
    compliance = data.get("compliance", {}) or {}
    rule = data.get("rule", {}) or {}
    loss = data.get("loss_assessment", {}) or {}
    payout = data.get("payout", {}) or {}
    contract = data.get("contract", {}) or {}
    materials = data.get("materials", {}) or {}
    report_meta = data.get("_report_meta", {}) or {}

    doc = Document()
    _configure_document(
        doc,
        claim_id,
        template_version,
        generated_at,
        report_title="农业保险查勘定损报告",
        generation_id=report_meta.get("generation_id"),
        snapshot_sha256=report_meta.get("snapshot_sha256"),
    )
    _add_cover(doc, claim_id, case, template_version, generated_at, report_meta)
    doc.add_page_break()

    _add_heading(doc, "一、案件基础信息")
    base_rows: list[tuple[str, Any]] = [
        ("案件编号", claim_id),
        ("保单号", case.get("policy_id")),
        ("投保人", case.get("holder_name")),
        ("承保地点", case.get("policy_address")),
        ("灾害类型", _disaster_name(case.get("disaster_type"))),
        ("灾害日期", case.get("loss_date")),
        ("作物类型", _crop_name(case.get("crop_type"))),
        ("地块编号", case.get("plot_id")),
    ]
    if _to_float(case.get("policy_area_mu")) is not None:
        base_rows.append(("保单登记面积（亩）", _format_number(case.get("policy_area_mu"), 2)))
    if case.get("policy_boundary_sha256"):
        base_rows.append(("在册承保边界 SHA-256", case.get("policy_boundary_sha256")))
    _add_kv_table(doc, base_rows)

    _add_contract_section(doc, contract)
    _add_material_section(doc, materials)
    _add_satellite_section(doc, sat)
    _add_growth_section(doc, growth, "五、NDVI 作物长势分析")
    _add_loss_section(doc, loss)
    _add_compliance_section(doc, compliance, contract)
    _add_rule_section(doc, rule)
    _add_payout_section(doc, payout, contract)
    _add_risk_section(doc, sat, growth, rule)
    _add_basis_section(doc, contract, str(case.get("crop_type") or ""))
    _add_reference_section(doc, contract, str(case.get("crop_type") or ""), "十二、依据与证据索引")
    _add_review_section(doc)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return str(out)


def generate_growth_appendix_report(
    claim_id: str,
    data: dict,
    output_path: str,
    *,
    template_version: str = "v1.1",
    generated_at: datetime | None = None,
) -> str:
    """从持久化长势摘要重建独立附件，供旧任务源 DOCX 丢失时安全降级。"""
    generated_at = generated_at or datetime.now(REPORT_TIMEZONE)
    case = data.get("case", {}) or {}
    growth = data.get("growth", {}) or {}
    report_meta = data.get("_report_meta", {}) or {}

    doc = Document()
    _configure_document(
        doc,
        claim_id,
        template_version,
        generated_at,
        report_title="作物长势监测报告",
        generation_id=report_meta.get("generation_id"),
        snapshot_sha256=report_meta.get("snapshot_sha256"),
    )
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(42)
    run = title.add_run("作物长势监测报告")
    _set_run_font(run, HEADING_FONT_CN, 22, bold=True, color=REPORT_GREEN)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.add_run("NDVI 遥感筛查附件"), BODY_FONT_CN, 12, color=MUTED_GREEN)

    _add_kv_table(
        doc,
        [
            ("案件编号", claim_id),
            ("保单号", case.get("policy_id")),
            ("投保人", case.get("holder_name")),
            ("作物", _crop_name(case.get("crop_type"))),
            ("模板版本", template_version),
            ("报告生成号", report_meta.get("generation_id")),
            ("数据快照 SHA-256", report_meta.get("snapshot_sha256")),
            ("生成时间", generated_at.strftime("%Y年%m月%d日 %H:%M（UTC+8）")),
        ],
    )
    _add_notice_box(doc, "本附件依据服务端已持久化的长势计算结果生成；NDVI 长势等级不直接等同于灾损等级。")
    _add_growth_section(doc, growth, "一、NDVI 作物长势分析")
    _add_heading(doc, "二、使用说明")
    doc.add_paragraph(
        "本报告用于遥感长势筛查和查勘辅助。赔付结论须结合承保条款、物候同期基线、"
        "多源灾损评估及人工现场复核形成。"
    )
    _add_reference_section(
        doc,
        data.get("contract", {}) or {},
        str(case.get("crop_type") or ""),
        "三、法规、合同与技术依据",
    )

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return str(out)


def _add_cover(
    doc: Document,
    claim_id: str,
    case: dict,
    template_version: str,
    generated_at: datetime,
    report_meta: dict,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(82)
    _set_run_font(p.add_run("农业保险灾后查勘定损报告"), HEADING_FONT_CN, 24, bold=True, color=REPORT_GREEN)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    _set_run_font(subtitle.add_run("遥感辅助评估 · 人工审核稿"), BODY_FONT_CN, 12, color=MUTED_GREEN)

    status_table = doc.add_table(rows=1, cols=1)
    status_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    status_cell = status_table.cell(0, 0)
    status_cell.text = "报告状态：草稿（须人工审核后生效）"
    _set_cell_shading(status_cell, "FFF1D6")
    for paragraph in status_cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            _set_run_font(run, BODY_FONT_CN, 10.5, bold=True, color="8A5A00")
    _set_cell_margins(status_cell, top=100, bottom=100, start=180, end=180)

    doc.add_paragraph()
    _add_kv_table(
        doc,
        [
            ("案件编号", claim_id),
            ("保单号", case.get("policy_id")),
            ("投保人", case.get("holder_name")),
            ("灾害类型", _disaster_name(case.get("disaster_type"))),
            ("灾害日期", case.get("loss_date")),
            ("模板版本", template_version),
            ("报告生成号", report_meta.get("generation_id")),
            ("数据快照 SHA-256", report_meta.get("snapshot_sha256")),
            ("生成时间", generated_at.strftime("%Y年%m月%d日 %H:%M（UTC+8）")),
        ],
    )
    _add_notice_box(
        doc,
        "本报告由服务端权威案件快照生成，用于辅助查勘定损。遥感筛查、模型估算和规则评级均须经人工复核，"
        "不得单独作为最终赔付决定。",
    )


def _add_satellite_section(doc: Document, sat: dict) -> None:
    _add_heading(doc, "四、卫星 SAR 遥感初筛")
    if not sat:
        _add_notice_box(doc, "未生成卫星遥感初筛结果，不能将缺失数据解释为 0 亩或低风险。", warning=True)
        return

    mock = _satellite_is_mock(sat)
    source_label = sat.get("source_label") or ("模拟 Sentinel-1 SAR" if mock else "GEE Sentinel-1 SAR")
    # Sentinel-1 引擎使用 GEE filterDate，截止日为排他边界；报告必须与计算语义一致。
    window = _date_window(sat.get("start_date"), sat.get("end_date"), inclusive_end=False)
    _add_kv_table(
        doc,
        [
            ("数据来源", source_label),
            ("观测窗口", window),
            ("疑似受灾面积（亩）", _format_number(sat.get("suspected_damage_area_mu"), 2)),
            ("初筛受损比例", _format_percent(sat.get("damage_ratio"))),
            ("置信度", _conf_label(sat.get("confidence"))),
            ("有效 Sentinel-1 影像", _format_count(sat.get("image_count"), "景")),
        ],
    )
    if mock:
        _add_notice_box(doc, "本节使用模拟卫星结果，仅供流程演示，不可作为定损证据。", warning=True)
    else:
        doc.add_paragraph(
            "系统依据承保地块范围开展 Sentinel-1 SAR 洪涝初筛。疑似受灾面积用于确定查勘重点，"
            "不等同于经承保红线裁剪后的合规受灾面积。[L-01§12][L-02§20][L-02§21]"
        )


def _add_material_section(doc: Document, materials: dict) -> None:
    _add_heading(doc, "三、材料理解与一致性核验")
    documents = materials.get("documents") or []
    fields = materials.get("fields") or []
    findings = materials.get("findings") or []
    open_findings = [item for item in findings if item.get("status") == "open"]
    high_findings = [item for item in open_findings if item.get("severity") == "high"]
    _add_kv_table(
        doc,
        [
            ("已登记材料", f"{len(documents)} 份"),
            ("结构化字段", f"{len(fields)} 项"),
            ("开放问题", f"{len(open_findings)} 项"),
            ("高风险阻断", f"{len(high_findings)} 项"),
        ],
    )
    if not documents:
        _add_notice_box(
            doc,
            "案件没有登记可追溯材料；报告仅保留历史流程结果，不能据此认定材料完整。",
            warning=True,
        )
        return

    table = doc.add_table(rows=1, cols=4, style="Table Grid")
    headers = ["材料类型", "文件名", "解析状态", "SHA-256"]
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True, size=9)
        _set_cell_shading(table.rows[0].cells[index], LIGHT_GREEN)
    for item in documents:
        row = table.add_row()
        values = [
            item.get("document_type"),
            item.get("original_filename"),
            item.get("parse_status"),
            str(item.get("sha256") or "")[:16],
        ]
        for index, value in enumerate(values):
            _set_cell_text(row.cells[index], value or "—", size=8.5)
        _prevent_row_split(row)

    if fields:
        doc.add_paragraph("关键字段及证据引用：")
        field_table = doc.add_table(rows=1, cols=4, style="Table Grid")
        field_headers = ["字段", "提取值", "置信度", "证据引用"]
        for index, header in enumerate(field_headers):
            _set_cell_text(field_table.rows[0].cells[index], header, bold=True, size=9)
            _set_cell_shading(field_table.rows[0].cells[index], LIGHT_GREEN)
        for item in fields:
            row = field_table.add_row()
            confidence = _to_float(item.get("confidence"))
            values = [
                item.get("field_name"),
                item.get("normalized_value"),
                f"{confidence * 100:.1f}%" if confidence is not None else "—",
                f"第 {item.get('page_number') or '—'} 页 · {item.get('source_ref') or '—'}",
            ]
            for index, value in enumerate(values):
                _set_cell_text(row.cells[index], value, size=8.5)
            _prevent_row_split(row)

    if open_findings:
        messages = "；".join(
            f"[{item.get('severity')}] {item.get('message')} ({item.get('source_ref') or '系统门禁'})"
            for item in open_findings[:8]
        )
        _add_notice_box(doc, f"材料核验待处理：{messages}", warning=True)
    else:
        _add_notice_box(doc, "材料字段已完成一致性核验，未发现开放的高风险冲突。")
    doc.add_paragraph(
        "材料、原始记录与证据摘要应保持真实、完整并可追溯。[L-01§22][L-02§23][L-02§46][L-03§22]"
    )


def _add_growth_section(doc: Document, growth: dict, heading: str) -> None:
    _add_heading(doc, heading)
    summary = growth.get("summary", []) or []
    if not summary:
        _add_notice_box(doc, "未执行长势分析或没有可用的 NDVI 长势结果。", warning=True)
        return

    raster = growth.get("raster", {}) or {}
    meta = raster.get("ndvi_meta", {}) or {}
    source_mode = raster.get("ndvi_source") or meta.get("source")
    source_label = raster.get("ndvi_source_label") or meta.get("source_label") or "服务端 NDVI 结果"
    method = growth.get("method") or "-"
    class_breaks = growth.get("class_breaks") or []

    formula = meta.get("formula")
    if not formula:
        if source_mode == "gee":
            formula = "NDVI = (Sentinel-2 B8 - B4) / (B8 + B4)"
        elif source_mode == "upload":
            formula = "源文件已提供 NDVI 值；原始波段公式未随文件提供"
        elif source_mode == "synthetic":
            formula = "本地模拟数值（非卫星波段反演）"

    provenance_rows: list[tuple[str, Any]] = [
        ("数据来源", source_label),
        ("观测窗口", _date_window(meta.get("start_date"), meta.get("end_date"))),
        ("合成与分级", _method_label(method)),
        ("NDVI 公式/口径", formula),
    ]
    if meta.get("collection"):
        provenance_rows.append(("影像集合", meta.get("collection")))
    if meta.get("source_filename"):
        provenance_rows.append(("源文件", meta.get("source_filename")))
    if meta.get("source_sha256"):
        provenance_rows.append(("源文件 SHA-256", meta.get("source_sha256")))
    if meta.get("boundary_source") or meta.get("boundary_filename"):
        provenance_rows.append(("地块边界来源", meta.get("boundary_source") or meta.get("boundary_filename")))
    if meta.get("boundary_sha256"):
        provenance_rows.append(("地块边界 SHA-256", meta.get("boundary_sha256")))
    if _to_float(meta.get("max_cloud_pct")) is not None:
        provenance_rows.append(("最大云量阈值", f"{_to_float(meta.get('max_cloud_pct')):g}%"))
    if meta.get("cloud_mask"):
        provenance_rows.append(("像元质量掩膜", meta.get("cloud_mask")))
    if _to_float(meta.get("scale")) is not None:
        provenance_rows.append(("请求空间分辨率", f"{_to_float(meta.get('scale')):g} m"))
    provenance_rows.extend(
        [
            ("有效影像数", _format_count(meta.get("image_count"), "景")),
            ("栅格坐标系", raster.get("crs")),
            ("NoData 值", raster.get("nodata")),
            (
                "面积折算口径",
                growth.get("area_estimation_method")
                or "各等级有效像元占比 × 地块总面积（对无效像元按有效像元分布作比例外推）",
            ),
            ("监测面积（亩）", _format_number(growth.get("total_area_mu"), 2)),
            ("有效像元数", _format_integer(growth.get("valid_pixel_count"))),
            ("地块内候选像元数", _format_integer(raster.get("roi_pixel_count"))),
            ("有效像元覆盖率", _format_percent(raster.get("valid_pixel_coverage"))),
        ]
    )

    _add_kv_table(doc, provenance_rows)
    doc.add_paragraph(
        f"本期采用{_method_label(method)}对 NDVI 进行五级分类。各等级面积由有效像元占比乘以地块总面积折算；"
        "存在云、NoData 等无效像元时，该口径相当于按有效像元分布作比例外推，并非逐像元实测面积。"
        "[L-02§21][L-02§26]"
    )
    if source_mode == "synthetic":
        _add_notice_box(doc, "本节使用模拟 NDVI，仅供流程演示，不可作为定损证据。", warning=True)

    table = doc.add_table(rows=1, cols=7, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["等级", "NDVI 范围", "长势评价", "面积（亩）", "占比", "像元数", "图例"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_text(cell, header, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(cell, LIGHT_GREEN)
    _mark_repeat_header(table.rows[0])
    _prevent_row_split(table.rows[0])

    qa_warnings: list[str] = []
    n_classes_number = _to_float(growth.get("n_classes"))
    expected_classes = int(n_classes_number) if n_classes_number is not None else len(summary)
    break_numbers = [_to_float(value) for value in class_breaks]
    if not class_breaks or (expected_classes > 0 and len(class_breaks) != expected_classes):
        qa_warnings.append("NDVI 分级阈值数量与等级数不一致，区间仅供核查。")
    if any(value is None for value in break_numbers):
        qa_warnings.append("NDVI 分级阈值包含非数值项。")
    valid_breaks = [value for value in break_numbers if value is not None]
    if any(value < -1 or value > 1 for value in valid_breaks):
        qa_warnings.append("NDVI 分级阈值超出 [-1, 1] 物理范围。")
    if len(valid_breaks) == len(break_numbers) and any(
        left >= right for left, right in zip(valid_breaks, valid_breaks[1:])
    ):
        qa_warnings.append("NDVI 分级阈值未严格递增。")

    total_area = 0.0
    total_ratio = 0.0
    total_count = 0
    all_areas_valid = True
    all_ratios_valid = True
    all_counts_valid = True
    level_counts: dict[int, int] = {}
    invalid_level_entries = 0
    for item in sorted(summary, key=lambda entry: _to_float(entry.get("value")) or 0, reverse=True):
        row = table.add_row()
        raw_level = _to_float(item.get("value"))
        value = int(raw_level) if raw_level is not None and raw_level.is_integer() and raw_level > 0 else None
        if value is not None:
            level_counts[value] = level_counts.get(value, 0) + 1
        else:
            invalid_level_entries += 1
        area = _to_float(item.get("area_mu"))
        ratio = _to_float(item.get("ratio"))
        raw_count = _to_float(item.get("count"))
        area_valid = area is not None and area >= 0
        ratio_valid = ratio is not None and 0 <= ratio <= 1
        count_valid = raw_count is not None and raw_count >= 0 and raw_count.is_integer()
        all_areas_valid = all_areas_valid and area_valid
        all_ratios_valid = all_ratios_valid and ratio_valid
        all_counts_valid = all_counts_valid and count_valid
        if area_valid:
            total_area += area
        if ratio_valid:
            total_ratio += ratio
        if count_valid:
            total_count += int(raw_count)
        values = [
            value if value is not None else "—",
            _ndvi_interval(value, class_breaks) if value is not None else "—",
            item.get("label"),
            f"{area:,.2f}" if area_valid else "—",
            f"{ratio * 100:.2f}%" if ratio_valid else "—",
            f"{int(raw_count):,}" if count_valid else "—",
            "",
        ]
        for index, value_text in enumerate(values):
            _set_cell_text(row.cells[index], value_text, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        color = LEVEL_COLORS.get(item.get("label"), LEVEL_COLORS.get(value, "CCCCCC"))
        _set_cell_shading(row.cells[6], color)
        _prevent_row_split(row)

    total_row = table.add_row()
    merged = total_row.cells[0].merge(total_row.cells[2])
    _set_cell_text(merged, "合计", bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_text(
        total_row.cells[3],
        f"{total_area:,.2f}" if all_areas_valid else "—",
        bold=True,
        size=9.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _set_cell_text(
        total_row.cells[4],
        f"{total_ratio * 100:.2f}%" if all_ratios_valid else "—",
        bold=True,
        size=9.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _set_cell_text(
        total_row.cells[5],
        f"{total_count:,}" if all_counts_valid else "—",
        bold=True,
        size=9.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _set_cell_text(total_row.cells[6], "—", bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in total_row.cells:
        _set_cell_shading(cell, "F3F6F4")
    _prevent_row_split(total_row)
    doc.add_paragraph()

    if not all_areas_valid:
        qa_warnings.append("部分长势等级的面积缺失、非数值或为负数，合计面积未计算。")
    if not all_ratios_valid:
        qa_warnings.append("部分长势等级的占比缺失或超出 0%–100%，合计占比未计算。")
    if not all_counts_valid:
        qa_warnings.append("部分长势等级的像元数缺失、非整数或为负数，合计像元数未计算。")
    if all_ratios_valid and abs(total_ratio - 1.0) > 0.02:
        qa_warnings.append(f"各等级占比合计为 {total_ratio * 100:.2f}%，与 100% 偏差超过 2 个百分点。")
    declared_area = _to_float(growth.get("total_area_mu"))
    if all_areas_valid and declared_area is not None:
        tolerance = max(0.1, abs(declared_area) * 0.01)
        if abs(total_area - declared_area) > tolerance:
            qa_warnings.append(
                f"各等级面积合计 {total_area:,.2f} 亩，与监测面积 {declared_area:,.2f} 亩不一致。"
            )
    declared_count = _to_float(growth.get("valid_pixel_count"))
    if all_counts_valid and declared_count is not None and (
        not declared_count.is_integer() or total_count != int(declared_count)
    ):
        qa_warnings.append(
            f"各等级像元数合计 {total_count:,}，与顶层有效像元数 {_format_integer(declared_count)} 不一致。"
        )
    valid_coverage = _to_float(raster.get("valid_pixel_coverage"))
    if valid_coverage is not None:
        if not 0 <= valid_coverage <= 1:
            qa_warnings.append("有效像元覆盖率超出 0%–100%，覆盖率元数据无效。")
        elif valid_coverage < 0.7:
            qa_warnings.append(
                f"有效像元覆盖率仅 {valid_coverage * 100:.2f}%，低于 70%；长势分级可能受云、阴影或 NoData 影响。"
            )
    expected_levels = set(range(1, expected_classes + 1)) if expected_classes > 0 else set()
    observed_levels = set(level_counts)
    duplicate_levels = sorted(level for level, count in level_counts.items() if count > 1)
    missing_levels = sorted(expected_levels - observed_levels)
    out_of_range_levels = sorted(observed_levels - expected_levels) if expected_levels else sorted(observed_levels)
    if duplicate_levels:
        qa_warnings.append(f"长势等级存在重复记录：{', '.join(map(str, duplicate_levels))} 级。")
    if missing_levels:
        qa_warnings.append(f"长势等级记录缺失：{', '.join(map(str, missing_levels))} 级。")
    if invalid_level_entries or out_of_range_levels:
        details = f"（{', '.join(map(str, out_of_range_levels))} 级）" if out_of_range_levels else ""
        qa_warnings.append(f"长势等级存在非整数、非正数或超出范围的记录{details}。")

    for warning in dict.fromkeys(qa_warnings):
        _add_notice_box(doc, warning, warning=True)

    poor_items = [
        item for item in summary
        if item.get("label") in ("差", "一般") or (_to_float(item.get("value")) or 99) <= 2
    ]
    summary_complete = bool(expected_levels) and all(level_counts.get(level) == 1 for level in expected_levels) and not (
        invalid_level_entries or out_of_range_levels or duplicate_levels
    )
    poor_values_valid = all(
        (_to_float(item.get("area_mu")) is not None and _to_float(item.get("area_mu")) >= 0)
        and (_to_float(item.get("ratio")) is not None and 0 <= _to_float(item.get("ratio")) <= 1)
        for item in poor_items
    )
    if summary_complete and poor_values_valid:
        poor_ratio = sum(_to_float(item.get("ratio")) for item in poor_items)
        poor_area = sum(_to_float(item.get("area_mu")) for item in poor_items)
        doc.add_paragraph(
            f"根据本期 NDVI 分类结果，长势“差/一般”区域合计 {poor_area:,.2f} 亩"
            f"（占比 {poor_ratio * 100:.2f}%）。该结果仅用于长势筛查，不直接等同于灾损；"
            "建议结合物候同期基线、多源灾损评估与现场查勘复核。"
        )
    else:
        _add_notice_box(doc, "因长势等级记录或数值不完整，无法可靠汇总“差/一般”区域。", warning=True)

    growth_outputs = growth.get("outputs", {}) or {}
    # 正式报告优先使用本次计算生成的本地分级图，避免把定位底图误认为观测期 Sentinel-2 证据。
    preview_path = growth_outputs.get("class_preview_png") or growth_outputs.get("report_growth_map_png")
    local_path = _resolve_image_path(preview_path) if preview_path else None
    if local_path:
        label = doc.add_paragraph()
        label.paragraph_format.keep_with_next = True
        label.paragraph_format.keep_together = True
        _set_run_font(label.add_run("NDVI 长势分级预览图"), BODY_FONT_CN, 10.5, bold=True)
        try:
            doc.add_picture(str(local_path), width=Inches(5.65))
            picture_paragraph = doc.paragraphs[-1]
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.paragraph_format.keep_with_next = True
            picture_paragraph.paragraph_format.keep_together = True
            shape = doc.inline_shapes[-1]
            shape._inline.docPr.set("title", "NDVI 长势分级预览图")
            shape._inline.docPr.set("descr", "承保地块 NDVI 五级长势分类图，颜色由红到绿表示差到优")
            caption = doc.add_paragraph("图 1  NDVI 作物长势分级图（红→绿：差→优）")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_together = True
            for run in caption.runs:
                _set_run_font(run, BODY_FONT_CN, 9, color=MUTED_GREEN)
        except Exception:
            _add_notice_box(doc, "长势预览图加载失败，统计结果仍保留。", warning=True)
    else:
        _add_notice_box(doc, "本次快照未包含可用的本地 NDVI 分级预览图，统计结果仍保留。", warning=True)


def _add_contract_section(doc: Document, contract: dict) -> None:
    _add_heading(doc, "二、保险合同与规则依据")
    if not contract:
        _add_notice_box(doc, "未找到冻结的保险合同版本，报告不能形成赔款结论。", warning=True)
        return
    _add_notice_box(
        doc,
        str(contract.get("simulation_disclosure") or "本合同为产品规则验证用模拟合同。"),
    )
    _add_kv_table(
        doc,
        [
            ("合同编号", contract.get("contract_number")),
            ("合同版本", contract.get("contract_version")),
            ("合同哈希（SHA-256）", contract.get("contract_sha256")),
            ("保险人", contract.get("insurer")),
            ("投保人/被保险人", contract.get("policyholder")),
            ("保险期间", " 至 ".join((contract.get("insurance_period") or {}).values())),
        ],
    )
    doc.add_paragraph(
        "合同要素、保险责任和赔偿办法以登记时上传并冻结的合同版本为准；报告生成阶段不得改写。"
        "[L-03§18][L-05§8][L-05§23]"
    )
    for clause in (contract.get("clauses") or [])[:4]:
        doc.add_paragraph(
            f"[C-{str(clause.get('id') or '').lstrip('C')}] {clause.get('title')}：{clause.get('text')}",
            style="List Bullet",
        )


def _add_loss_section(doc: Document, loss: dict) -> None:
    _add_heading(doc, "六、多源遥感灾损评估（减产率）")
    if not loss:
        _add_notice_box(doc, "未生成多源灾损评估。", warning=True)
        return

    severity = loss.get("severity") or {}
    _add_kv_table(
        doc,
        [
            ("综合减产率", _format_percent(loss.get("yield_loss_ratio"))),
            ("严重度", severity.get("label")),
            ("置信度", _conf_label(loss.get("confidence"))),
            ("主导驱动因子", "、".join(_safe_list(loss.get("dominant_drivers"))) or "—"),
            ("数据来源", "/".join(_safe_list(loss.get("data_sources"))) or "—"),
            ("评估方法", loss.get("method")),
        ],
    )
    doc.add_paragraph(
        "多源减产率用于综合判读灾害影响，仍需与保险条款约定、作物生育期、气象与现场证据共同复核。"
        "[L-01§12][L-02§26][L-02§31][T-01]"
    )

    breakdown = loss.get("product_breakdown") or []
    if breakdown:
        table = doc.add_table(rows=1, cols=6, style="Table Grid")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["遥感产品", "分辨率", "损失分量", "权重", "贡献", "来源"]
        for index, header in enumerate(headers):
            _set_cell_text(table.rows[0].cells[index], header, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_shading(table.rows[0].cells[index], LIGHT_GREEN)
        _mark_repeat_header(table.rows[0])
        _prevent_row_split(table.rows[0])
        for item in breakdown:
            row = table.add_row()
            resolution = item.get("native_res_m")
            values = [
                item.get("name_cn"),
                f"{_display(resolution)} m" if resolution not in (None, "") else "—",
                _format_percent(item.get("decline_score"), 1),
                _format_percent(item.get("weight"), 1),
                _format_percent(item.get("contribution"), 1),
                item.get("data_source"),
            ]
            for index, value in enumerate(values):
                _set_cell_text(row.cells[index], value, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            _prevent_row_split(row)
        doc.add_paragraph()

    for caveat in _safe_list(loss.get("caveats")):
        paragraph = doc.add_paragraph(style="Normal")
        _set_run_font(paragraph.add_run(f"注：{caveat}"), BODY_FONT_CN, 9, color=MUTED_GREEN)


def _add_compliance_section(doc: Document, compliance: dict, contract: dict | None = None) -> None:
    _add_heading(doc, "七、合规面积核验")
    if not compliance:
        _add_notice_box(doc, "未生成合规面积核验结果。", warning=True)
        return
    _add_kv_table(
        doc,
        [
            ("承保面积（亩）", _format_number(compliance.get("insured_area_mu"), 2)),
            ("合规受灾面积（亩）", _format_number(compliance.get("valid_damage_area_mu"), 2)),
            ("剔除/未计入面积（亩）", _format_number(compliance.get("excluded_area_mu"), 2)),
            ("合规受损比例", _format_percent(compliance.get("damage_ratio"))),
        ],
    )
    for log in _safe_list(compliance.get("clip_log")):
        doc.add_paragraph(log, style="List Bullet")
    doc.add_paragraph(
        "合规核验遵循合同约定与农业保险承保理赔的真实、完整、可追溯要求。"
        "[L-01§12][L-02§26][L-02§31][L-05§23]"
    )
    for check in compliance.get("contract_checks") or []:
        status = "通过" if check.get("passed") else "未通过"
        doc.add_paragraph(
            f"{check.get('clause')} {check.get('name')}：{status}。{check.get('detail')}",
            style="List Bullet",
        )


def _add_rule_section(doc: Document, rule: dict) -> None:
    _add_heading(doc, "八、规则引擎建议")
    if not rule:
        _add_notice_box(doc, "未生成规则评级；缺失规则结果不能解释为低风险。", warning=True)
        return
    risk_level = rule.get("risk_level")
    _add_kv_table(
        doc,
        [
            ("建议风险等级", _risk_label(risk_level)),
            ("是否必须人工复核", "是" if rule.get("review_required") is True else "否" if rule.get("review_required") is False else "—"),
            ("规则版本", rule.get("rule_version")),
        ],
    )
    for trace in _safe_list(rule.get("rule_trace")):
        doc.add_paragraph(trace, style="List Bullet")


def _add_payout_section(doc: Document, payout: dict, contract: dict | None = None) -> None:
    _add_heading(doc, "九、赔付测算")
    if not payout:
        _add_notice_box(doc, "未生成赔付测算。", warning=True)
        return
    payout_table, payout_spacer = _add_kv_table(
        doc,
        [
            ("作物", _crop_name(payout.get("crop_type"))),
            ("保额（元/亩）", _format_currency(payout.get("sum_insured_per_mu"), 2)),
            ("承保面积（亩）", _format_number(payout.get("insured_area_mu"), 2)),
            ("核定受灾面积（亩）", _format_number(payout.get("affected_area_mu"), 2)),
            ("总保额（元）", _format_currency(payout.get("total_sum_insured_yuan"), 2)),
            ("减产率", _format_percent(payout.get("yield_loss_ratio"))),
            ("起赔点", _format_percent(payout.get("deductible_threshold"))),
            ("赔付档", payout.get("tier_label")),
            ("赔付比例", _format_percent(payout.get("payout_factor"))),
            ("生育期与系数", f"{payout.get('growth_stage') or '—'} / {_format_number(payout.get('growth_stage_factor'), 2)}"),
            ("合同条款引用", "、".join(_safe_list(payout.get("contract_clause_refs"))) or "—"),
            ("预估赔款（元）", _format_currency(payout.get("payout_amount_yuan"), 2)),
        ],
    )
    if payout.get("calculation_formula"):
        doc.add_paragraph(
            f"合同计算公式：{payout.get('calculation_formula')} "
            "[C-7][L-01§15][L-03§55][L-05§23]"
        )
    traces = _safe_list(payout.get("trace"))
    if traces:
        for cell in payout_table.rows[-1].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        payout_spacer.paragraph_format.keep_with_next = True
        payout_spacer.paragraph_format.space_after = Pt(0)
    for index, trace in enumerate(traces):
        paragraph = doc.add_paragraph(trace, style="List Bullet")
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = index < len(traces) - 1


def _add_risk_section(doc: Document, sat: dict, growth: dict, rule: dict) -> None:
    _add_heading(doc, "十、风险提示")
    warnings: list[str] = []
    risk_level = rule.get("risk_level") if rule else None
    if risk_level == "high":
        warnings.append("该案件被评为高风险，必须启动高级别人工复核，不得自动结案。")
    elif risk_level == "medium":
        warnings.append("该案件被评为中风险，建议审核员重点核验关键证据后再形成结论。")
    elif risk_level == "low":
        warnings.append("该案件当前规则评级为低风险，仍须完成人工确认后方可归档。")
    else:
        warnings.append("规则评级缺失，当前报告不能形成风险结论。")
    if _satellite_is_mock(sat):
        warnings.append("卫星初筛包含模拟数据，不得作为定损证据。")
    raster = growth.get("raster", {}) or {}
    meta = raster.get("ndvi_meta", {}) or {}
    if (raster.get("ndvi_source") or meta.get("source")) == "synthetic":
        warnings.append("NDVI 长势分析包含模拟数据，不得作为定损证据。")
    for warning in warnings:
        _add_notice_box(doc, warning, warning=risk_level in {"high", "medium"} or "模拟" in warning or "缺失" in warning)


def _add_basis_section(doc: Document, contract: dict, crop_type: str) -> None:
    _add_heading(doc, "十一、结论与合规依据映射")
    table = doc.add_table(rows=1, cols=4, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ("业务环节", "本案结论口径", "引用依据", "适用边界")):
        _set_cell_text(cell, label, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(cell, LIGHT_GREEN)
    _mark_repeat_header(table.rows[0])
    for item in report_basis_rows(contract, crop_type):
        row = table.add_row()
        values = (
            item.get("stage"),
            item.get("conclusion"),
            " ".join(f"[{basis}]" for basis in item.get("basis_ids") or []),
            item.get("boundary"),
        )
        for index, value in enumerate(values):
            _set_cell_text(row.cells[index], value, size=8.5)
        _prevent_row_split(row)
    doc.add_paragraph(
        "以上映射由服务端固定规则生成。Agent 可检索并解释依据，但不能新增条款、改变合同参数或直接作出赔付决定。"
    )


def _add_reference_section(
    doc: Document, contract: dict, crop_type: str, heading: str
) -> None:
    _add_heading(doc, heading)
    table = doc.add_table(rows=1, cols=4, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ("编号", "类别", "名称", "来源/版本")):
        _set_cell_text(cell, label, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(cell, LIGHT_GREEN)
    _mark_repeat_header(table.rows[0])
    for item in reference_index(contract, crop_type):
        row = table.add_row()
        title = str(item.get("title") or "")
        if item.get("summary"):
            title = f"{title}\n要点：{item.get('summary')}"
        source = str(item.get("source") or item.get("issuer") or "")
        if item.get("applicability"):
            source = f"{item.get('applicability')}\n{source}"
        values = (
            item.get("id"),
            item.get("category"),
            title,
            source,
        )
        for index, value in enumerate(values):
            _set_cell_text(row.cells[index], value, size=8.5)
        _prevent_row_split(row)
    doc.add_paragraph(
        "引用编号说明：L=法律法规，C=已上传并冻结的保险合同及条款，T=遥感技术依据，D=服务端证据数据。"
    )


def _add_review_section(doc: Document) -> None:
    _add_heading(doc, "十三、人工审核意见区")
    doc.add_paragraph("审核结论：□ 通过    □ 退回补充材料    □ 驳回")
    doc.add_paragraph("审核意见：")
    for _ in range(3):
        paragraph = doc.add_paragraph("____________________________________________________________________")
        for run in paragraph.runs:
            _set_run_font(run, BODY_FONT_LATIN, 9, color="A0AAA3")
    doc.add_paragraph("审核人签名：____________________    审核日期：______年____月____日")
    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(final.add_run("模型负责理解 · 引擎负责计算 · 人类负责决策"), BODY_FONT_CN, 9, color=MUTED_GREEN)


def _configure_document(
    doc: Document,
    claim_id: str,
    template_version: str,
    generated_at: datetime,
    *,
    report_title: str,
    generation_id: str | None = None,
    snapshot_sha256: str | None = None,
) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    _set_style_font(normal, BODY_FONT_CN, 10.5, BODY_FONT_LATIN)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    if "Title" in doc.styles:
        _set_style_font(doc.styles["Title"], HEADING_FONT_CN, 24, BODY_FONT_LATIN)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        _set_style_font(style, HEADING_FONT_CN, size, BODY_FONT_LATIN)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(REPORT_GREEN)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        header.add_run(f"Agrisky AI｜{report_title}｜{claim_id}"),
        BODY_FONT_CN,
        8.5,
        color=MUTED_GREEN,
    )
    _set_paragraph_bottom_border(header, "D8E2DC")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_prefix = f"人工审核稿　·　模板 {template_version}"
    if snapshot_sha256:
        footer_prefix += f"　·　快照 {snapshot_sha256[:12]}"
    _set_run_font(footer.add_run(f"{footer_prefix}　·　第 "), BODY_FONT_CN, 8.5, color=MUTED_GREEN)
    _add_field(footer, "PAGE")
    _set_run_font(footer.add_run(" 页 / 共 "), BODY_FONT_CN, 8.5, color=MUTED_GREEN)
    _add_field(footer, "NUMPAGES")
    _set_run_font(footer.add_run(" 页"), BODY_FONT_CN, 8.5, color=MUTED_GREEN)

    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(first_footer.add_run("内部工作底稿 · 须经人工审核"), BODY_FONT_CN, 8.5, color=MUTED_GREEN)

    props = doc.core_properties
    props.title = f"{report_title} - {claim_id}"
    props.subject = "农业保险遥感辅助查勘定损" if "定损" in report_title else "农业保险 NDVI 作物长势监测"
    props.author = "Agrisky AI"
    props.keywords = "农业保险, 遥感, NDVI, 查勘, 定损"
    props.comments = f"模板版本 {template_version}；服务端案件快照生成"
    utc_time = (
        generated_at.astimezone(timezone.utc).replace(tzinfo=None)
        if generated_at.tzinfo is not None
        else generated_at.replace(tzinfo=None)
    )
    props.created = utc_time
    props.modified = utc_time

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_heading(text, level=1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def _add_kv_table(doc: Document, rows: list[tuple[str, Any]]) -> tuple[Any, Any]:
    table = doc.add_table(rows=len(rows), cols=2, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (label, value) in enumerate(rows):
        row = table.rows[index]
        label_cell, value_cell = row.cells
        label_cell.width = Cm(4.3)
        value_cell.width = Cm(11.6)
        _set_cell_text(label_cell, label, bold=True, size=10)
        _set_cell_shading(label_cell, "F0F4F0")
        _set_cell_text(value_cell, _display(value), size=10)
        _prevent_row_split(row)
    # 小型键值表作为一个语义块分页，避免“标签在上一页、数值在下一页”或只剩末行。
    if len(rows) <= 16:
        _keep_table_together(table)
    spacer = doc.add_paragraph()
    return table, spacer


def _add_notice_box(doc: Document, text: str, *, warning: bool = False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "FDECEC" if warning else "F3F6F4")
    _set_cell_text(cell, text, bold=warning, size=9.5, color=WARNING_RED if warning else MUTED_GREEN)
    _set_cell_margins(cell, top=110, bottom=110, start=150, end=150)
    _prevent_row_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _set_cell_text(
    cell: Any,
    value: Any,
    *,
    bold: bool = False,
    size: float = 10,
    color: str | None = None,
    align: Any = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = _display(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell, top=75, bottom=75, start=100, end=100)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            _set_run_font(run, BODY_FONT_CN, size, bold=bold, color=color)


def _set_run_font(
    run: Any,
    east_asia: str,
    size: float,
    *,
    bold: bool = False,
    color: str | None = None,
    latin: str = BODY_FONT_LATIN,
) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _set_style_font(style: Any, east_asia: str, size: float, latin: str) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _set_cell_shading(cell: Any, color: str) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    for existing in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    shading.set(qn("w:val"), "clear")
    tc_pr.append(shading)


def _set_cell_margins(cell: Any, *, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _prevent_row_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _mark_repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "true")
        tr_pr.append(marker)
    # 让表头至少与首条数据绑定，避免表头孤立在页尾。
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def _keep_table_together(table: Any) -> None:
    """通过段落 keep 链尽量让小表整体分页；超长表仍可由 Word 正常换页。"""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def _set_paragraph_bottom_border(paragraph: Any, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _add_field(paragraph: Any, field_name: str) -> None:
    run = paragraph.add_run()
    _set_run_font(run, BODY_FONT_LATIN, 8.5, color=MUTED_GREEN)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_name} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def _to_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _display(value: Any) -> str:
    if value is None:
        return "—"
    text = str(value).strip()
    return text if text else "—"


def _format_number(value: Any, decimals: int = 1) -> str:
    number = _to_float(value)
    return "—" if number is None else f"{number:,.{decimals}f}"


def _format_currency(value: Any, decimals: int = 2) -> str:
    return _format_number(value, decimals)


def _format_percent(value: Any, decimals: int = 1) -> str:
    number = _to_float(value)
    return "—" if number is None else f"{number * 100:.{decimals}f}%"


def _format_integer(value: Any) -> str:
    number = _to_float(value)
    return "—" if number is None else f"{int(number):,}"


def _format_count(value: Any, unit: str) -> str:
    number = _to_float(value)
    return "—" if number is None else f"{int(number):,} {unit}"


def _safe_list(value: Any) -> list[str]:
    if not isinstance(value, (list, tuple)):
        return []
    return [str(item) for item in value if item not in (None, "")]


def _date_window(start: Any, end: Any, *, inclusive_end: bool = True) -> str:
    start_text = _display(start)
    end_text = _display(end)
    if start_text == "—" and end_text == "—":
        return "—"
    semantics = "起止日期均含" if inclusive_end else "起始含、截止不含"
    return f"{start_text} 至 {end_text}（{semantics}）"


def _ndvi_interval(value: int, class_breaks: list[Any]) -> str:
    breaks = [number for item in class_breaks if (number := _to_float(item)) is not None]
    if value <= 0 or not breaks or value > len(breaks):
        return "—"
    upper = breaks[value - 1]
    if value == 1:
        return f"< {upper:.2f}"
    lower = breaks[value - 2]
    if value == len(breaks):
        return f"≥ {lower:.2f}"
    return f"[{lower:.2f}, {upper:.2f})"


def _method_label(method: Any) -> str:
    key = str(method or "").lower()
    return {
        "fixed": "固定阈值（0.30 / 0.45 / 0.60 / 0.75）",
        "jenks": "Jenks 自然断点",
        "equalinterval": "等距分级",
        "quantile": "分位数分级",
        "std": "标准差分级",
    }.get(key, _display(method))


def _satellite_is_mock(sat: dict) -> bool:
    if not sat:
        return False
    if sat.get("is_mock") is True or sat.get("source") in {"mock", "synthetic"}:
        return True
    if str(sat.get("confidence") or "").lower() == "mock":
        return True
    return any(str(asset).lower().startswith("mock_") for asset in sat.get("reference_assets", []) or [])


def _disaster_name(key: Any) -> str:
    return {
        "flood": "洪水",
        "drought": "干旱",
        "hail": "冰雹",
        "typhoon": "台风",
        "pest": "虫灾",
        "frost": "霜冻",
        "other": "其他",
    }.get(str(key or ""), _display(key))


def _crop_name(key: Any) -> str:
    return {
        "rice": "水稻",
        "wheat": "小麦",
        "corn": "玉米",
        "maize": "玉米",
        "soybean": "大豆",
        "cotton": "棉花",
        "peanut": "花生",
    }.get(str(key or "").lower(), _display(key))


def _conf_label(value: Any) -> str:
    return {
        "high": "高",
        "medium": "中",
        "low": "低",
        "mock": "模拟数据（无证据置信度）",
    }.get(str(value or "").lower(), _display(value))


def _risk_label(value: Any) -> str:
    return {"high": "高风险", "medium": "中风险", "low": "低风险"}.get(
        str(value or "").lower(), _display(value)
    )


def _resolve_image_path(url: str) -> Path | None:
    """将 /outputs URL 转为受约束的本地文件路径。"""
    if not url:
        return None
    base_dir = Path(__file__).resolve().parent.parent
    root = Path(os.getenv("AGRISKY_OUTPUT_ROOT", str(base_dir / "outputs")))
    if not root.is_absolute():
        root = base_dir / root
    root = root.resolve()
    if url.startswith("/outputs/"):
        candidate = root / url[len("/outputs/") :]
    else:
        raw = Path(url)
        candidate = raw if raw.is_absolute() else root / raw
    try:
        resolved = candidate.resolve()
    except OSError:
        return None
    return resolved if resolved.is_relative_to(root) and resolved.is_file() else None
