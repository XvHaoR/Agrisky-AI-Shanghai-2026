"""Agrisky AI — 受灾评估 Excel 报表。

工作簿只引用服务端案件快照，不在导出阶段重算业务结果。外部字符串统一按文本写入，
避免 Excel 公式注入；比例、面积和金额保留为数值并使用显示格式。
"""

from __future__ import annotations

import math
import re
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from compliance_library import reference_index, report_basis_rows


REPORT_TIMEZONE = timezone(timedelta(hours=8))
INVALID_XML_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
FORMULA_PREFIXES = ("=", "+", "-", "@")


def generate_excel_report(
    claim_id: str,
    data: dict,
    out_path: str,
    *,
    template_version: str = "v1.1",
    generated_at: datetime | None = None,
) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    generated_at = generated_at or datetime.now(REPORT_TIMEZONE)
    case = data.get("case", {}) or {}
    sat = data.get("satellite", {}) or {}
    growth = data.get("growth", {}) or {}
    loss = data.get("loss_assessment", {}) or {}
    comp = data.get("compliance", {}) or {}
    pay = data.get("payout", {}) or {}
    rule = data.get("rule", {}) or {}
    contract = data.get("contract", {}) or {}
    report_meta = data.get("_report_meta", {}) or {}

    wb = Workbook()
    wb.properties.creator = "Agrisky AI"
    wb.properties.lastModifiedBy = "Agrisky AI"
    wb.properties.title = f"农业保险受灾评估报告 - {claim_id}"
    wb.properties.subject = "农业保险遥感辅助查勘定损"
    wb.properties.description = "由服务端权威案件快照导出的结构化评估表"
    wb.properties.keywords = "农业保险,遥感,NDVI,查勘,定损"
    wb.properties.identifier = str(report_meta.get("snapshot_sha256") or claim_id)
    core_time = (
        generated_at.astimezone(timezone.utc).replace(tzinfo=None)
        if generated_at.tzinfo is not None
        else generated_at.replace(tzinfo=None)
    )
    wb.properties.created = core_time
    wb.properties.modified = core_time

    title_font = Font(name="微软雅黑", bold=True, size=16, color="1F2F26")
    subtitle_font = Font(name="微软雅黑", size=9, color="64736A")
    head_fill = PatternFill("solid", fgColor="2F7D4D")
    head_font = Font(name="微软雅黑", bold=True, color="FFFFFF")
    sec_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    sec_fill = PatternFill("solid", fgColor="227A80")
    label_fill = PatternFill("solid", fgColor="F0F4F0")
    warning_fill = PatternFill("solid", fgColor="FDECEC")
    warning_font = Font(name="微软雅黑", bold=True, color="B94B46", size=9)
    thin = Side(style="thin", color="D9E1DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def style_range(sheet: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
        for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.border = border
                cell.alignment = left_wrap

    def put(cell: Any, value: Any, number_format: str | None = None) -> None:
        if isinstance(value, str):
            cell.value = _safe_excel_text(value)
            cell.data_type = "s"
        elif value is None:
            cell.value = "—"
            cell.data_type = "s"
        else:
            cell.value = value
        if number_format and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            cell.number_format = number_format

    def wrapped_row_height(
        values: list[Any],
        column_widths: list[float],
        *,
        minimum: float = 21,
        maximum: float = 90,
    ) -> float:
        line_count = 1
        for value, width in zip(values, column_widths):
            text = "—" if value is None else str(value)
            usable_width = max(5, int(width) - 2)
            lines = sum(
                max(
                    1,
                    math.ceil(
                        sum(2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1 for character in part)
                        / usable_width
                    ),
                )
                for part in text.splitlines() or [""]
            )
            line_count = max(line_count, lines)
        return min(maximum, max(minimum, 15 * line_count + 4))

    def configure_sheet(sheet: Any, *, landscape: bool = True) -> None:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = "landscape" if landscape else "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.35
        sheet.page_margins.right = 0.35
        sheet.page_margins.top = 0.85
        sheet.page_margins.bottom = 0.65
        sheet.page_margins.header = 0.25
        sheet.page_margins.footer = 0.3
        sheet.oddHeader.center.text = f"Agrisky AI｜案件 {claim_id}"
        sheet.oddHeader.center.size = 9
        sheet.oddHeader.center.font = "微软雅黑,Regular"
        sheet.oddFooter.left.text = "人工审核稿"
        sheet.oddFooter.left.size = 8
        sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
        sheet.oddFooter.center.size = 9
        sheet.oddFooter.center.font = "微软雅黑,Regular"
        sheet.oddFooter.right.text = generated_at.strftime("%Y-%m-%d %H:%M UTC+8")
        sheet.oddFooter.right.size = 8

    ws = wb.active
    ws.title = "受灾评估"
    ws.sheet_properties.tabColor = "2F7D4D"
    configure_sheet(ws)
    ws.freeze_panes = "A4"
    ws.print_title_rows = "1:3"
    widths = {"A": 22, "B": 16, "C": 14, "D": 14, "E": 13, "F": 13, "G": 30}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    row_index = 1
    ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)
    ws.cell(row_index, 1, "Agrisky AI 农业保险受灾评估报告").font = title_font
    ws.cell(row_index, 1).alignment = center
    ws.row_dimensions[row_index].height = 27
    row_index += 1
    ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)
    subtitle = f"案件 {claim_id}｜模板 {template_version}｜生成于 {generated_at.strftime('%Y-%m-%d %H:%M UTC+8')}｜人工审核稿"
    put(ws.cell(row_index, 1), subtitle)
    ws.cell(row_index, 1).font = subtitle_font
    ws.cell(row_index, 1).alignment = center
    row_index += 2

    def section(name: str) -> None:
        nonlocal row_index
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)
        cell = ws.cell(row_index, 1, name)
        cell.fill = sec_fill
        cell.font = sec_font
        cell.alignment = left_wrap
        ws.row_dimensions[row_index].height = 22
        row_index += 1

    def kv(rows: list[tuple[str, Any, str | None]]) -> None:
        nonlocal row_index
        for label, value, number_format in rows:
            start_row = row_index
            ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=7)
            style_range(ws, row_index, row_index, 1, 7)
            label_cell = ws.cell(row_index, 1)
            put(label_cell, label)
            label_cell.font = Font(name="微软雅黑", bold=True, size=10, color="1F2F26")
            label_cell.fill = label_fill
            value_cell = ws.cell(row_index, 2)
            put(value_cell, value, number_format)
            value_cell.font = Font(name="微软雅黑", size=10)
            value_cell.alignment = left_wrap
            ws.row_dimensions[start_row].height = 21
            row_index += 1

    def notice(text: str, *, warning: bool = False) -> None:
        nonlocal row_index
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)
        cell = ws.cell(row_index, 1)
        put(cell, text)
        cell.alignment = left_wrap
        cell.fill = warning_fill if warning else label_fill
        cell.font = warning_font if warning else subtitle_font
        style_range(ws, row_index, row_index, 1, 7)
        ws.row_dimensions[row_index].height = 28
        row_index += 1

    section("一、案件基础信息")
    kv(
        [
            ("案件编号", claim_id, None),
            ("保单号", case.get("policy_id"), None),
            ("投保人", case.get("holder_name"), None),
            ("承保地点", case.get("policy_address"), None),
            ("灾害类型", _disaster_name(case.get("disaster_type")), None),
            ("受灾日期", case.get("loss_date"), None),
            ("作物", _crop_name(case.get("crop_type")), None),
            ("地块", case.get("plot_id"), None),
            ("保单登记面积（亩）", _number(case.get("policy_area_mu")), "#,##0.00"),
            ("在册承保边界 SHA-256", case.get("policy_boundary_sha256"), None),
        ]
    )
    row_index += 1

    section("二、保险合同与赔付规则依据")
    if contract:
        contract_period = contract.get("insurance_period") or {}
        terms = contract.get("payout_terms") or {}
        kv(
            [
                ("合同编号", contract.get("contract_number"), None),
                ("合同版本", contract.get("contract_version"), None),
                ("合同 SHA-256", contract.get("contract_sha256"), None),
                ("保险人", contract.get("insurer"), None),
                ("保险期间", f"{contract_period.get('start') or '—'} 至 {contract_period.get('end') or '—'}", None),
                ("每亩保险金额（元）", _number(terms.get("sum_insured_per_mu")), "#,##0.00"),
                ("起赔点", _number(terms.get("deductible_loss_ratio")), "0.0%"),
            ]
        )
        notice(str(contract.get("simulation_disclosure") or "该合同为系统冻结的规则依据。"))
    else:
        notice("未找到对应保单版本的合同规则附件，无法输出合同引用。", warning=True)
    row_index += 1

    section("三、卫星遥感初筛")
    if sat:
        sat_mock = _satellite_is_mock(sat)
        kv(
            [
                ("数据来源", sat.get("source_label") or ("模拟 Sentinel-1 SAR" if sat_mock else "GEE Sentinel-1 SAR"), None),
                ("观测窗口", _date_window(sat.get("start_date"), sat.get("end_date"), inclusive_end=False), None),
                ("疑似受灾面积（亩）", _number(sat.get("suspected_damage_area_mu")), "#,##0.00"),
                ("初筛受损比例", _number(sat.get("damage_ratio")), "0.0%"),
                ("置信度", _confidence_label(sat.get("confidence")), None),
                ("有效影像数（景）", _integer(sat.get("image_count")), "#,##0"),
            ]
        )
        if sat_mock:
            notice("警告：卫星初筛使用模拟数据，仅供流程演示，不可作为定损证据。", warning=True)
    else:
        notice("未生成卫星遥感初筛结果；缺失值不能解释为 0。", warning=True)
    row_index += 1

    section("四、多源遥感减产率评估")
    if loss:
        severity = loss.get("severity") or {}
        kv(
            [
                ("综合减产率", _number(loss.get("yield_loss_ratio")), "0.0%"),
                ("严重度", severity.get("label"), None),
                ("置信度", _confidence_label(loss.get("confidence")), None),
                ("主导因子", "、".join(_safe_list(loss.get("dominant_drivers"))) or "—", None),
                ("数据来源", "/".join(_safe_list(loss.get("data_sources"))) or "—", None),
                ("方法", loss.get("method"), None),
            ]
        )
    else:
        notice("未生成多源遥感减产率评估。", warning=True)

    breakdown = loss.get("product_breakdown") or []
    if breakdown:
        breakdown_header_row = row_index
        headers = ["遥感产品", "传感器", "分辨率（m）", "损失分量", "权重", "贡献", "来源"]
        for column, header in enumerate(headers, start=1):
            cell = ws.cell(row_index, column)
            put(cell, header)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = center
            cell.border = border
        row_index += 1
        for item in breakdown:
            values = [
                (item.get("name_cn"), None),
                (item.get("sensor"), None),
                (_number(item.get("native_res_m")), "0.##"),
                (_number(item.get("decline_score")), "0.0%"),
                (_number(item.get("weight")), "0.0%"),
                (_number(item.get("contribution")), "0.0%"),
                (item.get("data_source"), None),
            ]
            for column, (value, number_format) in enumerate(values, start=1):
                cell = ws.cell(row_index, column)
                put(cell, value, number_format)
                cell.border = border
                cell.alignment = center
            ws.row_dimensions[row_index].height = wrapped_row_height(
                [value for value, _number_format in values],
                [widths[get_column_letter(column)] for column in range(1, 8)],
                minimum=24,
            )
            row_index += 1
        ws.auto_filter.ref = f"A{breakdown_header_row}:G{row_index - 1}"
    row_index += 1

    section("四、合规面积核验")
    if comp:
        kv(
            [
                ("承保面积（亩）", _number(comp.get("insured_area_mu")), "#,##0.00"),
                ("合规受灾面积（亩）", _number(comp.get("valid_damage_area_mu")), "#,##0.00"),
                ("剔除/未计入面积（亩）", _number(comp.get("excluded_area_mu")), "#,##0.00"),
                ("合规受损比例", _number(comp.get("damage_ratio")), "0.0%"),
            ]
        )
    else:
        notice("未生成合规面积核验结果。", warning=True)
    row_index += 1

    # 赔付、规则与最终责任说明作为同一末页语义块，避免规则版本和审核说明孤立。
    ws.row_breaks.append(Break(id=max(1, row_index - 1)))
    section("五、赔付测算")
    if pay:
        kv(
            [
                ("保额（元/亩）", _number(pay.get("sum_insured_per_mu")), "¥#,##0.00"),
                ("承保面积（亩）", _number(pay.get("insured_area_mu")), "#,##0.00"),
                ("总保额（元）", _number(pay.get("total_sum_insured_yuan")), "¥#,##0.00"),
                ("减产率", _number(pay.get("yield_loss_ratio")), "0.0%"),
                ("起赔点", _number(pay.get("deductible_threshold")), "0.0%"),
                ("赔付比例", _number(pay.get("payout_factor")), "0.0%"),
                ("赔付档", pay.get("tier_label"), None),
                ("预估赔款（元）", _number(pay.get("payout_amount_yuan")), "¥#,##0.00"),
            ]
        )
    else:
        notice("未生成赔付测算。", warning=True)
    row_index += 1

    section("六、规则评级")
    if rule:
        review_required = rule.get("review_required")
        kv(
            [
                ("风险等级", _risk_label(rule.get("risk_level")), None),
                ("是否需人工复核", "是" if review_required is True else "否" if review_required is False else "—", None),
                ("规则版本", rule.get("rule_version"), None),
            ]
        )
    else:
        notice("未生成规则评级；缺失值不能解释为低风险。", warning=True)
    row_index += 2
    notice("本表数值由 Agrisky AI 服务端案件快照导出，仅供辅助查勘定损；最终理赔须经人工审核。")
    ws.print_area = f"A1:G{row_index}"

    if growth and (growth.get("summary") or []):
        gws = wb.create_sheet("长势分析")
        gws.sheet_properties.tabColor = "5A9F68"
        configure_sheet(gws)
        gws.sheet_view.showGridLines = False
        growth_widths = {"A": 12, "B": 18, "C": 15, "D": 16, "E": 14, "F": 16, "G": 13}
        for column, width in growth_widths.items():
            gws.column_dimensions[column].width = width
        gws.merge_cells("A1:G1")
        gws["A1"] = "NDVI 作物长势分析"
        gws["A1"].font = title_font
        gws["A1"].alignment = center
        gws.row_dimensions[1].height = 27
        gws.merge_cells("A2:G2")
        put(gws["A2"], f"案件 {claim_id}｜长势等级仅用于筛查，不直接等同于灾损")
        gws["A2"].font = subtitle_font
        gws["A2"].alignment = center

        raster = growth.get("raster", {}) or {}
        meta = raster.get("ndvi_meta", {}) or {}
        source = raster.get("ndvi_source_label") or meta.get("source_label") or "—"
        source_mode = raster.get("ndvi_source") or meta.get("source")
        formula = meta.get("formula")
        if not formula:
            formula = (
                "NDVI = (Sentinel-2 B8 - B4) / (B8 + B4)"
                if source_mode == "gee"
                else "源文件已提供 NDVI 值；原始波段公式未随文件提供"
                if source_mode == "upload"
                else "本地模拟数值（非卫星波段反演）"
                if source_mode == "synthetic"
                else "—"
            )
        info = [
            ("任务号", growth.get("task_id"), None),
            ("数据来源", source, None),
            ("观测窗口", _date_window(meta.get("start_date"), meta.get("end_date")), None),
            ("影像集合", meta.get("collection"), None),
            ("源文件", meta.get("source_filename"), None),
            ("源文件 SHA-256", meta.get("source_sha256"), None),
            ("地块边界来源", meta.get("boundary_source") or meta.get("boundary_filename"), None),
            ("地块边界 SHA-256", meta.get("boundary_sha256"), None),
            ("NDVI 公式/口径", formula, None),
            ("有效影像数（景）", _integer(meta.get("image_count")), "#,##0"),
            ("最大云量阈值", _number(meta.get("max_cloud_pct")), '0.0"%"'),
            ("像元质量掩膜", meta.get("cloud_mask"), None),
            ("空间分辨率（m）", _number(meta.get("scale")), "0.##"),
            ("分级方法", _method_label(growth.get("method")), None),
            ("分级阈值", _class_breaks_label(growth.get("class_breaks")), None),
            ("监测面积（亩）", _number(growth.get("total_area_mu")), "#,##0.00"),
            ("有效像元数", _integer(growth.get("valid_pixel_count")), "#,##0"),
            ("地块内候选像元数", _integer(raster.get("roi_pixel_count")), "#,##0"),
            ("有效像元覆盖率", _number(raster.get("valid_pixel_coverage")), "0.00%"),
            ("栅格坐标系", raster.get("crs"), None),
            ("NoData", _number(raster.get("nodata")), "0.####"),
            (
                "面积折算口径",
                growth.get("area_estimation_method")
                or "各等级有效像元占比 × 地块总面积（对无效像元按有效像元分布作比例外推）",
                None,
            ),
        ]
        grow_row = 4
        for label, value, number_format in info:
            # Provenance labels include hashes and quality terminology.  Giving the
            # label two columns prevents narrow-cell wrapping from producing
            # overlapping text in PDF/print exports, while the value still has five
            # merged columns for long source names and audit hashes.
            gws.merge_cells(start_row=grow_row, start_column=1, end_row=grow_row, end_column=2)
            gws.merge_cells(start_row=grow_row, start_column=3, end_row=grow_row, end_column=7)
            style_range(gws, grow_row, grow_row, 1, 7)
            put(gws.cell(grow_row, 1), label)
            gws.cell(grow_row, 1).font = Font(name="微软雅黑", bold=True)
            for label_column in (1, 2):
                gws.cell(grow_row, label_column).fill = label_fill
            put(gws.cell(grow_row, 3), value, number_format)
            gws.cell(grow_row, 3).alignment = left_wrap
            gws.row_dimensions[grow_row].height = wrapped_row_height(
                [label, value],
                [
                    growth_widths["A"] + growth_widths["B"],
                    sum(growth_widths[column] for column in "CDEFG"),
                ],
            )
            grow_row += 1

        if source_mode == "synthetic":
            gws.merge_cells(start_row=grow_row, start_column=1, end_row=grow_row, end_column=7)
            warning = gws.cell(grow_row, 1)
            put(warning, "警告：本表使用模拟 NDVI，仅供流程演示，不可作为定损证据。")
            warning.font = warning_font
            warning.fill = warning_fill
            warning.alignment = left_wrap
            style_range(gws, grow_row, grow_row, 1, 7)
            grow_row += 2
        else:
            grow_row += 1

        growth_header_row = grow_row
        growth_headers = ["等级", "NDVI 范围", "长势评价", "面积（亩）", "占比", "像元数", "图例"]
        for column, header in enumerate(growth_headers, start=1):
            cell = gws.cell(grow_row, column)
            put(cell, header)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = center
            cell.border = border
        grow_row += 1
        total_area = 0.0
        total_ratio = 0.0
        total_count = 0
        all_areas_valid = True
        all_ratios_valid = True
        all_counts_valid = True
        level_counts: dict[int, int] = {}
        invalid_level_entries = 0
        qa_warnings: list[str] = []
        class_breaks = growth.get("class_breaks") or []
        n_classes_value = _integer(growth.get("n_classes"))
        expected_classes = n_classes_value if n_classes_value is not None else len(growth.get("summary") or [])
        break_values = [_number(value) for value in class_breaks]
        if not class_breaks or (expected_classes > 0 and len(class_breaks) != expected_classes):
            qa_warnings.append("NDVI 分级阈值数量与等级数不一致。")
        if any(value is None for value in break_values):
            qa_warnings.append("NDVI 分级阈值包含非数值项。")
        valid_breaks = [value for value in break_values if value is not None]
        if any(value < -1 or value > 1 for value in valid_breaks):
            qa_warnings.append("NDVI 分级阈值超出 [-1, 1] 物理范围。")
        if len(valid_breaks) == len(break_values) and any(
            left >= right for left, right in zip(valid_breaks, valid_breaks[1:])
        ):
            qa_warnings.append("NDVI 分级阈值未严格递增。")
        for item in sorted(growth.get("summary") or [], key=lambda entry: _number(entry.get("value")) or 0, reverse=True):
            raw_level = _number(item.get("value"))
            level = int(raw_level) if raw_level is not None and raw_level.is_integer() and raw_level > 0 else None
            if level is not None:
                level_counts[level] = level_counts.get(level, 0) + 1
            else:
                invalid_level_entries += 1
            area_value = _number(item.get("area_mu"))
            ratio_value = _number(item.get("ratio"))
            raw_count = _number(item.get("count"))
            count_value = int(raw_count) if raw_count is not None and raw_count.is_integer() else None
            area_valid = area_value is not None and area_value >= 0
            ratio_valid = ratio_value is not None and 0 <= ratio_value <= 1
            count_valid = count_value is not None and count_value >= 0
            all_areas_valid = all_areas_valid and area_valid
            all_ratios_valid = all_ratios_valid and ratio_valid
            all_counts_valid = all_counts_valid and count_valid
            if area_valid:
                total_area += area_value
            if ratio_valid:
                total_ratio += ratio_value
            if count_valid:
                total_count += count_value
            values = [
                (level, "0"),
                (_ndvi_interval(level, class_breaks) if level is not None else "—", None),
                (item.get("label"), None),
                (area_value if area_valid else None, "#,##0.00"),
                (ratio_value if ratio_valid else None, "0.00%"),
                (count_value if count_valid else None, "#,##0"),
                ("", None),
            ]
            for column, (value, number_format) in enumerate(values, start=1):
                cell = gws.cell(grow_row, column)
                put(cell, value, number_format)
                cell.alignment = center
                cell.border = border
            color = str(item.get("color") or "CCCCCC").lstrip("#")
            if not re.fullmatch(r"[0-9A-Fa-f]{6}", color):
                color = "CCCCCC"
            gws.cell(grow_row, 7).fill = PatternFill("solid", fgColor=color.upper())
            grow_row += 1

        growth_data_end_row = grow_row - 1
        if not all_areas_valid:
            qa_warnings.append("部分等级面积缺失、非数值或为负数，合计面积未计算。")
        if not all_ratios_valid:
            qa_warnings.append("部分等级占比缺失或超出 0%–100%，合计占比未计算。")
        if not all_counts_valid:
            qa_warnings.append("部分等级像元数缺失、非整数或为负数，合计像元数未计算。")
        if all_ratios_valid and abs(total_ratio - 1.0) > 0.02:
            qa_warnings.append(f"各等级占比合计为 {total_ratio * 100:.2f}%，与 100% 偏差超过 2 个百分点。")
        declared_area = _number(growth.get("total_area_mu"))
        if all_areas_valid and declared_area is not None:
            tolerance = max(0.1, abs(declared_area) * 0.01)
            if abs(total_area - declared_area) > tolerance:
                qa_warnings.append(
                    f"各等级面积合计 {total_area:,.2f} 亩，与监测面积 {declared_area:,.2f} 亩不一致。"
                )
        declared_count = _number(growth.get("valid_pixel_count"))
        if all_counts_valid and declared_count is not None and (
            not declared_count.is_integer() or total_count != int(declared_count)
        ):
            qa_warnings.append(
                f"各等级像元数合计 {total_count:,}，与顶层有效像元数 {declared_count:,.0f} 不一致。"
            )
        valid_coverage = _number(raster.get("valid_pixel_coverage"))
        if valid_coverage is not None:
            if not 0 <= valid_coverage <= 1:
                qa_warnings.append("有效像元覆盖率超出 0%–100%，覆盖率元数据无效。")
            elif valid_coverage < 0.7:
                qa_warnings.append(
                    f"有效像元覆盖率仅 {valid_coverage * 100:.2f}%，低于 70%；长势分级可能受云、阴影或 NoData 影响。"
                )
        expected_levels = set(range(1, expected_classes + 1)) if expected_classes > 0 else set()
        observed_levels = set(level_counts)
        duplicate_levels = sorted(level for level, count in level_counts.items() if count > 1)
        missing_levels = sorted(expected_levels - observed_levels)
        out_of_range_levels = sorted(observed_levels - expected_levels) if expected_levels else sorted(observed_levels)
        if duplicate_levels:
            qa_warnings.append(f"长势等级存在重复记录：{', '.join(map(str, duplicate_levels))} 级。")
        if missing_levels:
            qa_warnings.append(f"长势等级记录缺失：{', '.join(map(str, missing_levels))} 级。")
        if invalid_level_entries or out_of_range_levels:
            details = f"（{', '.join(map(str, out_of_range_levels))} 级）" if out_of_range_levels else ""
            qa_warnings.append(f"长势等级存在非整数、非正数或超出范围的记录{details}。")

        gws.merge_cells(start_row=grow_row, start_column=1, end_row=grow_row, end_column=3)
        style_range(gws, grow_row, grow_row, 1, 7)
        put(gws.cell(grow_row, 1), "合计")
        gws.cell(grow_row, 1).font = Font(name="微软雅黑", bold=True)
        gws.cell(grow_row, 1).alignment = center
        put(gws.cell(grow_row, 4), total_area if all_areas_valid else None, "#,##0.00")
        put(gws.cell(grow_row, 5), total_ratio if all_ratios_valid else None, "0.00%")
        put(gws.cell(grow_row, 6), total_count if all_counts_valid else None, "#,##0")
        put(gws.cell(grow_row, 7), "—")
        for column in range(4, 8):
            gws.cell(grow_row, column).font = Font(name="微软雅黑", bold=True)
            gws.cell(grow_row, column).alignment = center
        grow_row += 1
        for warning_text in dict.fromkeys(qa_warnings):
            gws.merge_cells(start_row=grow_row, start_column=1, end_row=grow_row, end_column=7)
            warning_cell = gws.cell(grow_row, 1)
            put(warning_cell, f"数据质量警告：{warning_text}")
            warning_cell.font = warning_font
            warning_cell.fill = warning_fill
            warning_cell.alignment = left_wrap
            style_range(gws, grow_row, grow_row, 1, 7)
            gws.row_dimensions[grow_row].height = 28
            grow_row += 1
        grow_row += 1
        gws.merge_cells(start_row=grow_row, start_column=1, end_row=grow_row, end_column=7)
        put(gws.cell(grow_row, 1), "说明：NDVI 公式/来源以本页溯源信息为准；长势等级不直接等同于灾损等级。")
        gws.cell(grow_row, 1).font = subtitle_font
        gws.cell(grow_row, 1).alignment = left_wrap
        gws.freeze_panes = f"A{growth_header_row + 1}"
        gws.auto_filter.ref = f"A{growth_header_row}:G{growth_data_end_row}"
        gws.print_title_rows = f"{growth_header_row}:{growth_header_row}"
        gws.print_area = f"A1:G{grow_row}"
    else:
        gws = wb.create_sheet("长势分析")
        gws.sheet_properties.tabColor = "B94B46"
        configure_sheet(gws)
        gws.sheet_view.showGridLines = False
        for column, width in {"A": 18, "B": 18, "C": 16, "D": 16, "E": 16, "F": 16, "G": 18}.items():
            gws.column_dimensions[column].width = width
        gws.merge_cells("A1:G1")
        put(gws["A1"], "NDVI 作物长势分析")
        gws["A1"].font = title_font
        gws["A1"].alignment = center
        gws.merge_cells("A2:G2")
        put(gws["A2"], f"案件 {claim_id}｜长势等级仅用于筛查，不直接等同于灾损")
        gws["A2"].font = subtitle_font
        gws["A2"].alignment = center
        gws.merge_cells("A4:G4")
        put(gws["A4"], "未执行长势分析或没有可用的 NDVI 长势结果；缺失不能解释为 0 亩或低长势风险。")
        gws["A4"].font = warning_font
        gws["A4"].fill = warning_fill
        gws["A4"].alignment = left_wrap
        style_range(gws, 4, 4, 1, 7)
        gws.row_dimensions[4].height = 34
        gws.freeze_panes = "A4"
        gws.print_area = "A1:G4"

    materials = data.get("materials") or {}
    material_ws = wb.create_sheet("材料核验")
    configure_sheet(material_ws)
    material_ws.freeze_panes = "A4"
    for column, width in {"A": 20, "B": 30, "C": 18, "D": 18, "E": 52}.items():
        material_ws.column_dimensions[column].width = width
    material_ws.merge_cells("A1:E1")
    material_ws["A1"] = "理赔材料理解与一致性核验"
    material_ws["A1"].font = title_font
    material_ws["A1"].alignment = center
    material_ws.merge_cells("A2:E2")
    material_ws["A2"] = "所有提取字段均保留来源引用；开放问题须经人工核验。"
    material_ws["A2"].font = subtitle_font
    material_ws["A2"].alignment = center
    headers = ["材料/字段", "提取值或文件名", "状态/置信度", "证据引用", "说明/SHA-256"]
    for column, header in enumerate(headers, start=1):
        cell = material_ws.cell(3, column, header)
        cell.fill = sec_fill
        cell.font = sec_font
        cell.alignment = center
    material_row = 4
    for item in materials.get("documents") or []:
        values = [
            item.get("document_type"),
            item.get("original_filename"),
            item.get("parse_status"),
            "文件",
            item.get("sha256"),
        ]
        for column, value in enumerate(values, start=1):
            put(material_ws.cell(material_row, column), value)
            material_ws.cell(material_row, column).alignment = left_wrap
        material_row += 1
    for item in materials.get("fields") or []:
        confidence = _number(item.get("confidence"))
        values = [
            item.get("field_name"),
            item.get("normalized_value"),
            f"{confidence * 100:.1f}%" if confidence is not None else "—",
            f"第 {item.get('page_number') or '—'} 页 · {item.get('source_ref') or '—'}",
            f"提取器：{item.get('extractor') or '—'}",
        ]
        for column, value in enumerate(values, start=1):
            put(material_ws.cell(material_row, column), value)
            material_ws.cell(material_row, column).alignment = left_wrap
        material_row += 1
    for item in materials.get("findings") or []:
        if item.get("status") != "open":
            continue
        values = [
            item.get("code"),
            item.get("actual_value"),
            item.get("severity"),
            item.get("source_ref"),
            item.get("message"),
        ]
        for column, value in enumerate(values, start=1):
            put(material_ws.cell(material_row, column), value)
            material_ws.cell(material_row, column).alignment = left_wrap
            material_ws.cell(material_row, column).fill = warning_fill
        material_row += 1
    material_ws.auto_filter.ref = f"A3:E{max(3, material_row - 1)}"

    basis_ws = wb.create_sheet("结论依据映射")
    basis_ws.sheet_properties.tabColor = "176B57"
    configure_sheet(basis_ws)
    basis_ws.freeze_panes = "A4"
    for column, width in {"A": 24, "B": 58, "C": 42, "D": 48}.items():
        basis_ws.column_dimensions[column].width = width
    basis_ws.merge_cells("A1:D1")
    basis_ws["A1"] = "案件结论与合规依据映射"
    basis_ws["A1"].font = title_font
    basis_ws["A1"].alignment = center
    basis_ws.merge_cells("A2:D2")
    basis_ws["A2"] = "本表由服务端固定规则生成；Agent 可检索和解释依据，但不得改写合同参数或直接作出赔付决定。"
    basis_ws["A2"].font = subtitle_font
    basis_ws["A2"].alignment = center
    for column, header in enumerate(["业务环节", "本案结论口径", "引用依据", "适用边界"], start=1):
        cell = basis_ws.cell(3, column, header)
        cell.fill = sec_fill
        cell.font = sec_font
        cell.alignment = center
    basis_row = 4
    for item in report_basis_rows(contract, str(case.get("crop_type") or "")):
        values = [
            item.get("stage"),
            item.get("conclusion"),
            " ".join(f"[{basis}]" for basis in item.get("basis_ids") or []),
            item.get("boundary"),
        ]
        for column, value in enumerate(values, start=1):
            put(basis_ws.cell(basis_row, column), value)
            basis_ws.cell(basis_row, column).alignment = left_wrap
            basis_ws.cell(basis_row, column).border = border
        basis_ws.row_dimensions[basis_row].height = wrapped_row_height(
            values, [24, 58, 42, 48], minimum=34, maximum=100
        )
        basis_row += 1
    basis_ws.auto_filter.ref = f"A3:D{max(3, basis_row - 1)}"
    basis_ws.print_area = f"A1:D{max(3, basis_row - 1)}"

    reference_ws = wb.create_sheet("依据索引")
    reference_ws.sheet_properties.tabColor = "227A80"
    configure_sheet(reference_ws)
    reference_ws.freeze_panes = "A4"
    for column, width in {"A": 12, "B": 16, "C": 34, "D": 28, "E": 58, "F": 24}.items():
        reference_ws.column_dimensions[column].width = width
    reference_ws.merge_cells("A1:F1")
    reference_ws["A1"] = "法规、合同、技术与数据依据索引"
    reference_ws["A1"].font = title_font
    reference_ws["A1"].alignment = center
    reference_ws.merge_cells("A2:F2")
    reference_ws["A2"] = "L=法律法规，C=冻结合同及条款，T=技术依据，D=服务端证据数据。"
    reference_ws["A2"].font = subtitle_font
    reference_ws["A2"].alignment = center
    reference_headers = ["编号", "类别", "名称/条款", "发布/合同主体", "官方来源、摘要或条款内容", "SHA-256"]
    for column, header in enumerate(reference_headers, start=1):
        cell = reference_ws.cell(3, column, header)
        cell.fill = sec_fill
        cell.font = sec_font
        cell.alignment = center
    reference_row = 4
    for item in reference_index(contract, str(case.get("crop_type") or "")):
        values = [
            item.get("id"),
            item.get("category"),
            item.get("title"),
            item.get("issuer"),
            "\n".join(
                value
                for value in [
                    str(item.get("applicability") or ""),
                    str(item.get("summary") or ""),
                    str(item.get("source") or ""),
                ]
                if value
            ),
            item.get("sha256"),
        ]
        for column, value in enumerate(values, start=1):
            put(reference_ws.cell(reference_row, column), value)
            reference_ws.cell(reference_row, column).alignment = left_wrap
            reference_ws.cell(reference_row, column).border = border
        reference_ws.row_dimensions[reference_row].height = wrapped_row_height(
            values, [12, 16, 34, 28, 58, 24], minimum=28, maximum=80
        )
        reference_row += 1
    reference_ws.auto_filter.ref = f"A3:F{max(3, reference_row - 1)}"
    reference_ws.print_area = f"A1:F{max(3, reference_row - 1)}"

    meta_ws = wb.create_sheet("导出说明")
    meta_ws.sheet_properties.tabColor = "7D8B83"
    configure_sheet(meta_ws, landscape=False)
    meta_ws.sheet_view.showGridLines = False
    meta_ws.column_dimensions["A"].width = 24
    meta_ws.column_dimensions["B"].width = 62
    meta_ws.merge_cells("A1:B1")
    meta_ws["A1"] = "报告导出说明与追溯信息"
    meta_ws["A1"].font = title_font
    meta_ws["A1"].alignment = center
    meta_rows = [
        ("案件编号", claim_id),
        ("报告状态", "人工审核稿"),
        ("模板版本", template_version),
        ("生成时间", generated_at.strftime("%Y-%m-%d %H:%M:%S UTC+8")),
        ("数据快照 SHA-256", report_meta.get("snapshot_sha256")),
        ("报告生成号", report_meta.get("generation_id")),
        ("长势结果状态", "已包含服务端长势快照" if growth and (growth.get("summary") or []) else "未执行或无可用结果"),
        ("数据原则", "业务数值来自服务端案件快照，导出阶段不重新计算。"),
        ("NDVI 口径", "长势等级用于遥感筛查，不直接等同于灾损等级。"),
        ("最终责任", "最终理赔结论须经人工审核，并结合保险条款与现场证据。"),
    ]
    meta_row = 3
    for label, value in meta_rows:
        style_range(meta_ws, meta_row, meta_row, 1, 2)
        put(meta_ws.cell(meta_row, 1), label)
        meta_ws.cell(meta_row, 1).font = Font(name="微软雅黑", bold=True)
        meta_ws.cell(meta_row, 1).fill = label_fill
        put(meta_ws.cell(meta_row, 2), value)
        meta_ws.cell(meta_row, 2).alignment = left_wrap
        meta_ws.row_dimensions[meta_row].height = 28
        meta_row += 1
    meta_ws.freeze_panes = "A3"
    meta_ws.print_area = f"A1:B{meta_row - 1}"

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name is None:
                    cell.font = Font(name="微软雅黑", size=10)
                if cell.value is not None and cell.alignment == Alignment():
                    cell.alignment = left_wrap

    wb.active = 0
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)


def _safe_excel_text(value: str) -> str:
    """Write untrusted strings as literals and strip characters forbidden by OOXML."""
    text = INVALID_XML_CHARS.sub("", str(value))[:32767]
    stripped = text.lstrip()
    if len(stripped) > 1 and stripped.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _integer(value: Any) -> int | None:
    number = _number(value)
    return None if number is None else int(number)


def _safe_list(value: Any) -> list[str]:
    if not isinstance(value, (list, tuple)):
        return []
    return [str(item) for item in value if item not in (None, "")]


def _date_window(start: Any, end: Any, *, inclusive_end: bool = True) -> str:
    if start in (None, "") and end in (None, ""):
        return "—"
    semantics = "起止日期均含" if inclusive_end else "起始含、截止不含"
    return f"{start or '—'} 至 {end or '—'}（{semantics}）"


def _class_breaks_label(values: Any) -> str:
    if not isinstance(values, (list, tuple)):
        return "—"
    numbers = [_number(value) for value in values]
    return " / ".join(f"{value:.2f}" for value in numbers if value is not None) or "—"


def _ndvi_interval(level: int, values: Any) -> str:
    if not isinstance(values, (list, tuple)):
        return "—"
    breaks = [number for value in values if (number := _number(value)) is not None]
    if level <= 0 or level > len(breaks):
        return "—"
    if level == 1:
        return f"< {breaks[0]:.2f}"
    lower = breaks[level - 2]
    if level == len(breaks):
        return f"≥ {lower:.2f}"
    return f"[{lower:.2f}, {breaks[level - 1]:.2f})"


def _satellite_is_mock(sat: dict) -> bool:
    if sat.get("is_mock") is True or sat.get("source") in {"mock", "synthetic"}:
        return True
    if str(sat.get("confidence") or "").lower() == "mock":
        return True
    return any(str(asset).lower().startswith("mock_") for asset in sat.get("reference_assets", []) or [])


def _method_label(value: Any) -> str:
    key = str(value or "").lower()
    return {
        "fixed": "固定阈值（0.30 / 0.45 / 0.60 / 0.75）",
        "jenks": "Jenks 自然断点",
        "equalinterval": "等距分级",
        "quantile": "分位数分级",
        "std": "标准差分级",
    }.get(key, str(value or "—"))


def _disaster_name(value: Any) -> str:
    return {
        "flood": "洪水",
        "drought": "干旱",
        "hail": "冰雹",
        "typhoon": "台风",
        "pest": "虫灾",
        "frost": "霜冻",
        "other": "其他",
    }.get(str(value or ""), str(value or "—"))


def _crop_name(value: Any) -> str:
    return {
        "rice": "水稻",
        "wheat": "小麦",
        "corn": "玉米",
        "maize": "玉米",
        "soybean": "大豆",
        "cotton": "棉花",
        "peanut": "花生",
    }.get(str(value or "").lower(), str(value or "—"))


def _confidence_label(value: Any) -> str:
    return {
        "high": "高",
        "medium": "中",
        "low": "低",
        "mock": "模拟数据（无证据置信度）",
    }.get(str(value or "").lower(), str(value or "—"))


def _risk_label(value: Any) -> str:
    return {"high": "高风险", "medium": "中风险", "low": "低风险"}.get(
        str(value or "").lower(), str(value or "—")
    )
