"""Private claim-document parsing and evidence-grounded field extraction.

The module keeps document facts separate from authoritative policy and claim
records. Extracted values are evidence for review; they never overwrite a
policy or drive payout calculations without deterministic checks.
"""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import os
import re
import time
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from docx import Document

try:
    import pypdfium2 as pdfium
except ImportError:  # pragma: no cover - dependency check is surfaced at runtime
    pdfium = None

# PDFium is not thread safe, including across different documents.
_PDFIUM_LOCK = threading.RLock()

try:
    from rapidocr import RapidOCR
except ImportError:  # pragma: no cover - OCR remains optional for text-native files
    RapidOCR = None


MAX_DOCUMENT_BYTES = int(os.getenv("AGRISKY_DOCUMENT_MAX_BYTES", str(20 * 1024 * 1024)))
MAX_PDF_PAGES = int(os.getenv("AGRISKY_DOCUMENT_MAX_PAGES", "60"))
MAX_LLM_TEXT_CHARS = int(os.getenv("AGRISKY_DOCUMENT_LLM_MAX_CHARS", "24000"))

DOCUMENT_TYPES = {
    "policy_document": "保单或承保凭证",
    "claim_notice": "出险通知或理赔申请",
    "damage_certificate": "灾情证明或查勘记录",
    "onsite_photo": "现场照片",
    "other": "其他材料",
}

ALLOWED_SUFFIXES = {".pdf", ".docx", ".png", ".jpg", ".jpeg"}
ALLOWED_MEDIA_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "image/png",
    "image/jpeg",
}

_OCR_ENGINE: Any | None = None


@dataclass(frozen=True)
class EvidenceLine:
    line_id: str
    page: int
    text: str
    bbox: list[float] | None = None
    confidence: float = 1.0


def safe_filename(filename: str) -> str:
    name = Path(filename or "document").name
    name = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "_", name).strip("._")
    return (name or "document")[:180]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def detect_media_type(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError("仅支持 PDF、DOCX、PNG 和 JPEG 材料")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError(f"单个材料不能超过 {MAX_DOCUMENT_BYTES // (1024 * 1024)} MB")
    if not data:
        raise ValueError("材料文件为空")

    if data.startswith(b"%PDF-"):
        detected = "application/pdf"
    elif data.startswith(b"\x89PNG\r\n\x1a\n"):
        detected = "image/png"
    elif data.startswith(b"\xff\xd8\xff"):
        detected = "image/jpeg"
    elif data.startswith(b"PK\x03\x04") and suffix == ".docx":
        detected = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        guessed = mimetypes.guess_type(filename)[0]
        detected = guessed or "application/octet-stream"
    if detected not in ALLOWED_MEDIA_TYPES:
        raise ValueError("文件内容与允许的材料格式不匹配")
    expected_suffixes = {
        "application/pdf": {".pdf"},
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": {".docx"},
        "image/png": {".png"},
        "image/jpeg": {".jpg", ".jpeg"},
    }[detected]
    if suffix not in expected_suffixes:
        raise ValueError("文件扩展名与实际内容不一致")
    return detected


def _ocr_engine() -> Any:
    global _OCR_ENGINE
    if RapidOCR is None:
        raise RuntimeError("图片 OCR 组件未安装")
    if _OCR_ENGINE is None:
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_image(value: bytes | str | Path, page: int) -> list[EvidenceLine]:
    result = _ocr_engine()(value)
    texts = list(getattr(result, "txts", ()) or ())
    boxes = list(getattr(result, "boxes", ()) or ())
    scores = list(getattr(result, "scores", ()) or ())
    lines: list[EvidenceLine] = []
    for index, text in enumerate(texts, start=1):
        cleaned = " ".join(str(text).split())
        if not cleaned:
            continue
        box = boxes[index - 1].tolist() if index - 1 < len(boxes) else None
        flattened = [round(float(coord), 2) for point in (box or []) for coord in point]
        score = float(scores[index - 1]) if index - 1 < len(scores) else 0.5
        lines.append(
            EvidenceLine(
                line_id=f"p{page}-ocr{index}",
                page=page,
                text=cleaned,
                bbox=flattened or None,
                confidence=round(score, 4),
            )
        )
    return lines


def _extract_pdf(path: Path) -> tuple[list[EvidenceLine], str]:
    if pdfium is None:
        raise RuntimeError("pypdfium2 未安装，无法解析 PDF")
    lines: list[EvidenceLine] = []
    used_ocr = False
    with _PDFIUM_LOCK, pdfium.PdfDocument(path) as document:
        if len(document) > MAX_PDF_PAGES:
            raise ValueError(f"PDF 页数不能超过 {MAX_PDF_PAGES} 页")
        for page_index in range(len(document)):
            page_number = page_index + 1
            page = document[page_index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded()
                finally:
                    textpage.close()
                cleaned_lines = [" ".join(line.split()) for line in text.splitlines()]
                # Native PDF evidence retains page/line citations. Do not invent
                # precise rectangles when the extractor returns plain text.
                page_lines = [
                    EvidenceLine(f"p{page_number}-l{index}", page_number, line)
                    for index, line in enumerate(filter(None, cleaned_lines), start=1)
                ]
                if sum(len(item.text) for item in page_lines) < 20:
                    bitmap = page.render(scale=2)
                    try:
                        with io.BytesIO() as output:
                            bitmap.to_pil().save(output, format="PNG")
                            page_lines = _ocr_image(output.getvalue(), page_number)
                    finally:
                        bitmap.close()
                    used_ocr = True
                lines.extend(page_lines)
            finally:
                page.close()
    return lines, "pypdfium2+rapidocr" if used_ocr else "pypdfium2"


def _extract_docx(path: Path) -> tuple[list[EvidenceLine], str]:
    document = Document(path)
    lines: list[EvidenceLine] = []
    index = 0
    for paragraph in document.paragraphs:
        cleaned = " ".join(paragraph.text.split())
        if cleaned:
            index += 1
            lines.append(EvidenceLine(f"p1-l{index}", 1, cleaned))
    for table in document.tables:
        for row in table.rows:
            cleaned = " | ".join(
                " ".join(cell.text.split()) for cell in row.cells if cell.text.strip()
            )
            if cleaned:
                index += 1
                lines.append(EvidenceLine(f"p1-l{index}", 1, cleaned))
    return lines, "python-docx"


def extract_evidence(path: Path, media_type: str) -> tuple[list[EvidenceLine], str]:
    if media_type == "application/pdf":
        return _extract_pdf(path)
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(path)
    if media_type.startswith("image/"):
        return _ocr_image(path, 1), "rapidocr"
    raise ValueError(f"不支持解析的媒体类型: {media_type}")


def redact_sensitive_text(text: str) -> str:
    text = re.sub(r"(?<!\d)\d{17}[\dXx](?!\d)", "[身份证号已脱敏]", text)
    text = re.sub(r"(?<!\d)1[3-9]\d{9}(?!\d)", "[手机号已脱敏]", text)
    text = re.sub(r"(?<!\d)\d{12,19}(?!\d)", "[账户号已脱敏]", text)
    return text


def _normalized_date(value: str) -> str | None:
    match = re.search(r"(20\d{2})[年./\-](\d{1,2})[月./\-](\d{1,2})", value)
    if not match:
        return None
    year, month, day = map(int, match.groups())
    try:
        return f"{year:04d}-{month:02d}-{day:02d}"
    except ValueError:
        return None


def _number(value: str) -> float | None:
    match = re.search(r"\d+(?:,\d{3})*(?:\.\d+)?", value)
    return float(match.group(0).replace(",", "")) if match else None


def _field(
    name: str,
    value: Any,
    line: EvidenceLine,
    confidence: float,
    extractor: str,
) -> dict[str, Any]:
    return {
        "field_name": name,
        "normalized_value": value,
        "raw_text": line.text,
        "confidence": round(min(1.0, max(0.0, confidence * line.confidence)), 4),
        "page_number": line.page,
        "bbox": line.bbox,
        "source_ref": line.line_id,
        "extractor": extractor,
    }


def deterministic_fields(lines: list[EvidenceLine]) -> list[dict[str, Any]]:
    fields: dict[str, dict[str, Any]] = {}
    crop_terms = {
        "水稻": "rice",
        "稻谷": "rice",
        "小麦": "wheat",
        "玉米": "corn",
        "大豆": "soybean",
        "棉花": "cotton",
        "rice": "rice",
        "wheat": "wheat",
        "corn": "corn",
    }
    disaster_terms = {
        "洪涝": "flood",
        "洪水": "flood",
        "内涝": "flood",
        "干旱": "drought",
        "冰雹": "hail",
        "台风": "typhoon",
        "病虫害": "pest",
        "霜冻": "frost",
        "flood": "flood",
        "drought": "drought",
        "hail": "hail",
        "typhoon": "typhoon",
        "pest": "pest",
        "frost": "frost",
    }

    for line in lines:
        text = line.text
        lower = text.lower()
        if "policy" in lower or "保单" in text:
            matches = re.findall(r"\b[A-Z][A-Z0-9][A-Z0-9._\-]{4,}\b", text.upper())
            if matches and "policy_id" not in fields:
                fields["policy_id"] = _field("policy_id", matches[-1], line, 0.98, "regex")
        if any(label in text for label in ("出险日期", "灾害日期", "发生日期")):
            value = _normalized_date(text)
            if value and "loss_date" not in fields:
                fields["loss_date"] = _field("loss_date", value, line, 0.97, "regex")
        if any(label in text for label in ("承保面积", "保险面积")):
            value = _number(text)
            if value is not None and "insured_area_mu" not in fields:
                fields["insured_area_mu"] = _field(
                    "insured_area_mu", value, line, 0.94, "regex"
                )
        if any(label in text for label in ("申报受灾面积", "受灾面积", "损失面积")):
            value = _number(text)
            if value is not None and "reported_damage_area_mu" not in fields:
                fields["reported_damage_area_mu"] = _field(
                    "reported_damage_area_mu", value, line, 0.9, "regex"
                )
        if any(label in text for label in ("损失比例", "受损比例", "减产率")):
            value = _number(text)
            if value is not None:
                if "%" in text:
                    value /= 100
                if "reported_loss_ratio" not in fields:
                    fields["reported_loss_ratio"] = _field(
                        "reported_loss_ratio", value, line, 0.9, "regex"
                    )
        if any(label in text for label in ("地块编号", "地块号", "plot")):
            match = re.search(
                r"(?:地块编号|地块号|plot(?:\s*id)?)\s*[:：]?\s*([A-Za-z0-9._\-]+)",
                text,
                re.I,
            )
            if match and "plot_id" not in fields:
                fields["plot_id"] = _field("plot_id", match.group(1), line, 0.92, "regex")
        for term, normalized in crop_terms.items():
            if term.lower() in lower and "crop_type" not in fields:
                fields["crop_type"] = _field("crop_type", normalized, line, 0.88, "dictionary")
        for term, normalized in disaster_terms.items():
            if term.lower() in lower and "disaster_type" not in fields:
                fields["disaster_type"] = _field(
                    "disaster_type", normalized, line, 0.9, "dictionary"
                )
    return list(fields.values())


def _llm_json(messages: list[dict[str, str]]) -> dict[str, Any] | None:
    api_key = os.getenv("AGENT_API_KEY", "").strip()
    base_url = (
        os.getenv("AGENT_BASE_URL", "").strip() or "https://api.deepseek.com"
    ).rstrip("/")
    model = os.getenv("AGENT_MODEL", "").strip() or "deepseek-v4-pro"
    if not api_key:
        return None
    payload = json.dumps(
        {
            "model": model,
            "messages": messages,
            "response_format": {"type": "json_object"},
            "thinking": {"type": "disabled"},
            "temperature": 0,
            "max_tokens": 1600,
        },
        ensure_ascii=False,
    ).encode("utf-8")
    request = Request(
        f"{base_url}/chat/completions",
        data=payload,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
        method="POST",
    )
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            with urlopen(request, timeout=90) as response:
                body = json.load(response)
            content = body["choices"][0]["message"].get("content") or "{}"
            return json.loads(content)
        except (HTTPError, URLError, TimeoutError, KeyError, ValueError, json.JSONDecodeError) as exc:
            last_error = exc
            if isinstance(exc, HTTPError) and exc.code not in {429, 500, 502, 503, 504}:
                break
            time.sleep(0.75 * (2**attempt))
    if last_error:
        return {"_error": str(last_error)}
    return None


def llm_fields(lines: list[EvidenceLine], document_type: str) -> tuple[list[dict[str, Any]], str | None]:
    line_lookup = {line.line_id: line for line in lines}
    rendered = "\n".join(
        f"[{line.line_id}] {redact_sensitive_text(line.text)}" for line in lines
    )[:MAX_LLM_TEXT_CHARS]
    if not rendered.strip():
        return [], None
    system = """你是农业保险材料结构化抽取器。只输出 JSON，不做理赔结论。
仅从提供的证据行提取字段，不得猜测。每个字段必须引用一个真实 source_ref。
JSON 格式：
{"fields":[{"field_name":"policy_id|holder_name|crop_type|insured_area_mu|loss_date|disaster_type|plot_id|reported_damage_area_mu|reported_loss_ratio|survey_date","normalized_value":"string或number","raw_text":"原文短句","confidence":0到1,"source_ref":"p1-l1"}]}
没有证据的字段不要输出。面积统一为亩，比例统一为0到1，日期统一为YYYY-MM-DD。"""
    result = _llm_json(
        [
            {"role": "system", "content": system},
            {
                "role": "user",
                "content": f"材料类型：{DOCUMENT_TYPES.get(document_type, document_type)}\n证据行：\n{rendered}",
            },
        ]
    )
    if not result:
        return [], None
    if result.get("_error"):
        return [], str(result["_error"])
    extracted: list[dict[str, Any]] = []
    allowed_names = {
        "policy_id",
        "holder_name",
        "crop_type",
        "insured_area_mu",
        "loss_date",
        "disaster_type",
        "plot_id",
        "reported_damage_area_mu",
        "reported_loss_ratio",
        "survey_date",
    }
    for item in result.get("fields", []):
        if not isinstance(item, dict) or item.get("field_name") not in allowed_names:
            continue
        source_ref = str(item.get("source_ref") or "")
        line = line_lookup.get(source_ref)
        if not line:
            continue
        try:
            confidence = float(item.get("confidence", 0.5))
        except (TypeError, ValueError):
            confidence = 0.5
        extracted.append(
            _field(
                str(item["field_name"]),
                item.get("normalized_value"),
                line,
                confidence,
                "llm_json",
            )
        )
    return extracted, None


def merge_fields(
    deterministic: list[dict[str, Any]], semantic: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    merged: dict[str, dict[str, Any]] = {}
    for item in semantic + deterministic:
        name = item["field_name"]
        if name not in merged or float(item["confidence"]) > float(merged[name]["confidence"]):
            merged[name] = item
    return sorted(merged.values(), key=lambda item: item["field_name"])


def normalize_fields_for_document_type(
    fields: list[dict[str, Any]], document_type: str
) -> list[dict[str, Any]]:
    """Apply conservative type-specific semantics after model extraction."""
    normalized: list[dict[str, Any]] = []
    for item in fields:
        candidate = dict(item)
        if document_type == "damage_certificate" and candidate["field_name"] == "loss_date":
            candidate["field_name"] = "survey_date"
            candidate["confidence"] = min(float(candidate["confidence"]), 0.85)
        if document_type == "onsite_photo" and candidate["field_name"] in {
            "policy_id",
            "insured_area_mu",
            "reported_damage_area_mu",
            "reported_loss_ratio",
        }:
            continue
        normalized.append(candidate)
    best: dict[str, dict[str, Any]] = {}
    for item in normalized:
        name = item["field_name"]
        if name not in best or float(item["confidence"]) > float(best[name]["confidence"]):
            best[name] = item
    return sorted(best.values(), key=lambda item: item["field_name"])


def required_document_types(disaster_type: str | None) -> list[str]:
    required = ["policy_document", "claim_notice"]
    if disaster_type in {"flood", "drought", "hail", "typhoon", "pest", "frost"}:
        required.append("damage_certificate")
    return required


def compare_material_fields(
    case: dict[str, Any],
    policy: dict[str, Any] | None,
    fields: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    latest = {item["field_name"]: item for item in fields}
    findings: list[dict[str, Any]] = []

    def add(
        code: str,
        severity: str,
        field_name: str,
        expected: Any,
        actual: Any,
        source_ref: str | None,
        message: str,
    ) -> None:
        findings.append(
            {
                "code": code,
                "severity": severity,
                "field_name": field_name,
                "expected_value": expected,
                "actual_value": actual,
                "source_ref": source_ref,
                "message": message,
            }
        )

    comparisons = [
        ("policy_id", case.get("policy_id"), "POLICY_ID_MISMATCH", "材料保单号与案件不一致"),
        ("crop_type", case.get("crop_type"), "CROP_MISMATCH", "材料作物与案件不一致"),
        ("loss_date", str(case.get("loss_date") or "")[:10], "LOSS_DATE_MISMATCH", "材料出险日期与案件不一致"),
        ("disaster_type", case.get("disaster_type"), "DISASTER_TYPE_MISMATCH", "材料灾害类型与案件不一致"),
        ("plot_id", case.get("plot_id"), "PLOT_ID_MISMATCH", "材料地块编号与案件不一致"),
    ]
    for field_name, expected, code, message in comparisons:
        item = latest.get(field_name)
        actual = item.get("normalized_value") if item else None
        if expected and actual and str(expected).strip().lower() != str(actual).strip().lower():
            add(code, "high", field_name, expected, actual, item.get("source_ref"), message)

    insured = latest.get("insured_area_mu")
    policy_area = policy.get("area_mu") if policy else None
    if insured and policy_area is not None:
        actual_area = float(insured["normalized_value"])
        expected_area = float(policy_area)
        tolerance = max(0.5, expected_area * 0.01)
        if abs(actual_area - expected_area) > tolerance:
            add(
                "INSURED_AREA_MISMATCH",
                "high",
                "insured_area_mu",
                round(expected_area, 4),
                round(actual_area, 4),
                insured.get("source_ref"),
                "材料承保面积与在册保单面积不一致",
            )

    damage = latest.get("reported_damage_area_mu")
    ceiling = float(policy_area) if policy_area is not None else None
    if damage and ceiling is not None and float(damage["normalized_value"]) > ceiling * 1.01:
        add(
            "DAMAGE_AREA_EXCEEDS_POLICY",
            "high",
            "reported_damage_area_mu",
            round(ceiling, 4),
            damage["normalized_value"],
            damage.get("source_ref"),
            "申报受灾面积超过在册承保面积",
        )

    for item in fields:
        if float(item.get("confidence") or 0) < 0.7:
            add(
                f"LOW_CONFIDENCE_{item['field_name'].upper()}",
                "medium",
                item["field_name"],
                None,
                item.get("normalized_value"),
                item.get("source_ref"),
                "字段提取置信度较低，需要人工确认",
            )
    return findings
