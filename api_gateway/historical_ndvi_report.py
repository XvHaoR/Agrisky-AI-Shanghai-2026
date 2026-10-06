"""Charts and an A4 DOCX report for validated historical quarterly NDVI data."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from historical_ndvi import (
    LOW_VALID_COVERAGE,
    HistoricalNDVIOutputs,
    HistoricalNDVIResult,
    QuarterlyNDVIObservation,
    save_historical_ndvi_result,
)


NDVI_LEVELS = (
    (-1.0, 0.30, "差", "#d64545"),
    (0.30, 0.45, "一般", "#ef8f35"),
    (0.45, 0.60, "中", "#e4c441"),
    (0.60, 0.75, "良", "#8cc152"),
    (0.75, 1.00, "优", "#2f9e44"),
)
QUARTER_NAMES = {
    1: "第一季度（1—3月）",
    2: "第二季度（4—6月）",
    3: "第三季度（7—9月）",
    4: "第四季度（10—12月）",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalise_png_for_office(path: Path) -> None:
    """Remove the alpha channel that some Office/PDF renderers show as black."""

    from PIL import Image

    temporary = path.with_name(f".{path.name}.rgb.tmp.png")
    try:
        with Image.open(path) as source:
            rgba = source.convert("RGBA")
            background = Image.new("RGBA", rgba.size, "white")
            background.alpha_composite(rgba)
            background.convert("RGB").save(temporary, format="PNG")
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _configure_matplotlib() -> Any:
    import matplotlib

    matplotlib.use("Agg")
    from matplotlib import font_manager, rcParams

    installed = {font.name for font in font_manager.fontManager.ttflist}
    candidates = [
        "Microsoft YaHei",
        "Noto Sans CJK SC",
        "Noto Sans CJK JP",
        "SimHei",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    rcParams["font.sans-serif"] = [name for name in candidates if name in installed] or [
        "DejaVu Sans"
    ]
    rcParams["axes.unicode_minus"] = False
    return matplotlib


def _observations_for_year(
    result: HistoricalNDVIResult,
    year: int,
) -> list[QuarterlyNDVIObservation]:
    return [item for item in result.quarters if item.window.year == year]


def _verified_quarter_map(
    observation: QuarterlyNDVIObservation,
    artifact_root: Path,
) -> tuple[Path | None, str | None]:
    if not observation.quarter_map_png:
        return None, None
    root = artifact_root.resolve()
    candidate = (root / observation.quarter_map_png).resolve()
    if not candidate.is_relative_to(root):
        return None, f"{observation.window.label} 季度图路径越界，已拒绝嵌入"
    if not candidate.is_file():
        return None, f"{observation.window.label} 季度图文件缺失，已使用空白占位"
    if _sha256(candidate) != observation.quarter_map_sha256:
        return None, f"{observation.window.label} 季度图哈希不一致，已拒绝嵌入"
    return candidate, None


def render_quarter_trend_chart(
    result: HistoricalNDVIResult,
    output_path: Path,
) -> Path:
    """Render the cross-year Q1-Q4 curve without filling missing quarters."""
    _configure_matplotlib()
    import matplotlib.pyplot as plt
    import numpy as np

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure, axis = plt.subplots(figsize=(12.2, 6.8), dpi=180)

    for lower, upper, _label, color in NDVI_LEVELS:
        axis.axhspan(lower, upper, color=color, alpha=0.065, zorder=0)

    years = range(result.start_year, result.end_year + 1)
    preferred_colours = [
        "#35618d",
        "#cf7414",
        "#3e8d5a",
        "#b24842",
        "#7650a5",
        "#7b5b42",
        "#b65a87",
        "#5b777d",
    ]
    fallback_colours = plt.get_cmap("tab20")
    plotted_series = 0
    for index, year in enumerate(years):
        observations = _observations_for_year(result, year)
        values = [
            item.median_ndvi if item.usable else np.nan
            for item in observations
        ]
        if all(np.isnan(value) for value in values):
            continue
        has_partial = any(item.status == "partial" for item in observations)
        label = f"{year}（截至 {result.as_of_date.isoformat()}）" if has_partial else str(year)
        axis.plot(
            [1, 2, 3, 4],
            values,
            marker="o",
            linewidth=2.2,
            markersize=6,
            color=(
                preferred_colours[index]
                if index < len(preferred_colours)
                else fallback_colours(index % 20)
            ),
            label=label,
        )
        plotted_series += 1

    axis.set_title(
        f"{result.scope_label} {result.start_year}—{result.end_year} 年季度 NDVI 周期曲线",
        fontsize=17,
        color="#215b40",
        pad=18,
    )
    axis.set_xlabel("季度", fontsize=11)
    axis.set_ylabel("地块内季度中位 NDVI", fontsize=11)
    axis.set_xticks(
        [1, 2, 3, 4],
        ["第一季度\n1—3月", "第二季度\n4—6月", "第三季度\n7—9月", "第四季度\n10—12月"],
    )
    axis.set_xlim(0.75, 4.25)
    axis.set_ylim(-0.1, 1.0)
    axis.set_yticks([round(-0.1 + index * 0.1, 1) for index in range(12)])
    axis.grid(True, color="#aab8ae", alpha=0.35, linewidth=0.7)
    if plotted_series:
        axis.legend(
            loc="upper center",
            bbox_to_anchor=(0.5, 1.01),
            ncol=min(5, plotted_series),
        )
    else:
        axis.text(
            0.5,
            0.5,
            "所选期间没有可用的季度 NDVI 统计",
            transform=axis.transAxes,
            ha="center",
            va="center",
            color="#68736c",
            fontsize=13,
        )
    axis.spines[["top", "right"]].set_visible(False)
    figure.text(
        0.5,
        0.012,
        "季度值由季度 Sentinel-2 有效影像逐像元中位合成后取地块内像元中位数；缺失季度保持空白，不进行时间插值。",
        ha="center",
        fontsize=9,
        color="#5d655f",
    )
    figure.tight_layout(rect=(0.02, 0.05, 0.98, 0.98))
    figure.savefig(output_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    _normalise_png_for_office(output_path)
    return output_path


def _placeholder_text(observation: QuarterlyNDVIObservation) -> str:
    return {
        "not_reached": "尚未到达\n保持空白",
        "no_imagery": "无满足云量阈值的影像",
        "no_valid_pixels": "无有效地块内像元",
        "failed": "季度计算失败",
    }.get(observation.status, "季度图未生成")


def render_year_quarter_panel(
    result: HistoricalNDVIResult,
    year: int,
    output_path: Path,
    *,
    artifact_root: Path,
) -> tuple[Path, list[str]]:
    """Render one 2×2 yearly panel with identical NDVI level semantics."""
    if not result.start_year <= year <= result.end_year:
        raise ValueError("year 不在历史结果范围内")
    _configure_matplotlib()
    import matplotlib.image as mpimg
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    observations = _observations_for_year(result, year)
    warnings: list[str] = []
    figure, axes = plt.subplots(2, 2, figsize=(12.2, 9.2), dpi=180)
    figure.suptitle(
        f"{result.scope_label} {year} 年季度 NDVI 时序图",
        fontsize=18,
        color="#215b40",
        y=0.985,
    )

    for axis, observation in zip(axes.flat, observations, strict=True):
        map_path, warning = _verified_quarter_map(observation, Path(artifact_root))
        if warning:
            warnings.append(warning)
        if map_path:
            axis.imshow(mpimg.imread(map_path))
        else:
            axis.set_facecolor("#f3f5f3")
            axis.text(
                0.5,
                0.5,
                _placeholder_text(observation),
                ha="center",
                va="center",
                transform=axis.transAxes,
                fontsize=13,
                color="#68736c",
            )
        status_suffix = "（阶段性）" if observation.status == "partial" else ""
        value = f"中位 {observation.median_ndvi:.3f}" if observation.median_ndvi is not None else "无统计值"
        axis.set_title(
            f"{QUARTER_NAMES[observation.window.quarter]}{status_suffix}｜{value}",
            fontsize=11.5,
            color="#315b43",
            pad=8,
        )
        if observation.status != "not_reached":
            coverage = (
                f"有效 {observation.valid_pixel_coverage:.1%}"
                if observation.valid_pixel_coverage is not None
                else "有效覆盖率缺失"
            )
            axis.text(
                0.015,
                0.02,
                f"{observation.image_count} 景｜{coverage}",
                transform=axis.transAxes,
                fontsize=8.5,
                bbox={"facecolor": "white", "alpha": 0.82, "edgecolor": "none", "pad": 2.5},
            )
        axis.set_xticks([])
        axis.set_yticks([])
        for spine in axis.spines.values():
            spine.set_color("#d0d8d2")

    legend_labels = {
        "优": "优  NDVI ≥ 0.75",
        "良": "良  0.60—<0.75",
        "中": "中  0.45—<0.60",
        "一般": "一般  0.30—<0.45",
        "差": "差  NDVI < 0.30",
    }
    legend = [
        Patch(facecolor=color, label=legend_labels[label])
        for _lower, _upper, label, color in reversed(NDVI_LEVELS)
    ]
    figure.legend(
        handles=legend,
        loc="lower center",
        bbox_to_anchor=(0.5, 0.043),
        ncol=5,
        frameon=False,
        fontsize=9,
    )
    figure.text(
        0.5,
        0.012,
        "统一地块范围、统一五级阈值；白线为地块边界。季度合成口径详见报告方法说明。",
        ha="center",
        fontsize=8.8,
        color="#667169",
    )
    figure.tight_layout(rect=(0.015, 0.115, 0.985, 0.95), h_pad=2.2)
    figure.savefig(output_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    _normalise_png_for_office(output_path)
    return output_path, warnings


def _set_run_font(run: Any, size: float = 10.5, *, bold: bool = False, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_page_field(paragraph: Any) -> None:
    from docx.oxml import OxmlElement

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType", "end")
    run._r.extend([begin, instruction, end])


def _add_heading(document: Any, text: str, level: int = 1) -> Any:
    paragraph = document.add_heading(level=level)
    _set_run_font(paragraph.add_run(text), 15 if level == 1 else 12, bold=True, color="215B40")
    return paragraph


def _shade_cell(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _year_summary(observations: list[QuarterlyNDVIObservation], as_of: str) -> str:
    attempted = [item for item in observations if item.status != "not_reached"]
    usable = [item for item in observations if item.usable]
    image_count = sum(item.image_count for item in usable)
    pieces = [
        f"实际到达 {len(attempted)} 个季度，可用 {len(usable)} 个季度",
        f"季度合成共使用 {image_count} 景影像",
    ]
    if usable:
        peak = max(usable, key=lambda item: float(item.median_ndvi))
        low = min(usable, key=lambda item: float(item.median_ndvi))
        pieces.append(f"峰值为 Q{peak.window.quarter} 的 {peak.median_ndvi:.3f}")
        pieces.append(f"最低为 Q{low.window.quarter} 的 {low.median_ndvi:.3f}")
    if any(item.status == "partial" for item in observations):
        pieces.append(f"本年仅统计至 {as_of}；当前季度为阶段性结果，未来季度保持空白")
    return "；".join(pieces) + "。"


def _build_docx(
    result: HistoricalNDVIResult,
    *,
    trend_chart: Path,
    yearly_panels: dict[int, Path],
    report_path: Path,
    render_warnings: list[str],
) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(header.add_run("Agrisky AI｜穹野智保历史 NDVI 遥感筛查"), 8.5, color="68736C")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.add_run("第 "), 8.5, color="68736C")
    _add_page_field(footer)
    _set_run_font(footer.add_run(" 页｜人工复核前不得作为最终定损结论"), 8.5, color="68736C")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(42)
    _set_run_font(
        title.add_run(f"{result.scope_label}\n{result.start_year}—{result.end_year} 年季度 NDVI 时序监测报告"),
        21,
        bold=True,
        color="215B40",
    )
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.add_run(f"统计截至：{result.as_of_date.isoformat()}"), 12, color="587063")

    info = document.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    info_rows = [
        ("任务号", result.task_id),
        ("地块/范围", result.scope_label),
        ("parcel_id / feature_id", f"{result.parcel_id or '—'} / {result.feature_id or '—'}"),
        ("规范化边界 SHA-256", result.boundary_geometry_sha256 or "未绑定（不得作为最终证据）"),
        ("报告状态", "遥感辅助筛查草稿，待人工审核"),
    ]
    for label, value in info_rows:
        cells = info.add_row().cells
        _shade_cell(cells[0], "E9F2EC")
        _set_run_font(cells[0].paragraphs[0].add_run(label), 9.5, bold=True)
        _set_run_font(cells[1].paragraphs[0].add_run(str(value)), 9.5)

    notice = document.add_paragraph()
    notice.paragraph_format.space_before = Pt(12)
    _set_run_font(
        notice.add_run(
            "重要说明：季度 NDVI 描述植被长势筛查信号，不直接等同于灾损等级、减产率或保险赔付比例。"
        ),
        10,
        bold=True,
        color="9C4E18",
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    _add_heading(document, "一、跨年度季度 NDVI 周期曲线")
    document.add_picture(str(trend_chart), width=Cm(16.7))
    caption = document.add_paragraph(
        f"图 1  {result.scope_label} {result.start_year}—{result.end_year} 年季度 NDVI 周期曲线（地块内季度中位值）"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_heading(document, "二、年度季度统计摘要")
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["年度", "到达季度", "可用季度", "影像数", "峰值", "最低值"]
    for index, value in enumerate(headers):
        _shade_cell(table.rows[0].cells[index], "D9EBDD")
        _set_run_font(table.rows[0].cells[index].paragraphs[0].add_run(value), 9, bold=True)
    for year in range(result.start_year, result.end_year + 1):
        observations = _observations_for_year(result, year)
        reached = [item for item in observations if item.status != "not_reached"]
        usable = [item for item in observations if item.usable]
        peak = max(usable, key=lambda item: float(item.median_ndvi)) if usable else None
        low = min(usable, key=lambda item: float(item.median_ndvi)) if usable else None
        values = [
            str(year),
            str(len(reached)),
            str(len(usable)),
            str(sum(item.image_count for item in usable)),
            f"Q{peak.window.quarter} / {peak.median_ndvi:.3f}" if peak else "—",
            f"Q{low.window.quarter} / {low.median_ndvi:.3f}" if low else "—",
        ]
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _set_run_font(cells[index].paragraphs[0].add_run(value), 8.8)

    figure_number = 2
    for year in range(result.start_year, result.end_year + 1):
        document.add_section(WD_SECTION.NEW_PAGE)
        _add_heading(document, f"三.{year - result.start_year + 1}、{year} 年季度 NDVI 时序图")
        document.add_picture(str(yearly_panels[year]), width=Cm(16.7))
        caption = document.add_paragraph(
            f"图 {figure_number}  {result.scope_label} {year} 年 2×2 季度 NDVI 图（统一范围、色标与阈值）"
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        summary = document.add_table(rows=1, cols=1)
        summary.style = "Table Grid"
        _shade_cell(summary.cell(0, 0), "F1F6F2")
        _set_run_font(
            summary.cell(0, 0).paragraphs[0].add_run(
                _year_summary(_observations_for_year(result, year), result.as_of_date.isoformat())
            ),
            9.3,
        )
        figure_number += 1

    document.add_section(WD_SECTION.NEW_PAGE)
    _add_heading(document, "四、数据口径、质量控制与限制")
    methodology = [
        f"数据源：{result.source.collection}；{result.source.formula}。",
        "合成口径：季度内通过质量掩膜的影像逐像元取中位数，再对地块内有效像元取中位数。",
        f"云质量控制：{result.source.cloud_mask}（版本 {result.source.cloud_mask_version}）。",
        "日期口径：季度日期与阶段性截止日均为包含边界；GEE 查询采用下一日排他截止日。",
        "缺失策略：未来、无影像、无有效像元或失败季度均保持空白，不做月均替代或时间插值。",
        f"质量阈值：有效像元覆盖率低于 {LOW_VALID_COVERAGE:.0%} 时必须人工复核。",
    ]
    for text in methodology:
        paragraph = document.add_paragraph(style="List Bullet")
        _set_run_font(paragraph.add_run(text), 10)

    all_warnings = list(dict.fromkeys([*result.warnings, *render_warnings]))
    _add_heading(document, "五、质量警告与人工复核", level=2)
    if all_warnings:
        for warning in all_warnings:
            paragraph = document.add_paragraph(style="List Bullet")
            _set_run_font(paragraph.add_run(warning), 9.5, color="9C4E18")
    else:
        document.add_paragraph("本次自动校验未发现结构性警告；仍需结合农事物候、灾种和现场材料人工复核。")
    document.add_paragraph("审核结论：□ 通过   □ 退回补充影像   □ 退回核对地块   □ 其他：____________")
    document.add_paragraph("审核人：____________    审核日期：____________    备注：________________________")

    properties = document.core_properties
    properties.title = f"{result.scope_label} 历史季度 NDVI 时序监测报告"
    properties.subject = "农业保险 Sentinel-2 NDVI 历史季度遥感筛查"
    properties.keywords = "农业保险, Sentinel-2, NDVI, 季度时序, 遥感"
    properties.comments = f"schema={result.schema_version}; task_id={result.task_id}"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(report_path)


def generate_historical_ndvi_report(
    result: HistoricalNDVIResult,
    output_dir: Path,
    *,
    artifact_root: Path | None = None,
    persist_updated_result: bool = True,
) -> HistoricalNDVIResult:
    """Generate trend chart, yearly panels, DOCX, and optionally refresh JSON."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    artifact_root = Path(artifact_root or output_dir)

    trend_relative = "historical_ndvi_trend.png"
    trend_path = render_quarter_trend_chart(result, output_dir / trend_relative)
    yearly_panels: dict[int, Path] = {}
    render_warnings: list[str] = []
    yearly_relative: dict[str, str] = {}
    for year in range(result.start_year, result.end_year + 1):
        relative = f"yearly_panels/{year}-quarters.png"
        panel_path, panel_warnings = render_year_quarter_panel(
            result,
            year,
            output_dir / relative,
            artifact_root=artifact_root,
        )
        yearly_panels[year] = panel_path
        yearly_relative[str(year)] = relative
        render_warnings.extend(panel_warnings)

    report_relative = "historical_ndvi_report.docx"
    report_path = output_dir / report_relative
    updated = result.model_copy(deep=True)
    updated.outputs = HistoricalNDVIOutputs(
        result_json="result.json",
        trend_chart_png=trend_relative,
        yearly_panel_pngs=yearly_relative,
        report_docx=report_relative,
    )
    updated.warnings = list(dict.fromkeys([*updated.warnings, *render_warnings]))
    _build_docx(
        updated,
        trend_chart=trend_path,
        yearly_panels=yearly_panels,
        report_path=report_path,
        render_warnings=render_warnings,
    )
    if persist_updated_result:
        save_historical_ndvi_result(updated, output_dir / "result.json")
    return updated
